#!/usr/bin/env python3
"""
data_dictionary/data_dictionary_explorer_html_builder.py  (AEP Swiss Army Knife)
==========================================================
PROTOTYPE. Turns a Data Dictionary workbook into ONE web page people can use
without training: a search box, a page for every field, schema, dataset,
audience, journey and source, and lineage you can click along --

    source -> dataset -> schema -> field -> audience -> journey / destination

It only DRAWS what the workbook holds (the workbook stands alone): the lineage
comes from the Data Dictionary's Sources, Journeys and Destinations tabs and
the Audiences tab's 'Fields it reads', 'Builds on' / 'Used by' and 'Where it
goes' columns (v3.4.23). The PQL file beside the workbook supplies only the
full text of rules too long for a cell. No Adobe calls; seconds to build. The
page is a single .html file with everything inside it: no server, no install,
no internet needed. Open it in Edge or Chrome from the synced OneDrive folder
(SharePoint and Teams download .html files rather than show them).

Stdlib + openpyxl only.

Usage:
    python data_dictionary/data_dictionary_explorer_html_builder.py --sandbox=roi-prod
    python data_dictionary/data_dictionary_explorer_html_builder.py --workbook="...xlsx" [--pql="...jsonl"] [--out="...html"]
The newest "Data Dictionary - <sandbox> - *.xlsx" in AEP_OUTPUT_DIR/data_dictionary
(else output/data_dictionary) is used; the page is written next to it as
"Data Dictionary Explorer - <sandbox> - <date>.html".
"""

from __future__ import annotations

import html
import json
import os
import re
import sys
from datetime import datetime
from pathlib import Path

REPO = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(REPO))
from data_dictionary import dictionary_colours as colours  # noqa: E402  -- the workbook's colours

SCRIPT_NAME = "data_dictionary_explorer_html_builder"
SCRIPT_VERSION = "0.2.0"
SCRIPT_DATE = "2026-10-26"
SCRIPT_AUTHOR = "Barry Mann (barrymann.com)"

RAW_PQL_MAX = 60_000        # characters of raw rule kept per audience (the rest is in the PQL file)


def opt(name, default=None):
    return next((a.split("=", 1)[1] for a in sys.argv[1:] if a.startswith(f"--{name}=")), default)


# ----------------------------------------------------------------------------
# Reading the workbook
# ----------------------------------------------------------------------------
def _clean(v):
    if isinstance(v, str):
        v = v.strip()
        return v or None
    return v


def _table(ws, first_header: str) -> list[dict]:
    """Rows of a table sheet as dicts. The header is the first row whose first
    cell is `first_header`; blank rows and one-cell section rows are skipped."""
    rows = ws.iter_rows(values_only=True)
    for r in rows:
        if r and _clean(r[0]) == first_header:
            head = [str(c or "").strip() for c in r]
            break
    else:
        return []
    out = []
    for r in rows:
        vals = [_clean(v) for v in r]
        if sum(v is not None for v in vals) < 2:
            continue
        out.append({h: v for h, v in zip(head, vals) if h})
    return out


def _num(v):
    if v is None or isinstance(v, (int, float)):
        return v
    try:
        f = float(str(v).replace(",", ""))
        return int(f) if f.is_integer() else f
    except ValueError:
        return v


def read_summary(ws) -> dict:
    meta = {"retention": [], "gaps": []}
    lines = [[_clean(v) for v in r] for r in ws.iter_rows(values_only=True)]
    in_gaps = False
    for r in lines:
        a = str(r[0] or "") if r else ""
        if a.startswith("Data Dictionary v"):
            meta["title"] = a
        elif a.startswith("Generated "):
            meta["generated"] = a
        elif "Profile keeps" in a and not a.startswith("PROFILE RETENTION"):
            meta["retention"].append(a)
        elif a.startswith("DATA COMPLETENESS"):
            in_gaps = True
        elif in_gaps and a == "Schema":
            continue
        elif in_gaps and a and r[3]:
            meta["gaps"].append({"schema": a, "tab": r[1], "status": r[3]})
        elif in_gaps and not a:
            in_gaps = False
    return meta


def read_workbook(path: Path) -> dict:
    from openpyxl import load_workbook
    wb = load_workbook(path, read_only=True)
    meta = read_summary(wb["Summary"])
    schemas = _table(wb["Schemas"], "Tab")
    datasets = _table(wb["Datasets"], "Sandbox")
    fields = _table(wb["Field Index"], "Field (dot notation)")
    hist_ws = wb["Field History"] if "Field History" in wb.sheetnames else None
    hist = _table(hist_ws, "Field (dot notation)") if hist_ws else []
    hist_head = []
    if hist_ws:
        for r in hist_ws.iter_rows(values_only=True):
            if r and r[0] == "Field (dot notation)":
                hist_head = [str(c) for c in r[6:] if c]
                break
    audiences = _table(wb["Audiences"], "Sandbox") if "Audiences" in wb.sheetnames else []
    lineage = {n: _table(wb[n], "Sandbox") if n in wb.sheetnames else []
               for n in ("Sources", "Journeys", "Destinations")}
    tabs = {}
    for s in schemas:
        tab = s.get("Tab")
        if tab in wb.sheetnames:
            tabs[tab] = {r.get("Field (dot notation)"): r for r in _table(wb[tab], "Field (dot notation)")}
    wb.close()
    return {"meta": meta, "schemas": schemas, "datasets": datasets, "fields": fields,
            "hist": hist, "months": hist_head, "audiences": audiences, "tabs": tabs, **lineage}


def read_pql(path: Path | None) -> dict:
    if not path or not path.exists():
        return {}
    out = {}
    for line in path.read_text(encoding="utf-8").splitlines():
        if line.strip():
            r = json.loads(line)
            out[r.get("audience_id")] = r
    return out


# ----------------------------------------------------------------------------
# The model the page draws
# ----------------------------------------------------------------------------
def _split(v, sep) -> list[str]:
    return [x.strip() for x in str(v or "").split(sep) if x.strip()]


def _yes(v) -> bool:
    return str(v or "").strip().lower() in ("yes", "y", "true")


def build_model(wbk: dict, pql: dict, workbook_name: str) -> dict:
    schemas, s_ix = [], {}
    for r in wbk["schemas"]:
        title = str(r.get("Schema") or "").strip()
        s_ix[title] = len(schemas)
        schemas.append({
            "tab": r.get("Tab"), "title": title, "name": re.sub(r"\s+schema$", "", title, flags=re.I) or title,
            "desc": r.get("Description"),
            "cls": r.get("Class"), "kind": r.get("Kind"), "profile": _yes(r.get("In Profile")),
            "gb": _num(r.get("Data lake GB")), "rows": _num(r.get("Data lake rows")),
            "ttl": _num(r.get("Data lake TTL (days)")), "pgb": _num(r.get("Profile GB")),
            "pttl": _num(r.get("Profile TTL (days)")), "tables": r.get("SQL table names"),
            "nid": _num(r.get("Identities")), "nrel": _num(r.get("Relationships")),
            "mod": r.get("Last Modified"), "aa": _yes(r.get("Adobe Analytics")), "ds": [], "f": []})

    datasets = []
    for r in wbk["datasets"]:
        s = s_ix.get(str(r.get("Schema") or "").strip(), -1)
        if s >= 0:
            schemas[s]["ds"].append(len(datasets))
        datasets.append({
            "name": r.get("Friendly Name (dataset)"), "table": r.get("Table Name (SQL / system)"),
            "desc": r.get("Description"), "schema": r.get("Schema"), "s": s, "kind": r.get("Kind"),
            "profile": r.get("Profile"), "gb": _num(r.get("Data lake GB")),
            "rows": _num(r.get("Data lake rows")), "ttl": _num(r.get("Data lake TTL (days)")),
            "pgb": _num(r.get("Profile GB")), "pttl": _num(r.get("Profile TTL (days)")),
            "id": r.get("Dataset ID"), "fed": r.get("Fed by"), "src": []})

    month_col = next((k for k in (wbk["fields"][0] if wbk["fields"] else {}) if k.startswith("Month by month")), None)
    months = wbk["months"]
    hist = {(h.get("Field (dot notation)"), str(h.get("Schema") or "").strip()): h for h in wbk["hist"]}
    fields, f_ix = [], {}
    for r in wbk["fields"]:
        path = r.get("Field (dot notation)")
        title = str(r.get("Schema") or "").strip()
        s = s_ix.get(title, -1)
        if not path or s < 0:
            continue
        if path not in f_ix:
            f_ix[path] = len(fields)
            fields.append({"path": path, "name": r.get("Friendly Name"), "fpath": r.get("Friendly Path"),
                           "desc": r.get("Description"), "type": r.get("Data Type"), "on": [], "aud": []})
        f = fields[f_ix[path]]
        f["name"] = f["name"] or r.get("Friendly Name")
        f["desc"] = f["desc"] or r.get("Description")
        h = hist.get((path, title), {})
        tab_row = wbk["tabs"].get(schemas[s]["tab"], {}).get(path, {})
        f["on"].append({
            "s": s, "ident": r.get("Identity"), "req": _yes(tab_row.get("Required")),
            "rel": tab_row.get("Relationship → target"),
            "cov": _num(r.get("Coverage %")), "top": r.get("Top values (count)"),
            "months": _num(r.get("Have we got it? (months)")), "hdet": r.get("History detail"),
            "pdays": _num(r.get("In Profile now (days)")), "pdet": r.get("Profile detail"),
            "hist": [h.get(m) for m in months] if h else [],
            "spark": r.get(month_col) if month_col else None,
            "sql": r.get("SQL to check it (paste into Query Service)")})
        if f_ix[path] not in schemas[s]["f"]:
            schemas[s]["f"].append(f_ix[path])

    audiences, a_ix = [], {}
    for r in wbk["audiences"]:
        aid = r.get("Audience id")
        side = pql.get(aid, {})
        raw = side.get("raw_pql") or r.get("PQL (raw pql/json)") or ""
        prof, evt = _split(r.get("Profile fields it reads"), ";"), _split(r.get("Event fields it reads"), ";")
        if not prof and not evt:            # one column, kind not recorded: either
            prof = evt = _split(r.get("Fields it reads"), ";")
        reads = list(dict.fromkeys(prof + evt))
        a_ix[aid] = len(audiences)
        audiences.append({
            "id": aid, "name": r.get("Audience name"), "desc": r.get("Description"),
            "eval": r.get("Evaluation"), "life": r.get("Lifecycle"), "tags": r.get("Tags"),
            "origin": r.get("Origin"), "by": r.get("Created by"), "mod": r.get("Last modified by"),
            "rule": r.get("PQL (readable)"), "cut": _yes(r.get("PQL truncated in sheet")),
            "raw": raw[:RAW_PQL_MAX], "rawcut": len(raw) > RAW_PQL_MAX,
            "fmt": side.get("format", r.get("PQL format")) or "",
            "score": _num(r.get("Complexity score (0-100)")), "rag": r.get("Complexity RAG"),
            "why": r.get("Why this score"), "lookback": _num(r.get("Lookback (days)")),
            "base": r.get("Largest base audience"), "basen": _num(r.get("Largest base (profiles)")),
            "where": r.get("Where it goes"),
            "f": sorted({f_ix[p] for p in reads if p in f_ix}), "miss": [p for p in reads if p not in f_ix],
            "fp": sorted({f_ix[p] for p in prof if p in f_ix}), "fe": sorted({f_ix[p] for p in evt if p in f_ix}),
            "depids": _split(r.get("Builds on (audience ids)"), ","),
            "usedids": _split(r.get("Used by (audience ids)"), ","), "j": [], "x": []})
    for a in audiences:
        a["dep"] = [a_ix[d] for d in a.pop("depids") if d in a_ix]
        a["used"] = [a_ix[u] for u in a.pop("usedids") if u in a_ix]
    for i, a in enumerate(audiences):
        for f in a["f"]:
            fields[f]["aud"].append(i)

    journeys = []
    for r in wbk["Journeys"]:
        auds = [a_ix[x] for x in _split(r.get("Audience ids"), ",") if x in a_ix]
        for a in auds:
            audiences[a]["j"].append(len(journeys))
        journeys.append({"id": r.get("Journey id"), "name": r.get("Journey"), "version": r.get("Version"),
                         "status": r.get("Status"), "live": _yes(r.get("Live now")), "how": r.get("How"),
                         "reads": r.get("Audiences it reads"), "changed": r.get("Last changed"),
                         "by": r.get("Changed by"), "aud": auds})
    destinations = []
    for r in wbk["Destinations"]:
        auds = [a_ix[x] for x in _split(r.get("Audience ids"), ",") if x in a_ix]
        for a in auds:
            audiences[a]["x"].append(len(destinations))
        destinations.append({"id": r.get("Flow id"), "dest": r.get("Destination"), "name": r.get("Dataflow"),
                             "live": _yes(r.get("Live now")), "aud": auds})
    by_table = {d["table"]: i for i, d in enumerate(datasets) if d["table"]}
    sources = []
    for r in wbk["Sources"]:
        into = [by_table[x] for x in _split(r.get("Into dataset"), ",") if x in by_table]
        for d in into:
            datasets[d]["src"].append(len(sources))
        sources.append({"id": r.get("Flow id"), "kind": r.get("How it comes in"), "name": r.get("Dataflow"),
                        "live": _yes(r.get("Live now")), "into": into, "profile": _yes(r.get("Into Profile")),
                        "from": r.get("From"), "last": r.get("Last run (UTC)"), "state": r.get("Last run"),
                        "by": r.get("Created by")})

    meta = wbk["meta"]
    title = meta.get("title") or ""
    label = title.split("–", 1)[1].strip() if "–" in title else ""
    return {"meta": {"title": title, "label": label, "generated": meta.get("generated"),
                     "workbook": workbook_name, "months": months, "retention": meta["retention"],
                     "gaps": meta["gaps"], "built": datetime.now().strftime("%Y-%m-%d %H:%M"),
                     "builder": f"{SCRIPT_NAME} v{SCRIPT_VERSION}"},
            "schemas": schemas, "datasets": datasets, "fields": fields, "audiences": audiences,
            "journeys": journeys, "destinations": destinations, "sources": sources}


def palette_css() -> str:
    """CSS variables for the workbook's colour scheme (dictionary_colours), the
    same in light and dark: a pill looks the same as its cell in the workbook."""
    v = {"--header": colours.HEADER, "--link": colours.LINK}
    v.update({f"--k-{k}": colours.KIND[k][0] for k in colours.KIND})
    v.update({"--k-aud": colours.PILL["audience"][1], "--k-jrn": colours.PILL["journey"][1],
              "--k-dst": colours.PILL["destination"][1], "--k-src": colours.PILL["source"][1],
              "--k-ds": colours.PILL["neutral"][1]})
    for name, (bg, fg) in {**colours.PILL, **{k.lower(): c for k, c in colours.RAG.items()}}.items():
        v[f"--p-{name}-bg"], v[f"--p-{name}"] = bg, fg
    return " ".join(f"{k}:#{c};" for k, c in v.items())


def render_html(model: dict) -> str:
    model["meta"]["palette"] = {"coverage": colours.COVERAGE_SCALE,
                                "live": sorted(colours.LIVE_JOURNEY), "trying": sorted(colours.TRYING)}
    data = json.dumps(model, ensure_ascii=False, separators=(",", ":"), default=str)
    data = data.replace("</", "<\\/")
    label = model["meta"]["label"] or "Data Dictionary"
    return (PAGE.replace("{{TITLE}}", html.escape(f"Data Dictionary Explorer – {label}"))
                .replace("{{AUTHOR}}", html.escape(SCRIPT_AUTHOR))
                .replace("/*{{PALETTE}}*/", palette_css())
                .replace("/*{{DATA}}*/", data))


def main():
    if "-h" in sys.argv[1:] or "--help" in sys.argv[1:] or not (opt("sandbox") or opt("workbook")):
        print(__doc__)
        return
    out_dir = Path(os.environ.get("AEP_OUTPUT_DIR") or REPO / "output") / "data_dictionary"
    sandbox = opt("sandbox")
    wb_path = Path(opt("workbook")) if opt("workbook") else max(
        out_dir.glob(f"Data Dictionary - {sandbox} - *.xlsx"), key=lambda p: p.stat().st_mtime, default=None)
    if not wb_path or not wb_path.exists():
        sys.exit(f"No Data Dictionary workbook for {sandbox} in {out_dir} -- run the Data Dictionary first.")
    m = re.match(r"Data Dictionary - (.+) - (\d{4}-\d{2}-\d{2})", wb_path.stem)
    sandbox = sandbox or (m.group(1) if m else "")
    pql_path = Path(opt("pql")) if opt("pql") else (
        wb_path.parent / "pql" / f"audiences_pql_{sandbox}_{m.group(2).replace('-', '')}.jsonl" if m else None)
    out = Path(opt("out")) if opt("out") else wb_path.with_name(
        wb_path.name.replace("Data Dictionary - ", "Data Dictionary Explorer - ", 1)).with_suffix(".html")

    print(f"{SCRIPT_NAME} v{SCRIPT_VERSION} ({SCRIPT_DATE})  by {SCRIPT_AUTHOR}")
    print(f"workbook {wb_path.name}  |  PQL {pql_path.name if pql_path and pql_path.exists() else '(none)'}")
    model = build_model(read_workbook(wb_path), read_pql(pql_path), wb_path.name)
    page = render_html(model)
    out.write_text(page, encoding="utf-8")
    linked = sum(1 for a in model["audiences"] if a["f"])
    print(f"{len(model['schemas'])} schemas, {len(model['datasets'])} datasets, {len(model['fields'])} fields, "
          f"{len(model['audiences'])} audiences ({linked} linked to the fields they read), "
          f"{len(model['journeys'])} journeys, {len(model['destinations'])} destinations, "
          f"{len(model['sources'])} sources")
    print(f"written: {out}  ({len(page.encode('utf-8')) / 1e6:.1f} MB)")


# ----------------------------------------------------------------------------
# The page. Plain HTML, CSS and JavaScript; the data goes where /*{{DATA}}*/ is.
# ----------------------------------------------------------------------------
PAGE = r"""<!doctype html>
<html lang="en-GB">
<head>
<meta charset="utf-8">
<meta name="viewport" content="width=device-width, initial-scale=1">
<meta name="author" content="{{AUTHOR}}">
<title>{{TITLE}}</title>
<style>
:root{
  --bg:#f5f6f8; --panel:#ffffff; --ink:#1c2230; --muted:#5d6678; --faint:#8a93a3; --line:#e2e5eb;
  --accent:#0563C1; --accent-weak:#e9effc; --focus:#ffd257;
  --k-profile:#7a4cc2; --k-event:#2a72c8; --k-other:#21857a; --k-aud:#c4572a; --k-ds:#5d6678;
  --k-jrn:#0e7c86; --k-dst:#a23b72; --k-src:#8a6a12;
  --m-none:#eceef2; --m-unk:#d5d9e0; --m-zero:#f6d4c8; --m-hi:#2e9a57;
  --edge:#b8c0cc; --edge-hot:#1F4E78; --code-bg:#f2f4f7;
  --mono:Consolas,"Cascadia Mono",ui-monospace,monospace;
  --sans:"Segoe UI",system-ui,-apple-system,Roboto,Arial,sans-serif;
  /*{{PALETTE}}*/
}
@media (prefers-color-scheme: dark){
  :root:not([data-theme="light"]){
    --bg:#14171d; --panel:#1c2028; --ink:#e6e9ef; --muted:#a5adbb; --faint:#7d8696; --line:#2e3440;
    --accent:#7aa2ff; --accent-weak:#22304d; --focus:#8a6d1f;
    --m-none:#262b35; --m-unk:#3a404c; --m-zero:#5a3328; --m-hi:#4cc27a;
    --edge:#465063; --edge-hot:#7aa2ff; --code-bg:#232833;
  }
}
:root[data-theme="dark"]{
    --bg:#14171d; --panel:#1c2028; --ink:#e6e9ef; --muted:#a5adbb; --faint:#7d8696; --line:#2e3440;
    --accent:#7aa2ff; --accent-weak:#22304d; --focus:#8a6d1f;
    --m-none:#262b35; --m-unk:#3a404c; --m-zero:#5a3328; --m-hi:#4cc27a;
    --edge:#465063; --edge-hot:#7aa2ff; --code-bg:#232833;
}
*{box-sizing:border-box}
html,body{margin:0;background:var(--bg);color:var(--ink);font:15px/1.5 var(--sans)}
a{color:var(--accent);text-decoration:none} a:hover{text-decoration:underline}
code,.mono{font-family:var(--mono);font-size:13px}
header.top{position:sticky;top:0;z-index:20;background:var(--panel);border-bottom:4px solid var(--header)}
.bar{max-width:1280px;margin:0 auto;padding:10px 20px;display:flex;gap:18px;align-items:center;flex-wrap:wrap}
.brand{font-weight:700;white-space:nowrap;color:var(--ink)} .brand small{display:block;font-weight:400;color:var(--muted);font-size:12px}
.proto{font-size:11px;background:var(--p-amber-bg);color:var(--p-amber);border-radius:10px;padding:1px 8px;margin-left:6px;font-weight:600;vertical-align:middle}
.searchwrap{position:relative;flex:1;min-width:240px;max-width:560px}
.searchwrap input{width:100%;font:inherit;padding:8px 12px 8px 34px;border:1px solid var(--line);border-radius:8px;background:var(--bg);color:var(--ink)}
.searchwrap input:focus{outline:2px solid var(--accent);border-color:transparent}
.searchwrap svg{position:absolute;left:10px;top:50%;transform:translateY(-50%);color:var(--faint)}
.drop{position:absolute;left:0;right:0;top:calc(100% + 4px);background:var(--panel);border:1px solid var(--line);border-radius:8px;box-shadow:0 8px 24px rgba(0,0,0,.12);max-height:70vh;overflow:auto;display:none}
.drop a{display:flex;gap:10px;padding:7px 12px;color:var(--ink);align-items:baseline} .drop a.sel,.drop a:hover{background:var(--accent-weak);text-decoration:none}
.drop .sub{color:var(--muted);font-size:12px;overflow:hidden;text-overflow:ellipsis;white-space:nowrap}
.drop .all{border-top:1px solid var(--line);color:var(--accent);font-size:13px}
nav.main{display:flex;gap:4px;flex-wrap:wrap} nav.main a{padding:6px 10px;border-radius:6px;color:var(--muted);font-size:14px} nav.main a.on,nav.main a:hover{background:var(--accent-weak);color:var(--accent);text-decoration:none}
main{max-width:1280px;margin:0 auto;padding:22px 20px 60px;overflow-wrap:anywhere}
h1{font-size:26px;margin:4px 0 6px;line-height:1.25} h2{font-size:18px;margin:30px 0 10px} h3{font-size:15px;margin:18px 0 8px}
.lede{color:var(--muted);max-width:820px;margin:0 0 14px}
.crumb{font-size:13px;color:var(--muted);margin-bottom:2px}
.panel{background:var(--panel);border:1px solid var(--line);border-radius:10px;padding:16px 18px}
.grid{display:grid;gap:14px} .g4{grid-template-columns:repeat(auto-fit,minmax(220px,1fr))} .g3{grid-template-columns:repeat(auto-fit,minmax(280px,1fr))} .g2{grid-template-columns:repeat(auto-fit,minmax(380px,1fr))}
.tile{display:block;color:var(--ink)} .tile:hover{text-decoration:none;border-color:var(--accent)}
.tile .n{font-size:30px;font-weight:700;line-height:1.1} .tile .l{color:var(--muted);font-size:13px}
.card h3{margin-top:0} .card p{margin:4px 0 10px;color:var(--muted);font-size:14px}
.facts{display:grid;grid-template-columns:repeat(auto-fit,minmax(170px,1fr));gap:10px 18px;margin:12px 0 4px}
.facts div{font-size:13px;color:var(--muted)} .facts b{display:block;color:var(--ink);font-size:15px;font-weight:600;word-break:break-word}
.badge{display:inline-block;font-size:12px;font-weight:600;border-radius:10px;padding:1px 9px;margin-right:4px;vertical-align:middle;white-space:nowrap}
.badge{font-weight:700}
.b-profile{background:var(--p-profile-bg);color:var(--p-profile)} .b-event{background:var(--p-event-bg);color:var(--p-event)}
.b-other{background:var(--p-other-bg);color:var(--p-other)} .b-aud{background:var(--p-audience-bg);color:var(--p-audience)}
.b-in{background:var(--p-green-bg);color:var(--p-green)} .b-out{background:var(--p-grey-bg);color:var(--p-grey)}
.RED{background:var(--p-red-bg);color:var(--p-red)} .AMBER{background:var(--p-amber-bg);color:var(--p-amber)} .GREEN{background:var(--p-green-bg);color:var(--p-green)} .GREY{background:var(--p-grey-bg);color:var(--p-grey)}
.say{font-size:16px;margin:10px 0;max-width:900px} .say b{font-weight:600}
.big{font-size:28px;font-weight:700;line-height:1} .big small{font-size:14px;font-weight:400;color:var(--muted)}
.copyrow{display:flex;gap:8px;align-items:flex-start}
pre.code{flex:1;min-width:0;overflow-wrap:anywhere;margin:0;background:var(--code-bg);border:1px solid var(--line);border-radius:6px;padding:8px 10px;white-space:pre-wrap;word-break:break-word;font:12.5px/1.45 var(--mono);max-height:340px;overflow:auto}
#theme,#aa{white-space:nowrap} #aa.on{background:var(--p-neutral-bg);color:var(--p-neutral);font-weight:700}
button{font:inherit;font-size:13px;border:1px solid var(--line);background:var(--panel);color:var(--ink);border-radius:6px;padding:5px 10px;cursor:pointer} button:hover{border-color:var(--accent);color:var(--accent)}
details summary{cursor:pointer;color:var(--accent);font-size:14px;margin:6px 0}
table.t{width:100%;border-collapse:collapse;font-size:14px;background:var(--panel)}
table.t th{position:sticky;top:0;background:var(--header);text-align:left;font-size:12px;color:#fff;font-weight:600;border-bottom:2px solid var(--header);padding:7px 8px;cursor:pointer;white-space:nowrap;user-select:none}
table.t th.num,table.t td.num{text-align:right} table.t th:hover{color:#FFE699} table.t th .arr{opacity:.8}
table.t td{border-bottom:1px solid var(--line);padding:7px 8px;vertical-align:top}
table.t tr:hover td{background:color-mix(in srgb,var(--accent-weak) 45%,transparent)}
table.t td .sub{color:var(--muted);font-size:12px;word-break:break-all}
.tablewrap{border:1px solid var(--line);border-radius:10px;overflow:auto;max-height:75vh}
.tools{display:flex;gap:8px;align-items:center;flex-wrap:wrap;margin:8px 0 10px}
.tools input{font:inherit;font-size:14px;padding:6px 10px;border:1px solid var(--line);border-radius:6px;background:var(--panel);color:var(--ink);min-width:260px}
.chip{font-size:13px;border:1px solid var(--line);border-radius:14px;padding:3px 11px;background:var(--panel);cursor:pointer;color:var(--muted)} .chip.on{background:var(--accent-weak);border-color:var(--accent);color:var(--accent)}
.count{color:var(--muted);font-size:13px}
.covbar{display:inline-block;width:70px;height:8px;border-radius:4px;background:var(--m-none);vertical-align:middle;margin-right:6px;overflow:hidden} .covbar i{display:block;height:100%;background:var(--m-hi)}
/* month strip */
.strip{display:inline-grid;grid-auto-flow:column;grid-auto-columns:15px;gap:2px;vertical-align:middle}
.strip i{height:18px;border-radius:2px;background:var(--m-none);font-style:normal;font-size:10px;line-height:18px;text-align:center;color:var(--muted)}
.strip i.u{background:var(--m-unk)} .strip i.z{background:var(--m-zero)}
.striplab{display:flex;justify-content:space-between;font-size:11px;color:var(--faint);margin-top:2px}
.legend{display:flex;gap:14px;flex-wrap:wrap;font-size:12px;color:var(--muted);margin-top:8px} .legend span i{display:inline-block;width:12px;height:12px;border-radius:2px;vertical-align:-2px;margin-right:4px}
/* lineage */
.linwrap{position:relative;overflow-x:auto;padding:6px 2px 10px}
.lin{position:relative;display:grid;grid-auto-flow:column;grid-auto-columns:minmax(165px,1fr);column-gap:44px;min-width:min-content}
.lin svg{position:absolute;left:0;top:0;pointer-events:none;overflow:visible}
.lin .col{position:relative;z-index:1;min-width:0}
.lin h4{margin:0 0 8px;font-size:11px;letter-spacing:.06em;text-transform:uppercase;color:var(--faint);font-weight:700}
.node{display:block;position:relative;background:var(--panel);border:1px solid var(--line);border-left:4px solid var(--k-ds);border-radius:7px;padding:5px 9px;margin-bottom:7px;color:var(--ink);font-size:13px;line-height:1.3}
.node:hover{text-decoration:none;border-color:var(--accent);box-shadow:0 2px 8px rgba(0,0,0,.08)}
.node .l{display:block;overflow:hidden;text-overflow:ellipsis;white-space:nowrap;padding-right:10px}
.node .s{display:block;color:var(--muted);font-size:11px;overflow:hidden;text-overflow:ellipsis;white-space:nowrap}
.vals.mini{margin-top:0;gap:3px;max-width:340px} .vals.mini span{font-size:11px;padding:0 5px;max-width:150px;overflow:hidden;text-overflow:ellipsis;white-space:nowrap} .vals.mini em{font-size:11px;color:var(--faint);font-style:normal;align-self:center}
.vals{display:flex;flex-wrap:wrap;gap:4px;margin-top:6px} .vals span{font-size:12px;background:var(--code-bg);border:1px solid var(--line);border-radius:5px;padding:1px 7px;font-family:var(--mono)} .vals b{font-family:var(--sans);font-weight:400;color:var(--muted);margin-left:5px}
.node.k-jrn{border-left-color:var(--k-jrn)} .node.k-dst{border-left-color:var(--k-dst)} .node.k-src{border-left-color:var(--k-src)}
.b-jrn{background:var(--p-journey-bg);color:var(--p-journey)} .b-dst{background:var(--p-destination-bg);color:var(--p-destination)} .b-src{background:var(--p-source-bg);color:var(--p-source)} .b-neutral{background:var(--p-neutral-bg);color:var(--p-neutral)}
.node.k-profile{border-left-color:var(--k-profile)} .node.k-event{border-left-color:var(--k-event)} .node.k-other{border-left-color:var(--k-other)} .node.k-aud{border-left-color:var(--k-aud)}
.node.focus{border:2px solid var(--accent);border-left-width:5px;background:var(--accent-weak);font-weight:600}
.node .r{position:absolute;right:6px;top:5px;width:8px;height:8px;border-radius:50%}
.node .r.RED{background:var(--p-red)} .node .r.AMBER{background:#C9A100} .node .r.GREEN{background:#2E9A57}
.node.dim{opacity:.55}
.col .more{display:block;font-size:12px;padding:2px 4px}
.col .empty{font-size:12px;color:var(--faint);border:1px dashed var(--line);border-radius:7px;padding:8px 9px}
.col.ph .empty{background:repeating-linear-gradient(45deg,transparent 0 6px,color-mix(in srgb,var(--line) 50%,transparent) 6px 12px)}
path.e{fill:none;stroke:var(--edge);stroke-width:1.4} path.e.hot{stroke:var(--edge-hot);stroke-width:2.2}
.hint{font-size:12px;color:var(--faint);margin-top:4px}
.term{border-bottom:1px dotted var(--faint);cursor:help}
.foot{color:var(--faint);font-size:12px;margin-top:40px;border-top:1px solid var(--line);padding-top:12px}
.warn{background:var(--p-amber-bg);color:var(--ink);border-radius:8px;padding:10px 14px;font-size:14px;margin:10px 0}
.kv{display:grid;grid-template-columns:max-content 1fr;gap:4px 14px;font-size:14px} .kv dt{color:var(--muted)} .kv dd{margin:0}
.hero{padding:10px 0 6px} .hero .searchbig{max-width:720px;margin-top:12px}
.hero .searchbig input{font-size:18px;padding:13px 16px 13px 42px;border-radius:10px}
.hero .searchbig svg{left:14px}
.list-plain{margin:0;padding-left:18px} .list-plain li{margin:3px 0}
@media (max-width:760px){
  .lin{grid-auto-flow:row;grid-auto-columns:auto;row-gap:18px;min-width:0}
  .lin svg{display:none}
  .lin .col:not(:last-child)::after{content:"↓";display:block;text-align:center;color:var(--faint);font-size:18px}
  .bar{padding:8px 12px} main{padding:16px 12px 50px}
  .g2{grid-template-columns:1fr}
}
@media print{header.top{position:static} .drop,button{display:none}}
</style>
</head>
<body>
<header class="top"><div class="bar">
  <a class="brand" href="#/">Data Dictionary Explorer<span class="proto">PROTOTYPE</span><small id="brand-sub"></small></a>
  <div class="searchwrap" id="topsearch"></div>
  <nav class="main" id="nav">
    <a href="#/">Home</a><a href="#/fields">Fields</a><a href="#/schemas">Schemas</a><a href="#/datasets">Datasets</a><a href="#/audiences">Audiences</a><a href="#/journeys">Journeys</a><a href="#/sources">Sources</a><a href="#/help">How to read this</a>
  </nav>
  <button id="aa" hidden title="Adobe Analytics report-suite data: in the workbook, hidden here by default">AA hidden</button>
  <button id="theme" title="Theme: auto (click for light, dark, auto)">◐</button>
</div></header>
<main id="app"></main>
<script type="application/json" id="dd-data">/*{{DATA}}*/</script>
<script>
"use strict";
const D = JSON.parse(document.getElementById("dd-data").textContent);
const M = D.meta;
const app = document.getElementById("app");
const h = s => String(s ?? "").replace(/[&<>"']/g, c => ({"&":"&amp;","<":"&lt;",">":"&gt;",'"':"&quot;","'":"&#39;"}[c]));
const num = v => (v === null || v === undefined || v === "") ? null : (typeof v === "number" ? v : (isNaN(+v) ? null : +v));
const fmt = (v, dp) => { const n = num(v); return n === null ? "–" : n.toLocaleString("en-GB", {maximumFractionDigits: dp ?? 1}); };
const plural = (n, one, many) => `${fmt(n, 0)} ${n === 1 ? one : (many || one + "s")}`;
const enc = encodeURIComponent;
const SEARCH_ICON = '<svg width="16" height="16" viewBox="0 0 16 16" fill="none" stroke="currentColor" stroke-width="1.8"><circle cx="7" cy="7" r="5"/><path d="M11 11l3.5 3.5"/></svg>';

/* ---------- lookups ---------- */
const by = {field: {}, schema: {}, dataset: {}, audience: {}, journey: {}, source: {}, destination: {}};
const keyOf = (x, i) => x.id || "n" + i;
D.journeys.forEach((x, i) => by.journey[keyOf(x, i)] = i);
D.sources.forEach((x, i) => by.source[keyOf(x, i)] = i);
D.destinations.forEach((x, i) => by.destination[keyOf(x, i)] = i);
D.fields.forEach((f, i) => by.field[f.path] = i);
D.schemas.forEach((s, i) => by.schema[s.tab] = i);
D.datasets.forEach((d, i) => by.dataset[d.table || d.id] = i);
D.audiences.forEach((a, i) => by.audience[a.id] = i);
const audByName = {};
D.audiences.forEach((a, i) => { if (!(a.name in audByName)) audByName[a.name] = i; });
const href = {
  f: i => "#/field/" + enc(D.fields[i].path),
  s: i => "#/schema/" + enc(D.schemas[i].tab),
  d: i => "#/dataset/" + enc(D.datasets[i].table || D.datasets[i].id),
  a: i => "#/audience/" + enc(D.audiences[i].id),
  j: i => "#/journey/" + enc(keyOf(D.journeys[i], i)),
  r: i => "#/source/" + enc(keyOf(D.sources[i], i)),
  x: i => "#/destination/" + enc(keyOf(D.destinations[i], i)),
};
// The workbook's lineage tabs (Data Dictionary v3.4.23); older workbooks have none.
const HAS_LINEAGE = !!(D.journeys.length || D.sources.length || D.destinations.length);
const NOT_IN = "Not in this workbook – it needs a Data Dictionary run from v3.4.23.";
const liveWord = live => live ? '<span class="badge b-in">Live</span>' : '<span class="badge b-out">Not live</span>';
const WHERE_RAG = w => /^In use/.test(w || "") ? "GREEN" : /^Not live/.test(w || "") ? "AMBER" : /^Nothing found/.test(w || "") ? "RED" : "";
const whereBadge = a => a.where ? `<span class="badge ${WHERE_RAG(a.where)}">${h(a.where)}</span>` : "";
const kindOf = k => ({Profile: "profile", Event: "event"})[k] || "other";
const kindWord = k => ({Profile: "Profile record", Event: "Event"})[k] || "Lookup / other";
const fieldKind = f => kindOf(D.schemas[f.on[0].s].kind);
const fieldLabel = f => f.name || f.path.split(".").pop().replace("[]", "");
const profileWord = p => p ? '<span class="badge b-in">In Profile</span>' : '<span class="badge b-out">Not in Profile</span>';
const dsInProfile = d => /^profile/i.test(String(d.profile || ""));
const TYPE_WORDS = {"string": "text", "int": "whole number", "long": "whole number", "integer": "whole number",
  "number": "number", "double": "number", "boolean": "yes / no", "date": "date", "date-time": "date and time",
  "map": "set of named values", "object": "group of fields"};
const typeWords = t => !t ? "–" : (TYPE_WORDS[t] || (/^array<(.+)>$/.test(t) ? "list of " + (TYPE_WORDS[t.slice(6, -1)] || t.slice(6, -1)) : t));
const GLOSS = {
  schema: "The blueprint for one kind of data: which fields it has and what type each is.",
  dataset: "Where data of one schema is actually stored. A schema can have several datasets.",
  field: "One piece of information, such as an email address or an order total.",
  profile: "Unified Profile: the live, merged view of each customer that audiences are built from.",
  lake: "The data lake: Adobe's long-term store, which SQL (Query Service) reads.",
  audience: "A rule that picks customers, e.g. \"ordered in the last 30 days\".",
  coverage: "Out of a sample of recent records, the share that had this field filled in.",
};
const term = (word, key) => `<span class="term" title="${h(GLOSS[key])}">${h(word)}</span>`;

/* ---------- Adobe Analytics report-suite data: hidden unless asked for ---------- */
// The workbook flags the schemas (Schemas tab, 'Adobe Analytics'); hidden with them are
// fields found only there, their datasets, and the dataflows that feed only those.
const AA_S = new Set(D.schemas.flatMap((s, i) => s.aa ? [i] : []));
let SHOW_AA = false;
try { SHOW_AA = localStorage.getItem("dd-aa") === "show"; } catch (e) { }
const aaDs = d => AA_S.has(D.datasets[d].s);
const HIDE = {
  s: i => AA_S.has(i),
  f: i => D.fields[i].on.length > 0 && D.fields[i].on.every(o => AA_S.has(o.s)),
  d: i => aaDs(i),
  r: i => /^Adobe Analytics/i.test(D.sources[i].kind || "") || (D.sources[i].into.length > 0 && D.sources[i].into.every(aaDs)),
};
const shown = (t, i) => SHOW_AA || !(HIDE[t] && HIDE[t](i));
const showIdx = (t, xs) => xs.filter(i => shown(t, i));
const showRows = (t, list) => list.filter((_, i) => shown(t, i));
// A field's schemas, without the hidden ones -- unless that would leave none.
const visOn = f => { const o = SHOW_AA ? f.on : f.on.filter(x => !AA_S.has(x.s)); return o.length ? o : f.on; };
function aaNote(what) {
  if (!AA_S.size) return "";
  return SHOW_AA ? ` <a href="#" data-aa="hide">Hide Adobe Analytics</a>.`
    : ` Adobe Analytics report-suite ${what} are hidden (they are in the workbook) – <a href="#" data-aa="show">show them</a>.`;
}

/* ---------- search ---------- */
const IDX = [];
D.fields.forEach((f, i) => IDX.push({t: "f", i, title: fieldLabel(f), sub: f.path,
  hay: [f.name, f.path, f.fpath, f.desc, ...f.on.map(o => D.schemas[o.s].title), ...f.on.map(o => o.top)].join(" ").toLowerCase()}));
D.schemas.forEach((s, i) => IDX.push({t: "s", i, title: s.name, sub: `${kindWord(s.kind)} schema · ${plural(s.f.length, "field")}`,
  hay: [s.tab, s.title, s.desc, s.cls, s.tables].join(" ").toLowerCase()}));
D.datasets.forEach((d, i) => IDX.push({t: "d", i, title: d.name || d.table, sub: `dataset · ${d.table || ""}`,
  hay: [d.name, d.table, d.desc, d.schema, d.id].join(" ").toLowerCase()}));
D.audiences.forEach((a, i) => IDX.push({t: "a", i, title: a.name, sub: `audience · ${a.eval || ""} ${a.life || ""}`,
  hay: [a.name, a.desc, a.tags, a.id, a.by].join(" ").toLowerCase()}));
D.journeys.forEach((x, i) => IDX.push({t: "j", i, title: x.name, sub: `journey · ${x.status || ""}${x.version ? " · v" + x.version : ""}`,
  hay: [x.name, x.reads, x.by, x.id].join(" ").toLowerCase()}));
D.sources.forEach((x, i) => IDX.push({t: "r", i, title: x.name, sub: `source · ${x.kind || ""}`,
  hay: [x.name, x.kind, x.from, x.id, ...x.into.map(d => D.datasets[d].table)].join(" ").toLowerCase()}));
D.destinations.forEach((x, i) => IDX.push({t: "x", i, title: x.name || x.dest, sub: `destination · ${x.dest || ""}`,
  hay: [x.name, x.dest, x.id].join(" ").toLowerCase()}));
const TYPE_LABEL = {f: "Field", s: "Schema", d: "Dataset", a: "Audience", j: "Journey", r: "Source", x: "Destination"};
const TYPE_BADGE = {f: "b-event", s: "b-profile", d: "b-out", a: "b-aud", j: "b-jrn", r: "b-src", x: "b-dst"};
function search(q, limit) {
  const words = q.toLowerCase().split(/\s+/).filter(Boolean);
  if (!words.length) return [];
  const out = [];
  for (const e of IDX) {
    if (!shown(e.t, e.i)) continue;
    if (!words.every(w => e.hay.includes(w))) continue;
    const t = e.title.toLowerCase(), s = (e.sub || "").toLowerCase(), q0 = q.toLowerCase().trim();
    let sc = 0;
    if (t === q0 || s === q0) sc += 100; else if (t.startsWith(q0)) sc += 60; else if (t.includes(q0)) sc += 40;
    if (words.every(w => t.includes(w))) sc += 20;
    sc += {s: 6, d: 4, f: 3, a: 2, j: 1, x: 1, r: 0}[e.t];
    out.push([sc, e]);
  }
  out.sort((a, b) => b[0] - a[0] || a[1].title.length - b[1].title.length);
  return out.slice(0, limit || 400).map(x => x[1]);
}
function mountSearch(el, big) {
  el.innerHTML = SEARCH_ICON + `<input type="search" placeholder="${big ? "Search fields, schemas, datasets and audiences – try “slot” or “coupon”" : "Search…  ( / )"}" autocomplete="off" aria-label="Search"><div class="drop"></div>`;
  const inp = el.querySelector("input"), drop = el.querySelector(".drop");
  let res = [], sel = -1;
  const show = () => {
    const q = inp.value.trim();
    res = q ? search(q, 12) : [];
    sel = -1;
    if (!q) { drop.style.display = "none"; return; }
    drop.innerHTML = (res.length ? res.map(e => `<a href="${href[e.t](e.i)}"><span class="badge ${TYPE_BADGE[e.t]}">${TYPE_LABEL[e.t]}</span><span>${h(e.title)}<div class="sub">${h(e.sub)}</div></span></a>`).join("")
      : `<a><span class="sub">Nothing matches “${h(q)}”.</span></a>`) + `<a class="all" href="#/search/${enc(q)}">See all results for “${h(q)}” →</a>`;
    drop.style.display = "block";
  };
  inp.addEventListener("input", show);
  inp.addEventListener("focus", show);
  inp.addEventListener("keydown", ev => {
    const links = [...drop.querySelectorAll("a[href]")];
    if (ev.key === "ArrowDown" || ev.key === "ArrowUp") {
      ev.preventDefault();
      sel = Math.max(0, Math.min(links.length - 1, sel + (ev.key === "ArrowDown" ? 1 : -1)));
      links.forEach((a, k) => a.classList.toggle("sel", k === sel));
    } else if (ev.key === "Enter") {
      const q = inp.value.trim();
      location.hash = sel >= 0 && links[sel] ? links[sel].getAttribute("href") : "#/search/" + enc(q);
      drop.style.display = "none"; inp.blur();
    } else if (ev.key === "Escape") { drop.style.display = "none"; inp.blur(); }
  });
  document.addEventListener("click", ev => { if (!el.contains(ev.target)) drop.style.display = "none"; });
  drop.addEventListener("click", () => { drop.style.display = "none"; inp.value = ""; });
  return inp;
}
const topInput = mountSearch(document.getElementById("topsearch"), false);
document.addEventListener("keydown", ev => {
  if (ev.key === "/" && !/input|textarea/i.test(document.activeElement.tagName)) { ev.preventDefault(); topInput.focus(); }
});

/* ---------- small pieces ---------- */
function copyBtn(text) {
  const id = "c" + Math.random().toString(36).slice(2);
  COPY[id] = text;
  return `<button data-copy="${id}">Copy</button>`;
}
const COPY = {};
document.addEventListener("click", ev => {
  const b = ev.target.closest("button[data-copy]");
  if (!b) return;
  const text = COPY[b.dataset.copy];
  const done = () => { b.textContent = "Copied ✓"; setTimeout(() => b.textContent = "Copy", 1500); };
  const fallback = () => { const ta = document.createElement("textarea"); ta.value = text; document.body.appendChild(ta); ta.select(); document.execCommand("copy"); ta.remove(); done(); };
  if (navigator.clipboard && window.isSecureContext) navigator.clipboard.writeText(text).then(done, fallback); else fallback();
});
// The workbook's colour scale (Coverage %): red, through yellow, to green.
function scaleColour(scale, v) {
  const pts = scale.slice().sort((a, b) => a[0] - b[0]);
  if (v <= pts[0][0]) return "#" + pts[0][1];
  for (let k = 1; k < pts.length; k++) {
    const [v0, c0] = pts[k - 1], [v1, c1] = pts[k];
    if (v <= v1) {
      const f = v1 === v0 ? 1 : (v - v0) / (v1 - v0);
      return "#" + [0, 2, 4].map(i => Math.round(parseInt(c0.substr(i, 2), 16) + f * (parseInt(c1.substr(i, 2), 16) - parseInt(c0.substr(i, 2), 16))).toString(16).padStart(2, "0")).join("");
    }
  }
  return "#" + pts[pts.length - 1][1];
}
const covColour = v => scaleColour(M.palette.coverage, v);
const statusPill = s => { const x = String(s || "").toLowerCase(); return `<span class="badge ${M.palette.live.includes(x) ? "GREEN" : M.palette.trying.includes(x) ? "AMBER" : "GREY"}">${h(s || "–")}</span>`; };
const lifePill = s => s ? `<span class="badge ${/^published$/i.test(s) ? "GREEN" : "GREY"}">${h(s)}</span>` : "–";
const evalPill = s => s ? `<span class="badge b-neutral">${h(s)}</span>` : "–";
function covBar(c) {
  const n = num(c);
  return n === null ? '<span class="sub">–</span>' : `<span style="white-space:nowrap"><span class="covbar"><i style="width:${Math.max(4, Math.min(100, n))}%;background:${covColour(n)}"></i></span>${fmt(n, 0)}%</span>`;
}
function strip(hist) {
  if (!hist || !hist.length) return '<span class="sub">no history read</span>';
  const cells = hist.map((v, k) => {
    const m = M.months[k];
    if (v === null || v === undefined) return `<i title="${h(m)}: no data arrived this month"></i>`;
    if (v === "?") return `<i class="u" title="${h(m)}: data arrived but couldn't be read (see How to read this)">?</i>`;
    const n = num(v);
    return `<i style="background:${covColour(n)}" title="${h(m)}: ${n === 0 ? "data arrived, but this field was empty in what we read" : `filled in ${fmt(n, 0)}% of the records read`}"></i>`;
  }).join("");
  return `<div style="display:inline-block"><div class="strip">${cells}</div><div class="striplab"><span>${h(M.months[0] || "")}</span><span>${h(M.months[M.months.length - 1] || "")}</span></div></div>`;
}
const STRIP_LEGEND = `<div class="legend"><span><i style="background:${covColour(100)}"></i>filled in nearly always</span><span><i style="background:${covColour(50)}"></i>about half the time</span><span><i style="background:${covColour(0)}"></i>data arrived, field empty</span><span><i style="background:var(--m-unk)"></i>couldn't be read (“?”)</span><span><i style="background:var(--m-none)"></i>no data that month</span></div>`;
function ragBadge(a) { return a.rag ? `<span class="badge ${h(a.rag)}">${h(a.rag)} ${a.score !== null && a.score !== undefined ? fmt(a.score, 0) : ""}</span>` : ""; }

/* ---------- tables ---------- */
const TABLES = {};
function table(id, cols, rows, opts) {
  TABLES[id] = {cols, rows, sort: opts?.sort ?? null, dir: opts?.dir ?? 1, limit: opts?.limit ?? 500};
  return `<div class="tablewrap"><table class="t" id="${id}"><thead></thead><tbody></tbody></table></div><div class="hint" id="${id}-more"></div>`;
}
function drawTable(id) {
  const T = TABLES[id], el = document.getElementById(id);
  if (!T || !el) return;
  const rows = T.rows.slice();
  if (T.sort !== null) {
    const c = T.cols[T.sort];
    rows.sort((a, b) => {
      const x = c.v(a), y = c.v(b);
      const nx = x === null || x === undefined || x === "", ny = y === null || y === undefined || y === "";
      if (nx || ny) return nx - ny;
      return (typeof x === "number" && typeof y === "number" ? x - y : String(x).localeCompare(String(y))) * T.dir;
    });
  }
  el.querySelector("thead").innerHTML = "<tr>" + T.cols.map((c, k) => `<th class="${c.num ? "num" : ""}" data-k="${k}" title="${h(c.tip || "Click to sort")}">${h(c.label)} <span class="arr">${T.sort === k ? (T.dir > 0 ? "▲" : "▼") : ""}</span></th>`).join("") + "</tr>";
  el.querySelector("tbody").innerHTML = rows.slice(0, T.limit).map(r => "<tr>" + T.cols.map(c => `<td class="${c.num ? "num" : ""}">${c.html(r)}</td>`).join("") + "</tr>").join("")
    || `<tr><td colspan="${T.cols.length}" class="sub">Nothing to show.</td></tr>`;
  document.getElementById(id + "-more").textContent = rows.length > T.limit ? `Showing the first ${T.limit} of ${fmt(rows.length, 0)} – type in the filter box to narrow it down.` : "";
  el.querySelectorAll("th").forEach(th => th.onclick = () => {
    const k = +th.dataset.k;
    if (T.sort === k) T.dir = -T.dir; else { T.sort = k; T.dir = T.cols[k].num ? -1 : 1; }
    drawTable(id);
  });
}

/* ---------- lineage ---------- */
let LIN = [];
function linked(a, b) {
  if (a.t > b.t) [a, b] = [b, a];
  const k = a.t + b.t;
  if (k === "ds") return D.datasets[a.i].s === b.i;
  if (k === "fs") return D.fields[a.i].on.some(o => o.s === b.i);
  if (k === "af") return D.audiences[a.i].f.includes(b.i);
  if (k === "aa") return D.audiences[a.i].dep.includes(b.i) || D.audiences[b.i].dep.includes(a.i);
  if (k === "as") return readsFor(D.audiences[a.i], b.i).some(fi => D.fields[fi].on.some(o => o.s === b.i));
  if (k === "aj") return D.journeys[b.i].aud.includes(a.i);
  if (k === "ax") return D.destinations[b.i].aud.includes(a.i);
  if (k === "dr") return D.sources[b.i].into.includes(a.i);
  return false;
}
function nodeHtml(n) {
  let label, sub = "", kc = "ds", extra = "";
  if (n.t === "f") { const f = D.fields[n.i]; label = fieldLabel(f); sub = f.path; kc = fieldKind(f); }
  if (n.t === "s") { const s = D.schemas[n.i]; label = s.name; sub = `${kindWord(s.kind)}${s.profile ? " · in Profile" : ""}`; kc = kindOf(s.kind); }
  if (n.t === "d") { const d = D.datasets[n.i]; label = d.name || d.table; sub = dsInProfile(d) ? "in Profile" : "data lake only"; kc = "ds"; if (!dsInProfile(d)) extra = " dim"; }
  if (n.t === "a") { const a = D.audiences[n.i]; label = a.name; sub = [a.eval, a.life].filter(Boolean).join(" · "); kc = "aud"; if (a.rag) sub += `<span class="r ${h(a.rag)}"></span>`; }
  if (n.t === "j") { const x = D.journeys[n.i]; label = x.name; sub = `journey · ${x.status || ""}${x.version ? " · v" + x.version : ""}`; kc = "jrn"; if (!x.live) extra = " dim"; }
  if (n.t === "x") { const x = D.destinations[n.i]; label = x.name || x.dest; sub = x.dest || "destination"; kc = "dst"; if (!x.live) extra = " dim"; }
  if (n.t === "r") { const x = D.sources[n.i]; label = x.name || x.kind; sub = x.kind || "source"; kc = "src"; if (!x.live) extra = " dim"; }
  return `<a class="node k-${kc}${n.focus ? " focus" : ""}${extra}" data-n="${n.t}${n.i}" href="${href[n.t](n.i)}" title="${h(label)}"><span class="l">${h(label)}</span><span class="s">${n.t === "a" ? sub : h(sub)}</span></a>`;
}
const CAP = 10;
function lineage(cols, hint) {
  const id = "lin" + LIN.length;
  LIN.push({id, cols});
  const body = cols.map(c => {
    const items = c.items.slice(0, CAP);
    const more = c.items.length > CAP ? `<a class="more" href="${c.more || "#"}">+ ${fmt(c.items.length - CAP, 0)} more – see the list below</a>` : "";
    return `<div class="col${c.ph ? " ph" : ""}"><h4>${h(c.title)}</h4>${items.length ? items.map(nodeHtml).join("") : `<div class="empty">${c.empty || "None"}</div>`}${more}</div>`;
  }).join("");
  return `<div class="panel"><div class="linwrap"><div class="lin" id="${id}"><svg></svg>${body}</div></div><div class="hint">${hint || "Every box is a link – click one to make it the centre and keep following the data."}</div></div>`;
}
function drawLineage() {
  for (const L of LIN) {
    const el = document.getElementById(L.id);
    if (!el) continue;
    const svg = el.querySelector("svg");
    const box = el.getBoundingClientRect();
    svg.setAttribute("width", el.scrollWidth); svg.setAttribute("height", el.scrollHeight);
    const colEls = [...el.querySelectorAll(":scope > .col")];
    let paths = "";
    for (let c = 0; c + 1 < L.cols.length; c++) {
      const A = L.cols[c].items.slice(0, CAP), B = L.cols[c + 1].items.slice(0, CAP);
      for (const a of A) for (const b of B) {
        if (!linked(a, b)) continue;
        const ea = colEls[c].querySelector(`[data-n="${a.t}${a.i}"]`), eb = colEls[c + 1].querySelector(`[data-n="${b.t}${b.i}"]`);
        if (!ea || !eb) continue;
        const ra = ea.getBoundingClientRect(), rb = eb.getBoundingClientRect();
        const x1 = ra.right - box.left, y1 = ra.top + ra.height / 2 - box.top, x2 = rb.left - box.left, y2 = rb.top + rb.height / 2 - box.top;
        const mx = (x1 + x2) / 2, hot = a.focus || b.focus ? " hot" : "";
        paths += `<path class="e${hot}" data-a="${a.t}${a.i}" data-b="${b.t}${b.i}" d="M${x1},${y1} C${mx},${y1} ${mx},${y2} ${x2},${y2}"/>`;
      }
    }
    svg.innerHTML = paths;
    el.querySelectorAll(".node").forEach(nd => {
      nd.onmouseenter = () => svg.querySelectorAll(`path[data-a="${nd.dataset.n}"],path[data-b="${nd.dataset.n}"]`).forEach(p => p.classList.add("hot"));
      nd.onmouseleave = () => svg.querySelectorAll("path").forEach(p => p.classList.toggle("hot", /focus/.test(
        (el.querySelector(`[data-n="${p.dataset.a}"]`) || {}).className + " " + (el.querySelector(`[data-n="${p.dataset.b}"]`) || {}).className)));
    });
  }
}
window.addEventListener("resize", () => requestAnimationFrame(drawLineage));
const uniq = xs => [...new Set(xs)];
const N = (t, i, focus) => ({t, i, focus: !!focus});
// An audience reads a profile field from a profile schema and an event field from an
// event schema (the workbook says which); a lookup schema can be reached either way.
const readsFor = (a, s) => { const k = kindOf(D.schemas[s].kind); return k === "profile" ? a.fp : k === "event" ? a.fe : a.f; };
const schemaUsed = s => D.schemas[s].f.filter(f => D.fields[f].aud.some(a => readsFor(D.audiences[a], s).includes(f)));
const schemaAud = s => uniq(schemaUsed(s).flatMap(f => D.fields[f].aud.filter(a => readsFor(D.audiences[a], s).includes(f))));
const liveFirst = (xs, list) => xs.slice().sort((p, q) => (list[q].live ? 1 : 0) - (list[p].live ? 1 : 0));
// The two ends of the lineage, from the workbook's Sources / Journeys / Destinations tabs.
function srcCol(dsIdx, more) {
  return {title: "Where it comes from", items: liveFirst(showIdx("r", uniq(dsIdx.flatMap(d => D.datasets[d].src))), D.sources).map(r => N("r", r)), more,
          ph: !HAS_LINEAGE, empty: HAS_LINEAGE ? "No dataflow feeds it (loaded by an API call, or a one-off job)." : NOT_IN};
}
function goesCol(audIdx, title, more) {
  const js = liveFirst(uniq(audIdx.flatMap(a => D.audiences[a].j)), D.journeys).map(j => N("j", j));
  const xs = liveFirst(uniq(audIdx.flatMap(a => D.audiences[a].x)), D.destinations).map(x => N("x", x));
  return {title: title || "Where it goes", items: [...js, ...xs], more, ph: !HAS_LINEAGE,
          empty: HAS_LINEAGE ? "No journey or destination uses it." : NOT_IN};
}
const audSort = (x, y) => (D.audiences[y].score || 0) - (D.audiences[x].score || 0);

/* ---------- pages ---------- */
function pageHome() {
  const S = showRows("s", D.schemas), kinds = {profile: 0, event: 0, other: 0};
  S.forEach(s => kinds[kindOf(s.kind)]++);
  const DS = showRows("d", D.datasets), FS = showRows("f", D.fields);
  const inProf = DS.filter(dsInProfile).length;
  const withHist = FS.filter(f => f.on.some(o => num(o.months) > 0)).length;
  const red = D.audiences.filter(a => a.rag === "RED");
  const linkedAud = D.audiences.filter(a => a.f.length).length;
  const card = (s) => `<a class="panel tile" href="${href.s(D.schemas.indexOf(s))}" style="padding:10px 14px"><b>${h(s.name)}</b><div class="l">${kindWord(s.kind)} · ${plural(s.f.length, "field")} · ${s.profile ? "in Profile" : "not in Profile"}</div></a>`;
  const group = (k, title) => { const xs = S.filter(s => kindOf(s.kind) === k); return xs.length ? `<h3><span class="badge b-${k}">${title}</span> ${xs.length}</h3><div class="grid g4">${xs.map(card).join("")}</div>` : ""; };
  return `<section class="hero"><div class="crumb">${h(M.label)}</div><h1>What data do we have, and what uses it?</h1>
    <p class="lede">Search for anything below, or start from a box. Every name is a link: click through to follow the data from where it comes in, to the ${term("dataset", "dataset")} it lands in, its ${term("schema", "schema")} and ${term("fields", "field")}, the ${term("audiences", "audience")} that use it, and the journeys and destinations they go to.</p>
    <div class="searchwrap searchbig" id="bigsearch"></div></section>
    <div class="grid g4" style="margin-top:22px">
      <a class="panel tile" href="#/fields"><div class="n">${fmt(FS.length, 0)}</div><div class="l">fields · ${fmt(withHist, 0)} with history</div></a>
      <a class="panel tile" href="#/schemas"><div class="n">${fmt(S.length, 0)}</div><div class="l">schemas · ${kinds.profile} profile, ${kinds.event} event, ${kinds.other} other</div></a>
      <a class="panel tile" href="#/datasets"><div class="n">${fmt(DS.length, 0)}</div><div class="l">datasets · ${fmt(inProf, 0)} loaded into Profile</div></a>
      <a class="panel tile" href="#/audiences"><div class="n">${fmt(D.audiences.length, 0)}</div><div class="l">audiences · ${fmt(linkedAud, 0)} traced to their fields</div></a>
    </div>
    <div class="grid g2" style="margin-top:14px">
      <div class="panel card"><h3>How far back can an audience look?</h3><p>Audiences read ${term("Profile", "profile")}, and Profile keeps each event dataset for a set number of days. An audience can't look back further than that.</p>
        ${retentionList()}</div>
      <div class="panel card"><h3>Is a field good enough to build on?</h3><p>Open any field to see, in one place: how often it's filled in, how many months of it we hold, how many days of it Profile has now, and which audiences already use it.</p>
        <p><a href="#/fields">Browse all fields →</a> &nbsp; <a href="#/help">How to read a field page →</a></p>
        <h3 style="margin-top:16px">Heaviest audiences</h3><p>${red.length ? `${plural(red.length, "audience")} score RED: the heaviest for Adobe to work out.` : "None scores RED. The heaviest are:"}</p>
        <ul class="list-plain">${D.audiences.filter(a => num(a.score) !== null).sort((a, b) => b.score - a.score).slice(0, 6).map(a => `<li><a href="${href.a(D.audiences.indexOf(a))}">${h(a.name)}</a> ${ragBadge(a)}</li>`).join("")}</ul>
        <a href="#/audiences${red.length ? "/RED" : ""}">${red.length ? `All ${red.length} RED audiences →` : "All audiences →"}</a></div>
    </div>
    ${HAS_LINEAGE ? homeLineage() : ""}
    ${M.gaps.length ? `<div class="panel" style="margin-top:14px"><h3 style="margin-top:0">Where the picture is incomplete</h3><ul class="list-plain">${M.gaps.map(g => { const i = by.schema[g.tab]; return `<li>${i !== undefined ? `<a href="${href.s(i)}">${h(D.schemas[i].name)}</a>` : h(g.schema)}: ${h(g.status)}</li>`; }).join("")}</ul></div>` : ""}
    <h2>Every schema</h2>${group("profile", "Profile records")}${group("event", "Events")}${group("other", "Lookups and other")}`;
}

function homeLineage() {
  const pub = D.audiences.filter(a => /publish/i.test(a.life || ""));
  const n = re => pub.filter(a => re.test(a.where || "")).length;
  const kinds = {};
  showRows("r", D.sources).forEach(r => { kinds[r.kind] = kinds[r.kind] || {all: 0, live: 0}; kinds[r.kind].all++; kinds[r.kind].live += r.live ? 1 : 0; });
  return `<div class="grid g2" style="margin-top:14px">
    <div class="panel card"><h3>Where audiences go</h3><p>Of ${plural(pub.length, "published audience")}:</p>
      <ul class="list-plain"><li><span class="badge GREEN">In use</span> ${fmt(n(/^In use/), 0)} – a live journey or destination, or another audience built on it</li>
      <li><span class="badge AMBER">Not live</span> <a href="#/audiences/notlive">${fmt(n(/^Not live/), 0)}</a> – only in journeys not running or dataflows switched off</li>
      <li><span class="badge RED">Nothing found using it</span> <a href="#/audiences/unused">${fmt(n(/^Nothing found/), 0)}</a> – candidates to unpublish</li></ul>
      <p><a href="#/journeys">${plural(D.journeys.length, "journey version")} →</a> &nbsp; <a href="#/destinations">${plural(D.destinations.length, "destination dataflow")} →</a></p></div>
    <div class="panel card"><h3>How data comes in</h3><p>Every dataflow that fills a dataset, by kind:</p>
      <ul class="list-plain">${Object.entries(kinds).sort((p, q) => q[1].all - p[1].all).map(([k, v]) => `<li><span class="badge b-src">${h(k)}</span> ${fmt(v.all, 0)}${v.live < v.all ? ` <span class="sub">(${fmt(v.live, 0)} live)</span>` : ""}</li>`).join("")}</ul>
      <p><a href="#/sources">All sources →</a></p></div></div>`;
}
const UUID = "[0-9a-f]{8}-[0-9a-f]{4}-[0-9a-f]{4}-[0-9a-f]{4}-[0-9a-f]{12}";
function ruleHtml(rule) {
  // Audience ids become their names, linked: "is in «Mens_Audience»".
  const html = h(rule), unh = s => s.replace(/&quot;/g, '"').replace(/&#39;/g, "'").replace(/&lt;/g, "<").replace(/&gt;/g, ">").replace(/&amp;/g, "&");
  const link = key => { const i = by.audience[key] ?? audByName[unh(key)]; return i !== undefined ? `<a href="${href.a(i)}">«${h(D.audiences[i].name)}»</a>` : null; };
  // One pass, so the links it writes are not rewritten again. inSegment() names an audience by id or by name.
  const re = new RegExp(String.raw`segmentMembership\.\w+\.(${UUID})\.status in \[&quot;realized&quot;, &quot;existing&quot;\]|inSegment\(&quot;(.+?)&quot;\)|(${UUID})`, "g");
  return html.replace(re, (m, a, b, c) => { const l = link(a || b || c); return l ? (c ? l : "is in " + l) : m; });
}
function retentionList() {
  const rows = M.retention.map(r => {
    const t = r.split(":")[0].trim(), m = /Profile keeps (\d+) days/.exec(r), lake = /data lake: ([^)]+)/.exec(r);
    return {t, i: by.dataset[t], days: m ? +m[1] : null, lake: lake ? lake[1] : ""};
  });
  if (!rows.length) return '<p class="sub">Not in this workbook.</p>';
  const li = r => `<li>${r.i !== undefined ? `<a href="${href.d(r.i)}">${h(r.t)}</a>` : h(r.t)}: <b>${fmt(r.days, 0)} days</b>${r.lake ? ` <span class="sub">(data lake: ${h(r.lake)})</span>` : ""}</li>`;
  const short = rows.filter(r => r.days !== null && r.days < 90), rest = rows.filter(r => !short.includes(r));
  const common = rest.length && rest.every(r => r.days === rest[0].days) ? `${fmt(rest[0].days, 0)} days` : "90 days or more";
  return `${short.length ? `<p><b>Shorter than 90 days:</b></p><ul class="list-plain">${short.map(li).join("")}</ul>` : ""}
    ${rest.length ? `<details><summary>${short.length ? "The other " : ""}${plural(rest.length, "dataset")} keep ${common}</summary><ul class="list-plain">${rest.map(li).join("")}</ul></details>` : ""}`;
}
// The workbook's 'Top values (count)': "value(count) | value(count) | ..."; `limit` keeps the first few.
function topValues(s, limit) {
  if (!s) return "";
  const vals = String(s).split(" | ").map(v => { const m = /^(.*)\((\d+)\)$/.exec(v.trim()); return m ? [m[1], m[2]] : [v.trim(), ""]; });
  const shown = limit ? vals.slice(0, limit) : vals;
  return `<div class="vals${limit ? " mini" : ""}">${shown.map(([v, n]) => `<span title="${h(v)}${n ? ` – ${n} of the records read` : ""}">${h(v === "" ? "(blank)" : v)}${n ? `<b>×${fmt(+n, 0)}</b>` : ""}</span>`).join("")}${limit && vals.length > limit ? `<em>+${vals.length - limit}</em>` : ""}</div>`;
}

function pageField(i) {
  const f = D.fields[i];
  const schemaIdx = uniq(visOn(f).map(o => o.s));
  const ds = uniq(schemaIdx.flatMap(s => D.schemas[s].ds));
  // Audiences that go somewhere come first, so the lines to the last column show.
  const aud = f.aud.slice().sort((x, y) => (D.audiences[y].j.length + D.audiences[y].x.length > 0) - (D.audiences[x].j.length + D.audiences[x].x.length > 0) || audSort(x, y));
  const best = f.on.reduce((m, o) => Math.max(m, num(o.months) || 0), 0);
  const pd = f.on.reduce((m, o) => Math.max(m, num(o.pdays) || 0), 0);
  const cov = f.on.reduce((m, o) => Math.max(m, num(o.cov) ?? -1), -1);
  const ident = f.on.map(o => o.ident).find(Boolean), rel = f.on.map(o => o.rel).find(Boolean);
  const say = [];
  say.push(`It holds ${h(typeWords(f.type))}${f.on.length > 1 ? ` and appears in ${plural(f.on.length, "schema")}` : ` in the <a href="${href.s(f.on[0].s)}">${h(D.schemas[f.on[0].s].name)}</a> schema`}.`);
  if (cov >= 0) say.push(`In a recent sample it was filled in for <b>${fmt(cov, 0)}%</b> of records${f.on.length > 1 ? " (best schema)" : ""}.`);
  if (best > 0) say.push(`We hold about <b>${fmt(best)} months</b> of it.`);
  if (pd > 0) say.push(`Profile has <b>${fmt(pd, 0)} days</b> of it right now, so audiences can look back that far.`);
  say.push(aud.length ? `<b>${plural(aud.length, "audience")}</b> use it.` : "No audience uses it yet.");
  const rows = visOn(f).map(o => ({o, s: D.schemas[o.s]}));
  return `<div class="crumb"><a href="#/fields">Fields</a> › ${h(f.fpath || f.path)}</div>
    <h1>${h(fieldLabel(f))}</h1>
    <div class="copyrow" style="max-width:760px"><pre class="code">${h(f.path)}</pre>${copyBtn(f.path)}</div>
    ${f.desc ? `<p class="say">${h(f.desc)}</p>` : ""}
    <p class="say">${say.join(" ")}</p>
    ${ident || rel ? `<p>${ident ? `<span class="badge b-profile">Identity: ${h(ident)}</span>` : ""}${rel ? `<span class="badge b-other">Links to ${h(rel)}</span>` : ""}</p>` : ""}
    <h2>Lineage</h2>
    ${lineage([
      srcCol(ds),
      {title: "Stored in (datasets)", items: ds.map(d => N("d", d)), more: "#f-where"},
      {title: "Schemas", items: schemaIdx.map(s => N("s", s)), more: "#f-where"},
      {title: "This field", items: [N("f", i, true)]},
      {title: "Audiences that read it", items: aud.map(a => N("a", a)), more: "#f-aud", empty: "No audience reads this field."},
      goesCol(aud, "Where those go", "#f-aud")], "Greyed boxes are not live: a dataset only in the data lake (audiences can't read it), a journey not running, a dataflow switched off. Every box is a link.")}
    <h2 id="f-where">Where it lives, and how good the data is</h2>
    ${rows.map(({o, s}) => `<div class="panel" style="margin-bottom:12px">
      <h3 style="margin-top:0"><a href="${href.s(o.s)}">${h(s.name)}</a> <span class="badge b-${kindOf(s.kind)}">${kindWord(s.kind)}</span>${profileWord(s.profile)}${o.req ? '<span class="badge b-out">Required</span>' : ""}</h3>
      <div class="grid g3">
        <div><div class="sub">Filled in (recent sample)</div><div class="big">${num(o.cov) === null ? "–" : fmt(o.cov, 0) + "%"}</div>${o.top ? `<div class="sub" style="margin-top:6px">Most common values (personal data masked):</div>${topValues(o.top)}` : ""}</div>
        <div><div class="sub">Have we got it? (months of data)</div><div class="big">${num(o.months) === null ? "?" : fmt(o.months)} <small>months</small></div><div class="sub">${h(o.hdet || "")}</div></div>
        <div><div class="sub">In Profile now</div><div class="big">${num(o.pdays) === null ? (o.pdays ? h(o.pdays) : "–") : fmt(o.pdays, 0)} <small>days</small></div><div class="sub">${h(o.pdet || "")}</div></div>
      </div>
      <h3>Month by month</h3>${strip(o.hist)}${STRIP_LEGEND}
      ${o.sql ? `<details><summary>SQL to check it yourself (paste into Query Service)</summary><div class="copyrow"><pre class="code">${h(o.sql)}</pre>${copyBtn(o.sql)}</div></details>` : ""}
    </div>`).join("")}
    <h2 id="f-aud">Audiences that read this field</h2>
    ${aud.length ? table("t-faud", audCols(), aud.map(a => D.audiences[a]), {sort: 4, dir: -1}) : '<p class="sub">None.</p>'}`;
}

function audCols() {
  return [
    {label: "Audience", v: a => a.name, html: a => `<a href="${href.a(by.audience[a.id])}">${h(a.name)}</a>${a.desc ? `<div class="sub">${h(a.desc)}</div>` : ""}`},
    {label: "Evaluation", v: a => a.eval, html: a => evalPill(a.eval)},
    {label: "Status", v: a => a.life, html: a => lifePill(a.life)},
    {label: "Where it goes", v: a => a.where, html: a => a.where ? `<span class="badge ${WHERE_RAG(a.where)}">${h(a.where.split(":")[0].replace(/ \(.*/, ""))}</span><div class="sub">${h((a.where.split(":")[1] || "").trim())}</div>` : '<span class="sub">–</span>'},
    {label: "Complexity", num: true, v: a => num(a.score), html: a => ragBadge(a), tip: "0–100; RED is heaviest for Adobe to evaluate"},
    {label: "Fields read", num: true, v: a => a.f.length, html: a => fmt(a.f.length, 0)},
    {label: "Built on", num: true, v: a => a.dep.length, html: a => fmt(a.dep.length, 0), tip: "Other audiences this one includes"},
  ];
}
function jrnCols() {
  return [
    {label: "Journey", v: j => (j.name || "").toLowerCase(), html: j => `<a href="${href.j(D.journeys.indexOf(j))}">${h(j.name)}</a>${j.version ? `<div class="sub">version ${h(j.version)}</div>` : ""}`},
    {label: "Status", v: j => j.status, html: j => statusPill(j.status)},
    {label: "Live", v: j => j.live ? 1 : 0, html: j => liveWord(j.live)},
    {label: "How", v: j => j.how, html: j => h(j.how || "–")},
    {label: "Last changed", v: j => j.changed, html: j => h(j.changed || "–") + (j.by ? `<div class="sub">${h(j.by)}</div>` : "")},
  ];
}
function dstCols() {
  return [
    {label: "Dataflow", v: x => (x.name || "").toLowerCase(), html: x => `<a href="${href.x(D.destinations.indexOf(x))}">${h(x.name || x.dest)}</a>`},
    {label: "Destination", v: x => x.dest, html: x => h(x.dest || "–")},
    {label: "Live", v: x => x.live ? 1 : 0, html: x => liveWord(x.live)},
    {label: "Audiences", num: true, v: x => x.aud.length, html: x => fmt(x.aud.length, 0)},
  ];
}
function srcCols() {
  return [
    {label: "Dataflow", v: r => (r.name || "").toLowerCase(), html: r => `<a href="${href.r(D.sources.indexOf(r))}">${h(r.name || r.kind)}</a>${r.from ? `<div class="sub">${h(r.from)}</div>` : ""}`},
    {label: "How it comes in", v: r => r.kind, html: r => `<span class="badge b-src">${h(r.kind || "–")}</span>`},
    {label: "Live", v: r => r.live ? 1 : 0, html: r => liveWord(r.live)},
    {label: "Last run (UTC)", v: r => r.last, html: r => h(r.last || "–") + (r.state ? `<div class="sub">${h(r.state)}</div>` : "")},
    {label: "Into", v: r => r.into.length, html: r => r.into.map(d => `<a href="${href.d(d)}">${h(D.datasets[d].table || D.datasets[d].name)}</a>`).join(", ") || "–"},
  ];
}
function fieldCols(s) {
  const pick = f => s === undefined ? f.on.reduce((m, o) => (num(o.months) || 0) >= (num(m.months) || 0) ? o : m, f.on[0]) : f.on.find(o => o.s === s);
  return [
    {label: "Field", v: f => fieldLabel(f).toLowerCase(), html: f => `<a href="${href.f(by.field[f.path])}">${h(fieldLabel(f))}</a><div class="sub">${h(f.path)}</div>`},
    {label: "Holds", v: f => typeWords(f.type), html: f => h(typeWords(f.type))},
    ...(s === undefined ? [{label: "Schemas", num: true, v: f => f.on.length, html: f => fmt(f.on.length, 0)}] : []),
    {label: "Filled in", num: true, v: f => num(pick(f).cov), html: f => covBar(pick(f).cov), tip: "Share of a recent sample of records with this field filled in"},
    {label: "Most common", v: f => pick(f).top || "", html: f => topValues(pick(f).top, 3) || '<span class="sub">–</span>', tip: "Top values in the recent sample, with how many records had each (personal data masked)"},
    {label: "Months held", num: true, v: f => num(pick(f).months), html: f => num(pick(f).months) === null ? '<span class="sub">?</span>' : fmt(pick(f).months), tip: "Have we got it? Months of unbroken data"},
    {label: "Profile days", num: true, v: f => num(pick(f).pdays), html: f => num(pick(f).pdays) === null ? '<span class="sub">–</span>' : fmt(pick(f).pdays, 0), tip: "Days of it in Profile now"},
    {label: "Audiences", num: true, v: f => f.aud.length, html: f => f.aud.length ? `<b>${fmt(f.aud.length, 0)}</b>` : '<span class="sub">0</span>'},
  ];
}
function dsCols() {
  return [
    {label: "Dataset", v: d => (d.name || d.table || "").toLowerCase(), html: d => `<a href="${href.d(by.dataset[d.table || d.id])}">${h(d.name || d.table)}</a><div class="sub">${h(d.table || "")}</div>`},
    {label: "Schema", v: d => d.schema, html: d => d.s >= 0 ? `<a href="${href.s(d.s)}">${h(D.schemas[d.s].name)}</a>` : `<span class="sub">${h(d.schema || "–")}</span>`},
    {label: "In Profile", v: d => dsInProfile(d) ? 1 : 0, html: d => dsInProfile(d) ? '<span class="badge b-in">Yes</span>' : '<span class="sub">No</span>'},
    {label: "Rows (lake)", num: true, v: d => num(d.rows), html: d => fmt(d.rows, 0)},
    {label: "GB (lake)", num: true, v: d => num(d.gb), html: d => fmt(d.gb, 2)},
    {label: "Lake keeps (days)", num: true, v: d => num(d.ttl), html: d => num(d.ttl) === null ? '<span class="sub">no expiry</span>' : fmt(d.ttl, 0)},
    {label: "Profile keeps (days)", num: true, v: d => num(d.pttl), html: d => num(d.pttl) === null ? '<span class="sub">–</span>' : fmt(d.pttl, 0)},
  ];
}

function pageSchema(i) {
  const s = D.schemas[i];
  const used = schemaUsed(i).sort((x, y) => D.fields[y].aud.length - D.fields[x].aud.length);
  const aud = schemaAud(i).sort(audSort);
  const kind = kindOf(s.kind);
  const what = {profile: "one record per customer (or account) – the facts about them", event: "things that happened, each with a date and time", other: "reference data, such as products or stores, that other data links to"}[kind];
  return `<div class="crumb"><a href="#/schemas">Schemas</a> › ${h(s.title)}</div>
    <h1>${h(s.name)}</h1>
    <p><span class="badge b-${kind}">${kindWord(s.kind)}</span>${profileWord(s.profile)}</p>
    ${s.desc && ![s.tab, s.name, s.title].includes(s.desc) ? `<p class="say">${h(s.desc)}</p>` : ""}
    <p class="say">This ${term("schema", "schema")} describes ${what}. It has <b>${plural(s.f.length, "field")}</b> and <b>${plural(s.ds.length, "dataset")}</b>${s.profile ? ", and its data is in Profile, so audiences can use it" : ". None of its data is in Profile, so audiences can't use it directly (it's for SQL and reporting)"}.</p>
    <div class="facts">
      <div>Data class<b>${h(s.cls || "–")}</b></div><div>Data lake<b>${fmt(s.rows, 0)} rows · ${fmt(s.gb, 2)} GB</b></div>
      <div>Lake keeps<b>${num(s.ttl) === null ? "no expiry set" : fmt(s.ttl, 0) + " days"}</b></div><div>Profile keeps<b>${num(s.pttl) === null ? (s.profile ? "no expiry set" : "–") : fmt(s.pttl, 0) + " days"}</b></div>
      <div>Last changed<b>${h(s.mod || "–")}</b></div><div>SQL table names<b class="mono">${h(s.tables || "–")}</b></div>
    </div>
    <h2>Lineage</h2>
    ${lineage([srcCol(s.ds, "#s-ds"),
      {title: "Stored in (datasets)", items: s.ds.map(d => N("d", d)), more: "#s-ds"},
      {title: "This schema", items: [N("s", i, true)]},
      {title: "Fields audiences read", items: used.map(f => N("f", f)), more: "#s-fields", empty: "No audience reads any of its fields."},
      {title: "Audiences", items: aud.map(a => N("a", a)), more: "#s-aud", empty: "None."}])}
    <h2 id="s-fields">All ${plural(s.f.length, "field")}</h2>
    <div class="tools"><input type="search" placeholder="Filter these fields…" data-filter="t-sf"><span class="count" id="t-sf-count"></span></div>
    ${table("t-sf", fieldCols(i), s.f.map(f => D.fields[f]))}
    <h2 id="s-ds">Datasets</h2>${table("t-sds", dsCols(), s.ds.map(d => D.datasets[d]), {sort: 4, dir: -1})}
    <h2 id="s-aud">Audiences that read its fields</h2>${aud.length ? table("t-saud", audCols(), aud.map(a => D.audiences[a]), {sort: 4, dir: -1}) : '<p class="sub">None.</p>'}`;
}

function pageDataset(i) {
  const d = D.datasets[i], s = d.s >= 0 ? D.schemas[d.s] : null, inP = dsInProfile(d);
  const used = s && inP ? schemaUsed(d.s).sort((x, y) => D.fields[y].aud.length - D.fields[x].aud.length) : [];
  const aud = s && inP ? schemaAud(d.s).sort(audSort) : [];
  return `<div class="crumb"><a href="#/datasets">Datasets</a> › ${h(d.table || d.name)}</div>
    <h1>${h(d.name || d.table)}</h1><p>${inP ? '<span class="badge b-in">In Profile</span>' : '<span class="badge b-out">Data lake only</span>'}</p>
    ${d.desc ? `<p class="say">${h(d.desc)}</p>` : ""}
    <p class="say">${inP ? `This ${term("dataset", "dataset")} is loaded into ${term("Profile", "profile")}, so audiences can use its data${num(d.pttl) !== null ? `; Profile keeps it for <b>${fmt(d.pttl, 0)} days</b>` : ""}.` : `This ${term("dataset", "dataset")} is not loaded into Profile: audiences can't use it, but SQL in Query Service can.`}
      The ${term("data lake", "lake")} holds <b>${fmt(d.rows, 0)} rows</b> (${fmt(d.gb, 2)} GB)${num(d.ttl) !== null ? ` and keeps them for ${fmt(d.ttl, 0)} days` : ", with no expiry set"}.</p>
    ${d.fed ? `<p class="say">Fed by: ${h(d.fed)}.</p>` : ""}
    <div class="facts"><div>SQL table name<b class="mono">${h(d.table || "–")}</b></div><div>Schema<b>${s ? `<a href="${href.s(d.s)}">${h(s.name)}</a>` : h(d.schema || "–") + (d.s < 0 ? ' <span class="sub">(not in this dictionary)</span>' : "")}</b></div><div>Dataset ID<b class="mono">${h(d.id || "–")}</b></div><div>Type<b>${h(d.kind || "–")}</b></div></div>
    ${d.table ? `<div class="copyrow" style="max-width:760px;margin-top:8px"><pre class="code">SELECT * FROM ${h(d.table)} LIMIT 10</pre>${copyBtn(`/* Data Dictionary: look at ${d.table} */ SELECT * FROM ${d.table} LIMIT 10`)}</div>` : ""}
    <h2>Lineage</h2>
    ${lineage([srcCol([i], "#d-src"),
      {title: "This dataset", items: [N("d", i, true)]},
      {title: "Schema", items: s ? [N("s", d.s)] : [], empty: "Its schema is not in this dictionary."},
      {title: "Fields audiences read", items: used.map(f => N("f", f)), more: s ? href.s(d.s) : "", empty: inP ? "No audience reads its fields." : "Not in Profile, so no audience can read it."},
      {title: "Audiences", items: aud.map(a => N("a", a)), empty: "None."}],
      "Audiences are traced through the fields they read from this kind of data (profile or event). The same field can also come from another dataset of the schema.")}
    <h2 id="d-src">Where it comes from</h2>
    ${d.fed ? `<p class="say">${h(d.fed)}</p>` : `<p class="sub">${HAS_LINEAGE ? "No dataflow feeds it: it is loaded by an API call or a one-off job, or not at all." : NOT_IN}</p>`}
    ${d.src.length ? table("t-dsrc", srcCols(), d.src.map(r => D.sources[r]), {sort: 3, dir: -1}) : ""}
    ${aud.length ? `<h2>Audiences that read its fields</h2>${table("t-daud", audCols(), aud.map(a => D.audiences[a]), {sort: 4, dir: -1})}` : ""}`;
}

function pageAudience(i) {
  const a = D.audiences[i];
  // A field found in almost every schema (eventType, _id) says nothing about where the data is: use the specific ones.
  const specific = a.f.filter(f => D.fields[f].on.length <= 3);
  const schemas = uniq((specific.length ? specific : a.f).flatMap(f => visOn(D.fields[f]).map(o => o.s)
    .filter(s => readsFor(a, s).includes(f))));
  const reads = [...a.f.map(f => N("f", f)), ...a.dep.map(d => N("a", d))];
  const used = a.used.slice().sort(audSort);
  const noRule = !a.fmt && !a.rule;
  let pretty = a.raw;
  if (a.fmt === "json") { try { pretty = JSON.stringify(JSON.parse(a.raw), null, 1); } catch (e) { } }
  return `<div class="crumb"><a href="#/audiences">Audiences</a> › ${h(a.id)}</div>
    <h1>${h(a.name)}</h1>
    <p>${whereBadge(a)}${ragBadge(a)}${a.eval ? evalPill(a.eval) : ""}${a.life ? lifePill(a.life) : ""}${(a.tags || "").split(/\s*[,;|]\s*/).filter(Boolean).map(t => `<span class="badge b-aud">${h(t)}</span>`).join("")}</p>
    ${a.desc ? `<p class="say">${h(a.desc)}</p>` : ""}
    <p class="say">${noRule ? `This audience has no rule inside Adobe: it was ${a.origin ? `made by <b>${h(a.origin)}</b>` : "uploaded or made elsewhere"}, so there are no fields to trace.`
      : `It reads <b>${plural(a.f.length, "field")}</b>${a.dep.length ? ` and builds on <b>${plural(a.dep.length, "other audience")}</b>` : ""}. ${used.length ? `<b>${plural(used.length, "audience")}</b> build on it.` : "No other audience builds on it."}`}</p>
    <div class="facts"><div>Created by<b>${h(a.by || "–")}</b></div><div>Last changed by<b>${h(a.mod || "–")}</b></div><div>Origin<b>${h(a.origin || "–")}</b></div><div>Looks back<b>${num(a.lookback) ? plural(num(a.lookback), "day") : h(a.lookback || "–")}</b></div><div>Largest audience it builds on<b>${h(a.base || "–")}${num(a.basen) ? ` (${fmt(a.basen, 0)})` : ""}</b></div><div>Audience ID<b class="mono">${h(a.id)}</b></div></div>
    <h2>Lineage</h2>
    ${lineage([
      {title: "Schemas", items: schemas.map(s => N("s", s)), more: "#a-reads"},
      {title: "What it reads", items: reads, more: "#a-reads", empty: noRule ? "No rule – nothing to trace." : "No field we could match."},
      {title: "This audience", items: [N("a", i, true)]},
      {title: "Audiences built on it", items: used.map(x => N("a", x)), more: "#a-used", empty: "None."},
      goesCol([i], "Where it goes", "#a-goes")])}
    <h2 id="a-goes">Where it goes ${whereBadge(a)}</h2>
    ${HAS_LINEAGE ? (a.j.length || a.x.length ? `${a.j.length ? table("t-aj", jrnCols(), a.j.map(j => D.journeys[j]), {sort: 2, dir: -1}) : ""}
      ${a.x.length ? `<h3>Destinations</h3>${table("t-ax", dstCols(), a.x.map(x => D.destinations[x]))}` : ""}`
      : `<p class="say">No journey or destination uses it${a.used.length ? `, but ${plural(a.used.length, "audience")} ${a.used.length === 1 ? "is" : "are"} built on it` : ""}.</p>`)
      : `<p class="sub">${NOT_IN}</p>`}
    ${/^Nothing found/.test(a.where || "") ? '<p class="hint">"Nothing found using it" is not proof: calls from other systems, Data Distiller queries and other products are not seen. Ask the person who made it before unpublishing.</p>' : ""}
    ${a.why ? `<h2>Complexity ${ragBadge(a)}</h2><p class="say">${h(a.why)}</p><p class="hint">How heavy the rule is for Adobe to work out, 0–100. The formula is on the workbook's How to Use tab.</p>` : ""}
    ${a.rule || a.raw ? `<h2>The rule</h2>${a.rule ? `<div class="copyrow"><pre class="code">${ruleHtml(a.rule)}</pre></div>${/\[…\]/.test(a.rule) ? '<p class="hint">[…] marks parts too complex to show as text; the full rule is below.</p>' : ""}${a.cut ? '<p class="hint">Cut short in the workbook (too long for a cell); the full rule is below.</p>' : ""}` : ""}
      ${a.raw ? `<details><summary>Full rule, as Adobe stores it (technical)</summary><div class="copyrow"><pre class="code">${h(pretty)}</pre>${copyBtn(a.raw)}</div>${a.rawcut ? '<p class="hint">Cut short here; the whole rule is in the PQL file next to the workbook.</p>' : ""}</details>` : ""}` : ""}
    <h2 id="a-reads">Fields it reads</h2>
    ${a.f.length ? table("t-af", fieldCols(), a.f.map(f => D.fields[f])) : `<p class="sub">${a.dep.length ? "None directly – it is built from the audiences below." : noRule ? "None – it has no rule." : "None that is in this dictionary."}</p>`}
    ${a.miss.length ? `<p class="hint">Also reads ${a.miss.map(p => `<code>${h(p)}</code>${/^_experience\.customerJourneyManagement\./.test(p) ? " (Journey Optimizer's own email and push tracking)" : /^pushNotificationDetails/.test(p) ? " (Journey Optimizer's push-notification sign-ups)" : ""}`).join(", ")} – not in this dictionary: Adobe system data, or a schema that was left out.</p>` : ""}
    ${a.dep.length ? `<h2>Audiences it builds on</h2>${table("t-adep", audCols(), a.dep.map(d => D.audiences[d]), {sort: 4, dir: -1})}` : ""}
    <h2 id="a-used">Audiences built on it</h2>${used.length ? table("t-aused", audCols(), used.map(x => D.audiences[x]), {sort: 4, dir: -1}) : '<p class="sub">None.</p>'}`;
}

let FILTERS = [];
function listPage(title, lede, id, cols, rowsAll, chips, opts) {
  FILTERS.push({id, rowsAll, chips: chips || [], text: "", on: (opts && opts.on) || new Set()});
  return `<h1>${h(title)}</h1><p class="lede">${lede}</p>
    <div class="tools"><input type="search" placeholder="Filter…" data-filter="${id}" value="">${(chips || []).map((c, k) => `<span class="chip${opts && opts.on && opts.on.has(k) ? " on" : ""}" data-chip="${id}:${k}">${h(c.label)}</span>`).join("")}<span class="count" id="${id}-count"></span></div>
    ${table(id, cols, rowsAll, opts)}`;
}
function applyFilter(id) {
  const F = FILTERS.find(x => x.id === id), T = TABLES[id];
  if (!T) return;
  const words = (F ? F.text : "").toLowerCase().split(/\s+/).filter(Boolean);
  const rows = (F ? F.rowsAll : T.all || T.rows).filter(r => {
    if (F) for (const k of F.on) if (!F.chips[k].test(r)) return false;
    if (!words.length) return true;
    const hay = (r.__hay || (r.__hay = JSON.stringify([r.name, r.path, r.fpath, r.desc, r.tab, r.title, r.table, r.schema, r.tags, r.id]).toLowerCase()));
    return words.every(w => hay.includes(w));
  });
  T.rows = rows;
  const c = document.getElementById(id + "-count");
  if (c) c.textContent = `${fmt(rows.length, 0)} shown`;
  drawTable(id);
}
document.addEventListener("input", ev => {
  const id = ev.target.dataset && ev.target.dataset.filter;
  if (!id) return;
  let F = FILTERS.find(x => x.id === id);
  if (!F) { F = {id, rowsAll: TABLES[id].rows.slice(), chips: [], text: "", on: new Set()}; FILTERS.push(F); }
  F.text = ev.target.value;
  applyFilter(id);
});
document.addEventListener("click", ev => {
  const c = ev.target.closest("[data-chip]");
  if (!c) return;
  const [id, k] = c.dataset.chip.split(":"), F = FILTERS.find(x => x.id === id);
  if (F.on.has(+k)) F.on.delete(+k); else F.on.add(+k);
  c.classList.toggle("on");
  applyFilter(id);
});

function pageFields() {
  const FS = showRows("f", D.fields);
  return listPage("Fields", `Every field across every schema, ${fmt(FS.length, 0)} in all. Click a column heading to sort; click a field for its full story.${aaNote("fields")}`,
    "t-fields", fieldCols(), FS,
    [{label: "Used by an audience", test: f => f.aud.length > 0}, {label: "In Profile", test: f => f.on.some(o => D.schemas[o.s].profile)},
     {label: "12+ months held", test: f => f.on.some(o => num(o.months) >= 12)}, {label: "Profile records", test: f => f.on.some(o => kindOf(D.schemas[o.s].kind) === "profile")},
     {label: "Events", test: f => f.on.some(o => kindOf(D.schemas[o.s].kind) === "event")}], {sort: 7, dir: -1});
}
function pageSchemas() {
  const cols = [
    {label: "Schema", v: s => s.name.toLowerCase(), html: s => `<a href="${href.s(by.schema[s.tab])}">${h(s.name)}</a>${s.desc && ![s.tab, s.name, s.title].includes(s.desc) ? `<div class="sub">${h(s.desc)}</div>` : ""}`},
    {label: "Kind", v: s => kindWord(s.kind), html: s => `<span class="badge b-${kindOf(s.kind)}">${kindWord(s.kind)}</span>`},
    {label: "In Profile", v: s => s.profile ? 1 : 0, html: s => s.profile ? '<span class="badge b-in">Yes</span>' : '<span class="sub">No</span>'},
    {label: "Fields", num: true, v: s => s.f.length, html: s => fmt(s.f.length, 0)},
    {label: "Datasets", num: true, v: s => s.ds.length, html: s => fmt(s.ds.length, 0)},
    {label: "Rows (lake)", num: true, v: s => num(s.rows), html: s => fmt(s.rows, 0)},
    {label: "Audiences", num: true, v: s => schemaAud(by.schema[s.tab]).length, html: s => fmt(schemaAud(by.schema[s.tab]).length, 0), tip: "Audiences that read at least one of its fields"},
  ];
  return listPage("Schemas", `A ${term("schema", "schema")} is the blueprint for one kind of data. Purple are customer records, blue are events, green are lookups.${aaNote("schemas")}`,
    "t-schemas", cols, showRows("s", D.schemas),
    [{label: "In Profile", test: s => s.profile}, {label: "Profile records", test: s => kindOf(s.kind) === "profile"}, {label: "Events", test: s => kindOf(s.kind) === "event"}, {label: "Lookups / other", test: s => kindOf(s.kind) === "other"}],
    {sort: 6, dir: -1});
}
function pageDatasets() {
  return listPage("Datasets", `A ${term("dataset", "dataset")} is where data actually lands. All ${fmt(showRows("d", D.datasets).length, 0)} are listed, including those whose schema was left out of this dictionary.${aaNote("datasets")}`,
    "t-datasets", dsCols(), showRows("d", D.datasets),
    [{label: "In Profile", test: dsInProfile}, {label: "Schema in this dictionary", test: d => d.s >= 0}, {label: "Has rows", test: d => num(d.rows) > 0}],
    {sort: 3, dir: -1, on: new Set([1])});
}
function pageAudiences(rag) {
  const chips = [{label: "RED", test: a => a.rag === "RED"}, {label: "AMBER", test: a => a.rag === "AMBER"}, {label: "GREEN", test: a => a.rag === "GREEN"},
    {label: "Published", test: a => /publish/i.test(a.life || "")}, {label: "Streaming", test: a => /stream/i.test(a.eval || "")}, {label: "Batch", test: a => /batch/i.test(a.eval || "")},
    {label: "Has a rule", test: a => a.f.length || a.dep.length},
    {label: "In use", test: a => /^In use/.test(a.where || "")}, {label: "Not live", test: a => /^Not live/.test(a.where || "")},
    {label: "Nothing found using it", test: a => /^Nothing found/.test(a.where || "")}];
  const want = {RED: "RED", unused: "Nothing found using it", notlive: "Not live"}[rag];
  const on = new Set(want ? [chips.findIndex(c => c.label === want), ...(rag === "unused" ? [3] : [])] : []);
  return listPage("Audiences", `Every ${term("audience", "audience")} in ${h(M.label)}. Complexity is how heavy the rule is for Adobe to work out (RED is heaviest).`,
    "t-audiences", audCols(), D.audiences, chips, {sort: 4, dir: -1, on});
}
function pageJourneys() {
  return listPage("Journeys", `Every Adobe Journey Optimizer journey, one row per version, and the audiences it reads. ${HAS_LINEAGE ? "" : NOT_IN}`,
    "t-journeys", [...jrnCols(), {label: "Audiences", num: true, v: j => j.aud.length, html: j => fmt(j.aud.length, 0)}], D.journeys,
    [{label: "Live", test: j => j.live}, {label: "Not live", test: j => !j.live}, {label: "Reads an audience", test: j => j.aud.length > 0}],
    {sort: 2, dir: -1}) + `<p class="hint"><a href="#/destinations">Destinations →</a> the other place audiences go.</p>`;
}
function pageDestinations() {
  return listPage("Destinations", `Every dataflow sending audiences out of AEP. ${HAS_LINEAGE ? "" : NOT_IN}`,
    "t-dests", dstCols(), D.destinations, [{label: "Live", test: x => x.live}], {sort: 2, dir: -1});
}
function pageSources() {
  const kinds = {};
  showRows("r", D.sources).forEach(r => kinds[r.kind] = (kinds[r.kind] || 0) + 1);
  const chips = Object.keys(kinds).sort((p, q) => kinds[q] - kinds[p]).map(k => ({label: `${k} (${kinds[k]})`, test: r => r.kind === k}));
  return listPage("Sources", `Every dataflow that brings data into a dataset. ${HAS_LINEAGE ? "Pick a kind to narrow it down." : NOT_IN}${aaNote("feeds")}`,
    "t-sources", srcCols(), showRows("r", D.sources), [{label: "Live", test: r => r.live}, {label: "Into Profile", test: r => r.profile}, ...chips], {sort: 3, dir: -1});
}
function pageJourney(i) {
  const x = D.journeys[i];
  return `<div class="crumb"><a href="#/journeys">Journeys</a> › ${h(x.id || "")}</div>
    <h1>${h(x.name)}</h1><p><span class="badge b-jrn">Journey</span>${liveWord(x.live)}${statusPill(x.status)}${x.version ? `<span class="badge b-neutral">version ${h(x.version)}</span>` : ""}</p>
    <p class="say">${x.live ? "It is running now" : "It is not running"}${x.aud.length ? ` and reads <b>${plural(x.aud.length, "audience")}</b> (${h(x.how || "")})` : " and reads no audience: it starts from an event or an API call"}.</p>
    <div class="facts"><div>Last changed<b>${h(x.changed || "–")}</b></div><div>Changed by<b>${h(x.by || "–")}</b></div><div>Journey ID<b class="mono">${h(x.id || "–")}</b></div></div>
    <h2>Lineage</h2>
    ${lineage([{title: "Audiences it reads", items: x.aud.map(a => N("a", a)), more: "#j-aud", empty: "None."},
               {title: "This journey", items: [N("j", i, true)]}])}
    <h2 id="j-aud">Audiences it reads</h2>${x.aud.length ? table("t-jaud", audCols(), x.aud.map(a => D.audiences[a]), {sort: 4, dir: -1}) : '<p class="sub">None.</p>'}`;
}
function pageDestination(i) {
  const x = D.destinations[i];
  return `<div class="crumb"><a href="#/destinations">Destinations</a> › ${h(x.id || "")}</div>
    <h1>${h(x.name || x.dest)}</h1><p><span class="badge b-dst">${h(x.dest || "Destination")}</span>${liveWord(x.live)}</p>
    <p class="say">This dataflow sends <b>${plural(x.aud.length, "audience")}</b> to ${h(x.dest || "a destination")}${x.live ? "" : ", but it is switched off"}.</p>
    <h2>Lineage</h2>
    ${lineage([{title: "Audiences it sends", items: x.aud.map(a => N("a", a)), more: "#x-aud", empty: "None."},
               {title: "This destination", items: [N("x", i, true)]}])}
    <h2 id="x-aud">Audiences it sends</h2>${x.aud.length ? table("t-xaud", audCols(), x.aud.map(a => D.audiences[a]), {sort: 4, dir: -1}) : '<p class="sub">None.</p>'}`;
}
function pageSource(i) {
  const x = D.sources[i];
  const schemas = uniq(x.into.map(d => D.datasets[d].s).filter(s => s >= 0));
  return `<div class="crumb"><a href="#/sources">Sources</a> › ${h(x.id || "")}</div>
    <h1>${h(x.name || x.kind)}</h1><p><span class="badge b-src">${h(x.kind || "Source")}</span>${liveWord(x.live)}${x.profile ? '<span class="badge b-in">Into Profile</span>' : ""}</p>
    <p class="say">${h(x.kind || "This dataflow")} into <b>${x.into.map(d => h(D.datasets[d].name || D.datasets[d].table)).join(", ") || "a dataset not in this dictionary"}</b>${x.from ? ` – ${h(x.from)}` : ""}. ${x.last ? `It last ran at ${h(x.last)} UTC${x.state ? ` (${h(x.state)})` : ""}.` : "It has not run."}</p>
    <div class="facts"><div>Created by<b>${h(x.by || "–")}</b></div><div>Dataflow ID<b class="mono">${h(x.id || "–")}</b></div></div>
    <h2>Lineage</h2>
    ${lineage([{title: "This source", items: [N("r", i, true)]},
               {title: "Into (datasets)", items: x.into.map(d => N("d", d)), empty: "A dataset not in this dictionary."},
               {title: "Schemas", items: schemas.map(s => N("s", s)), empty: "Its schema is not in this dictionary."}])}`;
}
function pageSearch(q) {
  const res = search(q, 400);
  const groups = ["f", "s", "d", "a", "j", "x", "r"].map(t => [t, res.filter(e => e.t === t)]).filter(g => g[1].length);
  return `<h1>Results for “${h(q)}”</h1><p class="lede">${res.length ? plural(res.length, "match", "matches") : "Nothing matches. Try part of a word, or a friendly name such as “email” or “store”."}</p>
    ${groups.map(([t, xs]) => `<h2>${TYPE_LABEL[t]}s <span class="count">${xs.length}</span></h2><div class="panel" style="padding:6px 14px">${xs.slice(0, 100).map(e => `<div style="padding:5px 0;border-bottom:1px solid var(--line)"><a href="${href[e.t](e.i)}">${h(e.title)}</a> <span class="sub">${h(e.sub)}</span></div>`).join("")}${xs.length > 100 ? `<div class="hint">…and ${xs.length - 100} more; add another word to narrow it down.</div>` : ""}</div>`).join("")}`;
}
function pageHelp() {
  return `<h1>How to read this</h1><p class="lede">Everything here was read from Adobe Experience Platform (${h(M.label)}) by the Data Dictionary. It shows what data exists, how good it is, and what uses it. Nothing here changes anything in Adobe.</p>
  <div class="grid g2">
   <div class="panel"><h3 style="margin-top:0">The words</h3><dl class="kv">${Object.entries({"Schema": GLOSS.schema, "Dataset": GLOSS.dataset, "Field": GLOSS.field, "Profile": GLOSS.profile, "Data lake": GLOSS.lake, "Audience": GLOSS.audience, "Filled in": GLOSS.coverage}).map(([k, v]) => `<dt><b>${k}</b></dt><dd>${h(v)}</dd>`).join("")}</dl></div>
   <div class="panel"><h3 style="margin-top:0">Following the data (lineage)</h3><p>Each page has a lineage strip. Data flows left to right:</p>
    <p><b>source → dataset → schema → field → audience → journey or destination</b></p>
    <p>Click any box to move there; the lines show which boxes connect. Hover a box to highlight its lines. A greyed box is not live: a dataset only in the data lake (audiences can't read it), a journey that is not running, a dataflow that is switched off.</p>
    <p><b>Where it goes</b> on an audience: <span class="badge GREEN">In use</span> a live journey or destination, or another audience built on it; <span class="badge AMBER">Not live</span> only in journeys not running or dataflows switched off; <span class="badge RED">Nothing found using it</span> a candidate to unpublish – but ask its owner first, as calls from other systems are not seen.</p></div>
   <div class="panel"><h3 style="margin-top:0">A field's three numbers</h3>
    <p><b>Filled in</b>: out of a recent sample of records, the share that had the field. 0% can mean the field is new or unused.</p>
    <p><b>Have we got it? (months)</b>: how many months of unbroken data we hold, read month by month from the data lake.</p>
    <p><b>In Profile now (days)</b>: how far back an audience can look. Audiences read Profile, not the data lake, and Profile keeps each event dataset for a set number of days.</p></div>
   <div class="panel"><h3 style="margin-top:0">The month-by-month strip</h3>${strip([null, 0, "?", 40, 75, 100])}${STRIP_LEGEND}
    <p style="margin-top:10px">A <b>“?”</b> month means data arrived but couldn't be read from the data lake – usually because it went straight into Profile, or the lake keeps only empty placeholder files for it. It is <b>not</b> the same as “no data”. To check by hand, open the field and use <b>SQL to check it yourself</b>.</p></div>
  </div>
  <p class="foot">Built from <b>${h(M.workbook)}</b> (${h(M.generated || "")}) by ${h(M.builder)} on ${h(M.built)}. Prototype – tell Barry what's missing or confusing.</p>`;
}

/* ---------- router ---------- */
function route() {
  LIN = []; FILTERS = [];
  for (const k in TABLES) delete TABLES[k];
  const parts = location.hash.replace(/^#\/?/, "").split("/");
  const [kind, ...rest] = parts, key = decodeURIComponent(rest.join("/"));
  let html, nav = kind || "";
  const miss = what => `<h1>Not found</h1><p class="lede">There is no ${what} called “${h(key)}” in this dictionary. <a href="#/">Back to the start</a>.</p>`;
  if (!kind) html = pageHome();
  else if (kind === "field") { html = by.field[key] !== undefined ? pageField(by.field[key]) : miss("field"); nav = "fields"; }
  else if (kind === "schema") { html = by.schema[key] !== undefined ? pageSchema(by.schema[key]) : miss("schema"); nav = "schemas"; }
  else if (kind === "dataset") { html = by.dataset[key] !== undefined ? pageDataset(by.dataset[key]) : miss("dataset"); nav = "datasets"; }
  else if (kind === "audience") { html = by.audience[key] !== undefined ? pageAudience(by.audience[key]) : miss("audience"); nav = "audiences"; }
  else if (kind === "journey") { html = by.journey[key] !== undefined ? pageJourney(by.journey[key]) : miss("journey"); nav = "journeys"; }
  else if (kind === "source") { html = by.source[key] !== undefined ? pageSource(by.source[key]) : miss("source"); nav = "sources"; }
  else if (kind === "destination") { html = by.destination[key] !== undefined ? pageDestination(by.destination[key]) : miss("destination"); nav = "journeys"; }
  else if (kind === "journeys") html = pageJourneys();
  else if (kind === "sources") html = pageSources();
  else if (kind === "destinations") { html = pageDestinations(); nav = "journeys"; }
  else if (kind === "fields") html = pageFields();
  else if (kind === "schemas") html = pageSchemas();
  else if (kind === "datasets") html = pageDatasets();
  else if (kind === "audiences") html = pageAudiences(rest[0]);
  else if (kind === "search") html = pageSearch(key);
  else if (kind === "help") html = pageHelp();
  else html = miss("page");
  app.innerHTML = html + `<div class="foot">Data Dictionary Explorer (prototype) · ${h(M.label)} · from ${h(M.workbook)} · ${h(M.generated || "")} · Barry Mann (barrymann.com)</div>`;
  document.querySelectorAll("#nav a").forEach(a => a.classList.toggle("on", a.getAttribute("href") === "#/" + nav));
  const big = document.getElementById("bigsearch");
  if (big) mountSearch(big, true);
  Object.keys(TABLES).forEach(id => FILTERS.find(f => f.id === id) ? applyFilter(id) : drawTable(id));
  window.scrollTo(0, 0);
  requestAnimationFrame(drawLineage);
  document.title = (app.querySelector("h1") ? app.querySelector("h1").textContent + " – " : "") + "Data Dictionary Explorer – " + M.label;
}
/* light / dark / follow the computer; ?theme=light also works */
const THEMES = ["auto", "light", "dark"];
let THEME = "auto";
function setTheme(t) {
  if (t === "auto") document.documentElement.removeAttribute("data-theme"); else document.documentElement.setAttribute("data-theme", t);
  document.getElementById("theme").title = `Theme: ${t} (click for light, dark, auto)`;
  THEME = t;
  try { localStorage.setItem("dd-theme", t); } catch (e) { }
  requestAnimationFrame(drawLineage);
}
let theme = new URLSearchParams(location.search).get("theme");
if (!THEMES.includes(theme)) { try { theme = localStorage.getItem("dd-theme"); } catch (e) { theme = null; } }
setTheme(THEMES.includes(theme) ? theme : "auto");
function setAA(show) {
  SHOW_AA = show;
  try { localStorage.setItem("dd-aa", show ? "show" : "hide"); } catch (e) { }
  const b = document.getElementById("aa");
  b.textContent = show ? "AA shown" : "AA hidden";
  b.classList.toggle("on", show);
  b.title = show ? "Adobe Analytics report-suite data is shown; click to hide it" : "Adobe Analytics report-suite data is hidden (it is in the workbook); click to show it";
}
if (AA_S.size) { document.getElementById("aa").hidden = false; setAA(SHOW_AA); }
document.getElementById("aa").onclick = () => { setAA(!SHOW_AA); route(); };
document.addEventListener("click", ev => {
  const a = ev.target.closest("[data-aa]");
  if (!a) return;
  ev.preventDefault();
  setAA(a.dataset.aa === "show");
  route();
});
document.getElementById("theme").onclick = () => setTheme(THEMES[(THEMES.indexOf(THEME) + 1) % 3]);
document.getElementById("brand-sub").textContent = `${M.label} · ${M.generated ? M.generated.split("|")[0].replace("Generated", "").trim() : ""}`;
window.addEventListener("hashchange", route);
document.addEventListener("click", ev => {
  const a = ev.target.closest('a[href^="#f-"],a[href^="#s-"],a[href^="#a-"],a[href^="#d-"],a[href^="#j-"],a[href^="#x-"]');
  if (!a) return;
  ev.preventDefault();
  const t = document.getElementById(a.getAttribute("href").slice(1));
  if (t) t.scrollIntoView({behavior: "smooth", block: "start"});
});
route();
/* ?selftest renders every page once and lists any that fail (used by the build check) */
if (new URLSearchParams(location.search).has("selftest")) {
  const pages = ["#/", "#/fields", "#/schemas", "#/datasets", "#/audiences", "#/audiences/RED", "#/audiences/unused",
    "#/journeys", "#/destinations", "#/sources", "#/help", "#/search/coupon",
    ...D.fields.map((_, i) => href.f(i)), ...D.schemas.map((_, i) => href.s(i)),
    ...D.datasets.map((_, i) => href.d(i)), ...D.audiences.map((_, i) => href.a(i)),
    ...D.journeys.map((_, i) => href.j(i)), ...D.sources.map((_, i) => href.r(i)), ...D.destinations.map((_, i) => href.x(i))];
  const bad = [];
  for (const p of pages) {
    try {
      history.replaceState(null, "", p);
      route(); drawLineage();
      if (/^Not found/.test((app.querySelector("h1") || {}).textContent || "")) bad.push(p + "  not found");
    } catch (e) { bad.push(p + "  " + e.message); }
  }
  document.body.insertAdjacentHTML("beforeend", `<pre id="selftest">SELFTEST ${pages.length} pages, ${bad.length} problems\n${h(bad.slice(0, 40).join("\n"))}</pre>`);
}
</script>
</body>
</html>
"""


if __name__ == "__main__":
    main()
