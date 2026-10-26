"""Offline tests: the Data Dictionary workbook and its Explorer share ONE colour
scheme (data_dictionary/dictionary_colours.py), and categorical columns are
pills in both."""
import sys
from pathlib import Path

from openpyxl import Workbook

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))   # the repo root
import data_dictionary_v3 as dd  # noqa: E402
from data_dictionary import dictionary_colours as colours  # noqa: E402
from data_dictionary import data_dictionary_explorer_html_builder as ex  # noqa: E402

results = []


def check(name, ok, got=""):
    results.append(bool(ok))
    print(("PASS" if ok else "FAIL"), name, got)


# ---- the workbook takes its colours from the palette
check("tab colours by class", dd.TAB_COLOURS[("profile", True)] == colours.KIND["profile"][0]
      and dd.TAB_COLOURS[("event", False)] == colours.KIND["event"][1])
check("header blue", dd._HEADER_BG == colours.HEADER)
check("bands", dd.BAND_DATA[1:] == colours.BANDS["data"])
check("colour scales", dd.SCALE_COLUMNS["Coverage %"] == colours.COVERAGE_SCALE)

# ---- which pill a value gets
check("live: green, not live: grey", colours.pill_for("Live now", "Yes") == colours.RAG["GREEN"]
      and colours.pill_for("Live now", "No") == colours.RAG["GREY"])
check("journey status: running green, trying amber, stopped grey",
      [colours.pill_for("Status", s) for s in ("Deployed", "Dry Run", "Stopped")]
      == [colours.RAG["GREEN"], colours.RAG["AMBER"], colours.RAG["GREY"]])
check("kind by class", colours.pill_for("Kind", "Event") == colours.PILL["event"]
      and colours.pill_for("Kind", "Custom") == colours.PILL["other"])
check("no pill for an empty cell or another column",
      colours.pill_for("Live now", None) is None and colours.pill_for("Audience name", "x") is None)
check("scale colour matches the ends and the middle", colours.scale_colour(colours.COVERAGE_SCALE, 0) == "F8CBAD"
      and colours.scale_colour(colours.COVERAGE_SCALE, 50) == "FFE699"
      and colours.scale_colour(colours.COVERAGE_SCALE, 100) == "A9D08E")

# ---- the finishing pass draws pills in the workbook
wb = Workbook()
ws = wb.active
ws.title = "Journeys"
for c, h in enumerate(["Sandbox", "Journey", "Status", "Live now"], 1):
    ws.cell(5, c, h)
ws._dd_table = (5, 4)
ws.append(["ROI Prod", "Weekly", "Deployed", "Yes"])
ws.append(["ROI Prod", "Old", "Stopped", "No"])
ws.auto_filter.ref = "A5:D7"
dd._polish_workbook(wb)
fill = lambda c: ws[c].fill.fgColor.rgb[-6:]          # noqa: E731
check("workbook: Live now 'Yes' is a green pill", fill("D6") == colours.RAG["GREEN"][0]
      and ws["D6"].font.bold and ws["D6"].font.color.rgb[-6:] == colours.RAG["GREEN"][1])
check("workbook: Stopped is a grey pill", fill("C7") == colours.RAG["GREY"][0])
check("workbook: a name is not a pill", ws["B6"].fill.fgColor.rgb[-6:] in ("000000", "FFFFFF") or ws["B6"].fill.fill_type is None)

# ---- the Explorer uses the same hexes
css = ex.palette_css()
check("explorer: class colours", f"--k-profile:#{colours.KIND['profile'][0]};" in css)
check("explorer: traffic lights", f"--p-green-bg:#{colours.RAG['GREEN'][0]};" in css and f"--p-red:#{colours.RAG['RED'][1]};" in css)
check("explorer: pills for journeys, sources, destinations",
      all(f"--p-{k}-bg:#{colours.PILL[k][0]};" in css for k in ("journey", "source", "destination")))
check("explorer: header blue", f"--header:#{colours.HEADER};" in css)
page = ex.render_html({"meta": {"label": "ROI Prod"}, "schemas": [], "datasets": [], "fields": [], "audiences": [],
                       "journeys": [], "destinations": [], "sources": []})
check("explorer: palette and coverage scale in the page", css in page and '"coverage":[[0,"F8CBAD"]' in page)

print(f"\n{sum(results)}/{len(results)} passed")
sys.exit(0 if all(results) else 1)
