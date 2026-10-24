#!/usr/bin/env python3
"""
tests/field_history_independent_check.py  (AEP Swiss Army Knife)
================================================================
Is the Data Dictionary's field history RIGHT? After a `--history` run, this
re-asks Adobe the same questions a different way and says, verdict by verdict,
whether the answer holds. Read-only: Catalog and Data Access GETs only.

  A  every "Profile only" month   walk EVERY page of the month: is there really
                                  no readable data-lake batch with records?
  B  every "data-lake files       read up to 3 of the biggest and 3 of the newest
     empty" month                 batches the run did not try: still no rows?
  C  every dataset's start date   the month before the first one with data has
                                  no batch with records (or Profile-side) at all
  D  every empty month between    walk every page: really nothing readable?
     two months with data
  E  "all N days" in Profile      a sample of fields: in every month of the last
                                  N days a Profile-enabled dataset keeping >= N
                                  days recorded the field, and re-reading the
                                  batches it read finds the field again
  F  Summary retention lines      match the dataset's raw Catalog record

Run it after a run, against the same sandbox:
    python tests/field_history_independent_check.py --sandbox=roi-prod
    python tests/field_history_independent_check.py --sandbox=prod --credential=aep-prod --limit=40
--limit=N checks at most N verdicts per kind (a big sandbox has hundreds);
--workbook= points at a workbook other than the newest for the sandbox.
Exit code 0 when every verdict checked holds, 1 when any does not.
"""

from __future__ import annotations

import json
import os
import random
import sys
from datetime import datetime, timedelta, timezone
from pathlib import Path

REPO = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(REPO))
import data_dictionary_v3 as dd  # noqa: E402
from toolkit_internals import credential_loader  # noqa: E402

SCRIPT_NAME = "field_history_independent_check"
SCRIPT_VERSION = "1.0.0"
SCRIPT_DATE = "2026-10-24"
SCRIPT_AUTHOR = "Barry Mann (barrymann.com)"


def opt(name, default=None):
    return next((a.split("=", 1)[1] for a in sys.argv[1:] if a.startswith(f"--{name}=")), default)


def main():
    if "-h" in sys.argv[1:] or "--help" in sys.argv[1:] or not opt("sandbox"):
        print(__doc__)
        return
    sandbox, limit = opt("sandbox"), int(opt("limit", "0") or 0)
    cache = dd.HISTORY_CACHE_DIR / sandbox
    out = Path(os.environ.get("AEP_OUTPUT_DIR") or REPO / "output") / "data_dictionary"
    wb_path = Path(opt("workbook")) if opt("workbook") else max(
        out.glob(f"Data Dictionary - {sandbox} - *.xlsx"), key=lambda p: p.stat().st_mtime, default=None)
    if not cache.is_dir() or not wb_path:
        sys.exit(f"No history cache in {cache} or no workbook in {out} -- run the Data "
                 f"Dictionary with --history on {sandbox} first.")
    print(f"{SCRIPT_NAME} v{SCRIPT_VERSION} ({SCRIPT_DATE})  by {SCRIPT_AUTHOR}")
    print(f"sandbox {sandbox}  |  cache {cache}  |  workbook {wb_path.name}  (read-only)\n")

    conf = credential_loader.load_creds(credential_loader.pick_service(opt("credential", "aep-prod")))
    tok = dd.authenticate(conf)["access_token"]
    H = dd.aep_headers(tok, conf, sandbox)
    _, ds = dd.get_all_datasets(tok, conf, sandbox)
    name = {d["id"]: d["table"] for d in ds}
    by_table = {d["table"]: d for d in ds}
    months = {f.stem: json.loads(f.read_text(encoding="utf-8"))["months"] for f in cache.glob("*.json")}
    tally = {}
    random.seed(1)

    def score(part, good, line=None):
        t = tally.setdefault(part, [0, 0])
        t[0 if good else 1] += 1
        if line and (not good or part in "AD"):
            print(f"   {line}")

    def capped(items):
        items = list(items)
        return random.sample(items, limit) if limit and len(items) > limit else items

    def window(lab):
        y, m = map(int, lab.split("-"))
        return (dd._ms(datetime(y, m, 1, tzinfo=timezone.utc)),
                dd._ms(datetime(y + (m == 12), m % 12 + 1, 1, tzinfo=timezone.utc)))

    def q(dsid, lab, extra):
        a, b = window(lab)
        d = json.loads(dd.http(f"{dd.CATALOG_BATCHES_URL}?dataSet={dsid}&status=success"
                               f"&createdAfter={a}&createdBefore={b}{extra}", headers=H, timeout=60)[0])
        return [(bid, x) for bid, x in d.items() if a <= int(x.get("created") or 0) < b]

    def every_batch(dsid, lab):
        got, start = [], 0
        while True:
            page = q(dsid, lab, f"&orderBy=desc:created&limit=100&start={start}")
            got += page
            if len(page) < 100 or start >= 5000:
                return got
            start += 100

    def with_records(bs):
        return [(bid, x) for bid, x in bs if dd._lake_readable(bid, x) and dd._batch_record_count(x) > 0]

    def rows_of(bid, n=200):
        return dd._download_batch_rows(H, bid, n, smallest_first=True, quiet=True,
                                       listing=dd.HISTORY_LISTING)[0]

    def nm(dsid):
        return name.get(dsid, dsid)[:46]

    print("A. 'Profile only' months: every page walked")
    for dsid, lab in capped((k, lab) for k, ms in months.items() for lab, r in sorted(ms.items())
                            if str(r.get("why") or "").startswith("Profile only")):
        rb = with_records(every_batch(dsid, lab))
        score("A", not rb, f"{nm(dsid):<46} {lab}: " + ("CONFIRMED" if not rb else
                                                         f"WRONG: {len(rb)} readable batch(es) with records"))

    print("\nB. 'data-lake files empty' months: untried biggest and newest batches")
    for dsid, lab in capped((k, lab) for k, ms in months.items() for lab, r in sorted(ms.items())
                            if r.get("why") == "data-lake files empty"):
        tried = set(months[dsid][lab].get("batches") or [])
        try_ = [kv for kv in with_records(q(dsid, lab, "&orderBy=desc:metrics.outputRecordCount&limit=20"))
                if kv[0] not in tried][:3]
        try_ += [kv for kv in q(dsid, lab, "&orderBy=desc:created&limit=50")
                 if dd._lake_readable(*kv) and dd._batch_record_count(kv[1]) != 0
                 and kv[0] not in tried and kv not in try_][:3]
        hit = next(((bid, n) for bid, _ in try_ if (n := len(rows_of(bid, 20)))), None)
        score("B", not hit, f"{nm(dsid):<46} {lab}: WRONG, {hit[1]} row(s) in {hit[0]}" if hit else None)
    print(f"   {tally.get('B', [0, 0])[0]} confirmed empty, {tally.get('B', [0, 0])[1]} wrong")

    print("\nC. each dataset's start: the month before has nothing")
    for dsid, ms in capped(months.items()):
        labs = sorted(lab for lab, r in ms.items() if r.get("state") in ("sampled", "unsure")
                      and (r.get("batches") or str(r.get("why") or "").startswith("Profile only")))
        if not labs or labs[0] == min(ms):
            continue
        y, m = map(int, labs[0].split("-"))
        prev = f"{y - (m == 1):04d}-{(m - 2) % 12 + 1:02d}"
        before = [kv for kv in every_batch(dsid, prev) if dd._lake_readable(*kv)
                  and dd._batch_record_count(kv[1]) > 0 or dd._is_profile_side_batch(kv[1])]
        score("C", not before, f"{nm(dsid):<46} starts {labs[0]}: WRONG, {prev} has "
                               f"{len(before)} batch(es) with data" if before else None)
    print(f"   {tally.get('C', [0, 0])[0]} confirmed, {tally.get('C', [0, 0])[1]} wrong")

    print("\nD. empty months between months with data: every page walked")
    gaps = [(k, c) for k, ms in months.items() for p, c, n in zip(sorted(ms), sorted(ms)[1:], sorted(ms)[2:])
            if ms[c].get("state") == "none" and ms[p].get("state") != "none" and ms[n].get("state") != "none"]
    for dsid, lab in capped(gaps):
        rb = with_records(every_batch(dsid, lab))
        score("D", not rb, f"{nm(dsid):<46} {lab}: " + ("CONFIRMED empty" if not rb else
                                                         f"WRONG: {len(rb)} batch(es) with records"))

    print("\nE. 'all N days' in Profile: the field recorded every month, and re-read")
    from openpyxl import load_workbook
    wb = load_workbook(wb_path, read_only=True)
    rows = list(wb["Field Index"].iter_rows(min_row=5, values_only=True))
    ix = {h: i for i, h in enumerate(rows[0])}
    full = [r for r in rows[1:] if str(r[ix["Profile detail (BETA)"]] or "").startswith("all ")]
    now = datetime.now(timezone.utc)
    for r in random.sample(full, min(limit or 10, len(full))):
        path, days = r[0], int(r[ix["In Profile now (days, BETA)"]] or 0)
        labs = sorted({f"{now - timedelta(days=d):%Y-%m}" for d in range(0, days + 1, 5)})
        tables = [t.strip() for t in str(r[ix["SQL table name(s)"]] or "").split(",")]
        proof = None
        for d in (by_table[t] for t in tables if t in by_table):
            if not (str(d.get("profile") or "").startswith("Profile")
                    and (d.get("profile_ttl_days") or 0) >= days and d["id"] in months):
                continue
            ms = months[d["id"]]
            done = [lab for lab in labs if lab in ms]          # the current month is not cached
            if done and all((ms[lab].get("fields") or {}).get(path) for lab in done):
                again = sum(1 for bid in ms[done[-1]].get("batches") or []
                            for row in rows_of(bid, dd.HISTORY_ROWS) if dd._extract_values(row, path))
                if again:
                    proof = f"{d['table']} ({done[0]} > {done[-1]}; re-read {done[-1]}: {again} rows)"
                    break
        score("E", bool(proof), None)
        print(f"   {path[-52:]:<52} {'CONFIRMED via ' + proof if proof else 'NOT CONFIRMED'}")

    print("\nF. Summary retention lines vs the raw Catalog record")
    for row in wb["Summary"].iter_rows(values_only=True):
        line = str(row[0] or "")
        if "Profile keeps" not in line or line.startswith("PROFILE RETENTION"):
            continue
        table = line.split(":")[0].split("  (")[0].strip()
        raw = json.loads(dd.http(f"{dd.PLATFORM}/data/foundation/catalog/dataSets/{by_table[table]['id']}",
                                 headers=H, timeout=30)[0])
        ext = next(iter(raw.values())).get("extensions") or {}
        p = ((ext.get("adobe_unifiedProfile") or {}).get("rowExpiration") or {}).get("ttlValue")
        lk = ((ext.get("adobe_lakeHouse") or {}).get("rowExpiration") or {}).get("ttlValue")
        good = (f"Profile keeps {dd._iso_duration_days(p)} days" in line
                and f"data lake: {dd._ttl_words(lk)}" in line)
        score("F", good, f"{table[:46]:<46} WRONG: Catalog says Profile {p}, lake {lk}" if not good else None)
    print(f"   {tally.get('F', [0, 0])[0]} confirmed, {tally.get('F', [0, 0])[1]} wrong")

    wrong = sum(v[1] for v in tally.values())
    print("\nTOTAL  " + "  |  ".join(f"{k}: {v[0]} held, {v[1]} wrong" for k, v in sorted(tally.items())))
    sys.exit(1 if wrong else 0)


if __name__ == "__main__":
    main()
