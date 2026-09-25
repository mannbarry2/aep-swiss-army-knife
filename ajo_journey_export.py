"""
ajo_journey_export.py  --  AJO journey definitions, every node, for a sandbox
==========================================================================

Python port of Craig's ajo-journey-export/ajo-journey-export.mjs (kept in the
repo for provenance; the two produce interchangeable JSON + index.csv). Pulls
Adobe Journey Optimizer journey definitions -- every node/step, i.e. the same
content as the UI's "Copy technical details" -- for a whole sandbox, straight
from the AJO authoring API, and writes:

    output/ajo-journeys_<sandbox>_<date>/
        <JourneyName>__<uid>.json     full definition per journey (all nodes)
        index.csv                     Craig's summary (unchanged columns)
    output/AJO Journeys - <sandbox> - <date>.xlsx
        the human-readable version, in the Data Dictionary house style:
        Summary / Journeys / Steps / Audiences / Failed (see below)

!!  UNSUPPORTED ENDPOINT. journey-private.adobe.io/authoring is the private API
    behind the AJO UI. It is not an official/public Adobe API and can change
    without notice. Use it for inspection/reporting only; do not build critical
    automation on it. Strictly read-only (GET only) -- it never modifies a
    journey. There is a supported Journey Public API in development at Adobe
    (GET /ajo/journey/{id}?include=...) which is the long-term route.

!!  This is a UI hack and porting it to Python does not change that: the
    endpoint only accepts a USER (shell) bearer token, paired with the UI's own
    client id 'voyager_ui' as x-api-key. The service-account credentials in the
    keyring are NOT entitled, so the token has to be lifted from the browser.

HOW TO GET THE TOKEN (every run -- it expires within ~24h)
    1. Sign in to Adobe Experience Cloud in Chrome/Edge and open Journey
       Optimizer -> Journeys, in the sandbox you want to export.
    2. Press F12 (DevTools) -> Network tab -> tick "Fetch/XHR". If the list is
       empty, reload the Journeys page so requests appear.
    3. In the filter box type   journey-private   and click any request whose
       URL starts with journey-private.adobe.io/authoring/ (e.g. journeys/).
    4. In the right-hand pane open "Headers" -> "Request Headers".
         authorization:      Bearer eyJhbGci....   <- copy everything AFTER "Bearer "
         x-gw-ims-org-id:    XXXXXXXX@AdobeOrg     <- only needed without the keyring
       Tip: right-click the request -> Copy -> Copy as cURL also gives you
       both headers in one go.
    5. Give the token to this script one of three ways:
         set AJO_BEARER_TOKEN=eyJ...          (PowerShell: $env:AJO_BEARER_TOKEN="eyJ...")
         --token=eyJ...
         or just run it and paste at the masked prompt (pasting "Bearer eyJ..." is fine).
    A 401 / 403 means the token has expired or belongs to a different org --
    grab a fresh one. Never commit the token anywhere; it is your user session.

What IS reused from the rest of the knife: the keyring credential (org id for
the private endpoint, and -- because the audiences API accepts the service
account -- resolving every audience UUID to its name, description and
evaluation type), the sandbox picker, the http helper, colours and logging.

Usage:
    set AJO_BEARER_TOKEN=eyJ...            (or --token=..., or paste at the prompt)
    python ajo_journey_export.py                         # creds menu; sandbox menu; ALL journeys
    python ajo_journey_export.py aep-prod --sandbox=prod
    python ajo_journey_export.py aep-prod --sandbox=prod --journey=<journeyUid>
    python ajo_journey_export.py aep-prod --sandbox=dev --out=./export-2026-09
    python ajo_journey_export.py --org-id=XXXX@AdobeOrg --sandbox=prod          # no keyring at all
    python ajo_journey_export.py aep-prod --sandbox=dev --from-dir=output/ajo-journeys_dev_2026-09-25
                                                         # rebuild the workbook from an earlier
                                                         # export -- no token needed
    --no-xlsx   skip the workbook (JSON + index.csv only, exactly Craig's output)

The workbook:
    Summary    what was exported, journeys by state, steps by type, channels,
               audiences referenced (resolved / not found), the failures.
    Journeys   one row per journey: name, state, description, how it starts,
               channels used, the audiences it reads (NAMES, with the UUIDs
               alongside), step counts by type, who created / changed /
               deployed / stopped it and when, the stale list name if any.
    Steps      one row per node, in flow order from the start node: type, name,
               a plain-English detail (wait 15 min; email message; condition
               branches with the expression rendered and audiences named), the
               branches and where they go.
    Audiences  one row per audience referenced by any journey: name, description,
               evaluation (Batch / Streaming / Edge), lifecycle, UUID, and which
               journeys use it.
    Failed     journeys the list knew about but /latest could not return
               (abandoned duplications, typically).

Gotchas (Craig's, all still true):
  * Sandbox name is case-sensitive and must be lowercase -- "Prod" silently
    returns an empty result set. This script lowercases it and warns.
  * The list endpoint can return a stale journey-level name (..._Copy_Copy90);
    the real name is result.name from /latest. Both are reported.
  * Authoring format 1.0 and 2.0 both carry inSegment/inAudience in steps.
"""

from __future__ import annotations

import base64
import csv
import getpass
import json
import os
import re
import sys
import urllib.error
import urllib.parse
from collections import Counter, deque
from concurrent.futures import ThreadPoolExecutor
from datetime import datetime, timezone
from pathlib import Path

import aep_creds
import journey_audience_census as census   # http(), pick_sandbox(), audiences, ANSI, logger

SCRIPT_VERSION = "1.1"
SCRIPT_DATE = "2026-09-25"
SCRIPT_DIR = Path(__file__).resolve().parent
OUTPUT_DIR = SCRIPT_DIR / "output"

HOST = os.environ.get("AJO_HOST") or "https://journey-private.adobe.io/authoring"
DEFAULT_API_KEY = "voyager_ui"          # the AJO UI's own client id
CONCURRENCY = 5
PAGE_SIZE = 100
MAX_PAGES = 500
UUID_RE = re.compile(r"^[0-9a-f]{8}-[0-9a-f]{4}-")
CONFIDENTIAL = "STRICTLY CONFIDENTIAL"
_HEADER_BG = "1F4E78"                   # the Data Dictionary's header blue

ANSI = census.ANSI
logger = census.logger

# Printed at the paste prompt and again on a 401/403, so nobody has to open
# the file to remember the dance. Same steps as the header docstring.
TOKEN_HELP = f"""
  {ANSI['bold']}How to get the token{ANSI['reset']}  {ANSI['dim']}(browser session token; expires within ~24h){ANSI['reset']}
    1. Sign in to Experience Cloud, open Journey Optimizer -> Journeys (in the sandbox you want).
    2. F12 -> Network -> tick Fetch/XHR (reload the Journeys page if the list is empty).
    3. Filter on  journey-private  and click any journey-private.adobe.io/authoring/ request.
    4. Headers -> Request Headers -> authorization: copy everything AFTER "Bearer ".
       (x-gw-ims-org-id is on the same list if you are not using the keyring for the org id.)
    5. Paste it below, or set AJO_BEARER_TOKEN, or pass --token=...  (never commit it).
"""


# ----------------------------------------------------------------------------
# Helpers (1:1 with the .mjs)
# ----------------------------------------------------------------------------
def b64(obj) -> str:
    """Byte-identical to Node's Buffer.from(JSON.stringify(obj)).toString('base64')."""
    return base64.b64encode(json.dumps(obj, separators=(",", ":")).encode()).decode()


def safe_name(s) -> str:
    return re.sub(r"[^\w.-]+", "_", str(s or "journey"))[:120]


def extract_segment_refs(node, out: set) -> None:
    """Recursively pull inSegment / inAudience references out of a definition."""
    if isinstance(node, dict):
        if node.get("function") in ("inSegment", "inAudience"):
            for a in node.get("args") or []:
                v = a.get("value") if isinstance(a, dict) else None
                if isinstance(v, str) and UUID_RE.match(v):
                    out.add(v)
        for v in node.values():
            extract_segment_refs(v, out)
    elif isinstance(node, list):
        for v in node:
            extract_segment_refs(v, out)


def get_json(url, headers, timeout=60):
    try:
        body, _ = census.http(url, headers=headers, timeout=timeout)
    except urllib.error.HTTPError as e:
        detail = e.read().decode("utf-8", "replace")[:300]
        if e.code in (401, 403):
            print(f"\n  {ANSI['red']}HTTP {e.code} -- the token was rejected: expired (~24h), "
                  f"wrong org, or not a browser session token.{ANSI['reset']}", file=sys.stderr)
            print(TOKEN_HELP, file=sys.stderr)
        raise RuntimeError(f"HTTP {e.code} for {url}\n{detail}") from None
    return json.loads(body)


# ----------------------------------------------------------------------------
# 1) enumerate journeys
# ----------------------------------------------------------------------------
def list_journeys(headers, one_journey=None) -> list[dict]:
    if one_journey:
        return [{"uid": one_journey, "name": one_journey}]
    fields = b64([["uid"], ["name"]])
    sorts = b64([{"direction": "ascending", "fields": ["name"]}])
    page, out = 0, []
    for _ in range(MAX_PAGES):
        data = get_json(f"{HOST}/journeys/", {**headers,
                                              "x-vyg-query-fields": fields,
                                              "x-vyg-query-page": str(page),
                                              "x-vyg-query-pagesize": str(PAGE_SIZE),
                                              "x-vyg-query-sorts": sorts})
        results = data.get("results") if isinstance(data.get("results"), list) else []
        out.extend(results)
        pg = data.get("pagination") or {}
        size = pg.get("pageSize") or PAGE_SIZE
        total = pg.get("totalCount") or 0
        sys.stdout.write(f"\r  listed {len(out)}/{pg.get('totalCount', '?')} journeys")
        sys.stdout.flush()
        if (page + 1) * size >= total or not results:
            break
        page += 1
    sys.stdout.write("\n")
    return out


# ----------------------------------------------------------------------------
# 2) fetch each definition + write JSON / index.csv
# ----------------------------------------------------------------------------
def export(headers, sandbox, out_dir: Path, one_journey=None):
    """Returns (records, failures). records: [{uid, list_name, definition, file}]."""
    out_dir.mkdir(parents=True, exist_ok=True)
    print(f"  AJO Journey Exporter -> {out_dir}  (sandbox: {sandbox})")
    journeys = list_journeys(headers, one_journey)
    records, failures = [], []

    def one(j):
        uid = j.get("uid")
        try:
            definition = get_json(f"{HOST}/journeys/{uid}/latest", headers)
            r = definition.get("result") or definition
            file = f"{safe_name(r.get('name') or j.get('name'))}__{uid}.json"
            (out_dir / file).write_text(json.dumps(definition, indent=2), encoding="utf-8")
            records.append({"uid": uid, "list_name": j.get("name"), "definition": definition, "file": file})
            sys.stdout.write(f"\r  exported {len(records)}/{len(journeys)}")
            sys.stdout.flush()
        except Exception as e:  # noqa: BLE001 -- one bad journey must not sink the run
            failures.append({"uid": uid, "list_name": j.get("name"), "error": str(e).splitlines()[0]})
            print(f"\n  ! {uid} ({j.get('name')}): {e}", file=sys.stderr)

    with ThreadPoolExecutor(max_workers=min(CONCURRENCY, max(1, len(journeys)))) as ex:
        list(ex.map(one, journeys))
    sys.stdout.write("\n")
    write_index_csv(records, out_dir)
    print(f"\n  Done. {len(records)} exported, {len(failures)} failed.")
    print(f"    Per-journey JSON: {out_dir}{os.sep}*.json")
    print(f"    Summary:          {out_dir}{os.sep}index.csv")
    return records, failures


def write_index_csv(records, out_dir: Path):
    """index.csv -- Craig's columns, CRLF, same quoting rules as the .mjs."""
    rows = []
    for rec in records:
        r = rec["definition"].get("result") or rec["definition"]
        steps = r.get("steps") if isinstance(r.get("steps"), list) else []
        segs: set = set()
        extract_segment_refs(rec["definition"], segs)
        rows.append([rec["uid"], r.get("name") or rec["list_name"], rec["list_name"],
                     r.get("authoringFormatVersion") or "", len(steps), len(segs),
                     " ".join(sorted(segs)), rec["file"]])
    rows.sort(key=lambda x: (str(x[1]).lower(), x[0]))
    with open(out_dir / "index.csv", "w", newline="", encoding="utf-8") as f:
        w = csv.writer(f, lineterminator="\r\n")
        w.writerow(["journeyUid", "displayName", "listName", "authoringFormatVersion",
                    "nodeCount", "segmentRefCount", "segmentIds", "file"])
        w.writerows(rows)


def load_export_dir(out_dir: Path):
    """Rebuild records from an earlier export (no token needed)."""
    records = []
    for p in sorted(out_dir.glob("*__*.json")):
        try:
            definition = json.loads(p.read_text(encoding="utf-8"))
        except Exception as e:  # noqa: BLE001
            logger.warning(f"skipping {p.name}: {e}")
            continue
        uid = p.stem.rsplit("__", 1)[-1]
        records.append({"uid": uid, "list_name": None, "definition": definition, "file": p.name})
    try:
        for row in csv.DictReader(open(out_dir / "index.csv", encoding="utf-8")):
            for rec in records:
                if rec["uid"] == row["journeyUid"]:
                    rec["list_name"] = row["listName"]
    except FileNotFoundError:
        pass
    return records


# ----------------------------------------------------------------------------
# 3) make it readable
# ----------------------------------------------------------------------------
NODE_LABEL = {"start": "Start", "event": "Event", "intermediate": "Event (wait for)", "condition": "Condition",
              "timer": "Wait", "message": "Message", "campaign": "Campaign", "action": "Action", "end": "End"}


def duration_text(iso: str) -> str:
    """PT15M -> 15 min, P3D -> 3 days, PT1H30M -> 1 hour 30 min."""
    m = re.fullmatch(r"P(?:(\d+)Y)?(?:(\d+)M)?(?:(\d+)W)?(?:(\d+)D)?(?:T(?:(\d+)H)?(?:(\d+)M)?(?:(\d+)S)?)?", iso or "")
    if not m:
        return iso or ""
    parts = []
    for val, unit in zip(m.groups(), ("year", "month", "week", "day", "hour", "min", "sec")):
        if val and int(val):
            n = int(val)
            parts.append(f"{n} {unit}" + ("s" if n != 1 and unit not in ("min", "sec") else ""))
    return " ".join(parts) or iso


def rule_text(rule: dict) -> str:
    t = (rule or {}).get("type")
    if t == "delay":
        return f"wait {duration_text(rule.get('delay', ''))}"
    if t == "fixedDate":
        return f"at {str(rule.get('date', '')).replace('T', ' ')}"
    if t:
        extras = {k: v for k, v in rule.items() if k != "type"}
        return f"{t} {json.dumps(extras)[:80]}" if extras else t
    return "wait (no rule)"


def field_tail(ref: str) -> str:
    """'json://.../_tescostores/x/y' -> 'x.y' (last two segments)."""
    segs = [s for s in re.split(r"[/.]", str(ref or "")) if s and not s.startswith("json:") and not s.startswith("_")]
    return ".".join(segs[-2:]) if segs else str(ref)


INFIX = {"equal": "=", "equalIgnoreCase": "=", "notEqual": "!=", "greater": ">", "greaterEqual": ">=",
         "less": "<", "lessEqual": "<=", "in": "in", "startWith": "starts with", "contain": "contains"}


def render_expr(node, aud_names: dict, depth=0) -> str:
    """Plain-English rendering of a journey condition expression tree."""
    if depth > 8:
        return "…"
    if not isinstance(node, dict):
        return json.dumps(node)[:40] if node is not None else ""
    t = node.get("type")
    if t == "constant":
        v = node.get("value")
        return f'"{v}"' if isinstance(v, str) else str(v)
    if t in ("entityFieldRef", "eventFieldRef", "contextFieldRef", "variableRef"):
        return field_tail(node.get("fieldRef") or node.get("name") or "")
    if t == "function" or "function" in node:
        fn = node.get("function")
        args = [render_expr(a, aud_names, depth + 1) for a in (node.get("args") or [])]
        raw = node.get("args") or []
        if fn in ("inAudience", "inSegment"):
            ids = [a.get("value") for a in raw if isinstance(a, dict) and isinstance(a.get("value"), str) and UUID_RE.match(a["value"])]
            names = [f'"{aud_names.get(i, i)}"' for i in ids]
            return "in audience " + (", ".join(names) or "?")
        if fn in ("and", "or"):
            return "(" + f" {fn.upper()} ".join(a for a in args if a) + ")"
        if fn == "not":
            inner = args[0] if args else "?"
            if inner.startswith("in audience "):
                return "NOT " + inner
            return f"NOT ({inner})"
        if fn in INFIX and len(args) == 2:
            return f"{args[0]} {INFIX[fn]} {args[1]}"
        if fn == "between" and len(args) == 3:
            return f"{args[0]} between {args[1]} and {args[2]}"
        if fn == "now":
            return "now"
        if fn == "nowWithDelta" and args:
            return f"now {args[0]}{(' ' + args[1]) if len(args) > 1 else ''}"
        if fn in ("toDateOnly", "toDateTimeOnly", "toString") and args:
            return args[0]
        if fn == "count" and args:
            return f"count({args[0]})"
        if fn == "random":
            return "random()"
        return f"{fn}({', '.join(args)})"
    return json.dumps(node)[:60]


def flow_order(steps: list, initial: str | None) -> list:
    """Steps in canvas order: breadth-first from the start node along the
    transitions, then anything unreachable in list order."""
    by_id = {s.get("nodeId") or s.get("uid"): s for s in steps}
    start = initial if initial in by_id else next((k for k, s in by_id.items() if s.get("nodeType") == "start"), None)
    order, seen, q = [], set(), deque([start] if start else [])
    while q:
        nid = q.popleft()
        if nid in seen or nid not in by_id:
            continue
        seen.add(nid)
        order.append(by_id[nid])
        for t in by_id[nid].get("transitions") or []:
            tgt = t.get("targetStep")
            if tgt and tgt not in seen:
                q.append(tgt)
    order += [s for k, s in by_id.items() if k not in seen]
    return order


def describe_step(step: dict, names_by_id: dict, aud_names: dict) -> tuple[str, str, list]:
    """(detail, branches, audience_ids) for one node."""
    nt = step.get("nodeType")
    trans = step.get("transitions") or []
    action = step.get("action") or {}
    scheds = [sa for sa in step.get("schedulerActions") or [] if isinstance(sa, dict)]
    auds: set = set()
    extract_segment_refs(step, auds)
    bits = []
    if nt == "start":
        for sa in scheds:
            bits.append(rule_text(sa.get("rule") or {}))
        for t in trans:
            if t.get("segmentMembershipDef"):
                bits.append("on audience qualification")
            elif t.get("elementRef") or t.get("name"):
                if t.get("eventId") != "scheduledNotificationReceived":
                    bits.append(f"on event {t.get('name') or ''}".strip())
        detail = "; ".join(dict.fromkeys(b for b in bits if b)) or "start"
    elif nt in ("event", "intermediate"):
        evs = [t.get("name") for t in trans if t.get("elementRef") or t.get("segmentMembershipDef")]
        waits = [rule_text(sa.get("rule") or {}) for sa in scheds if (sa.get("rule") or {}).get("type")]
        detail = "wait for event " + ", ".join(e for e in evs if e) if evs else "event"
        if waits:
            detail += f" (timeout: {', '.join(waits)})"
    elif nt == "condition":
        if step.get("nodeSubType") == "pathExperimentation":
            detail = "A/B path experiment" + (f" (decision policy {str((step.get('decision') or {}).get('decisionPolicyId', ''))[:8]}…)" if step.get("decision") else "")
        else:
            parts = []
            for t in trans:
                cond = t.get("condition")
                nm = t.get("name") or ""
                if isinstance(cond, dict):
                    parts.append(f"{nm}: {render_expr(cond, aud_names)}" if nm else render_expr(cond, aud_names))
                elif nm:
                    parts.append(f"{nm}: otherwise")
            detail = " | ".join(parts) or "condition"
    elif nt == "timer":
        detail = "; ".join(rule_text(sa.get("rule") or {}) for sa in scheds) or "wait"
    elif nt == "message":
        surf = action.get("surface") or {}
        detail = f"{surf.get('channel', 'message')} message ({action.get('category', '')})".replace(" ()", "")
        if action.get("type") == "inboundActivationAction":
            detail = f"inbound activation ({surf.get('channel', '')})"
        if action.get("messageId"):
            detail += f" -- message {action['messageId'][:8]}…"
        tr = action.get("trackingOptions") or {}
        if tr:
            detail += " -- tracking: " + ", ".join(k.replace("TrackingEnabled", "") for k, v in tr.items() if v)
    elif nt == "campaign":
        detail = "campaign action"
    elif nt == "action":
        at = action.get("type")
        if at == "customActionRef":
            detail = f"custom action {str((action.get('elementRef') or {}).get('uid', ''))[:8]}…"
        elif at == "datasetLookupAction":
            detail = f"dataset lookup {action.get('lookupTarget', '')}"
        else:
            detail = at or "action"
    elif nt == "end":
        detail = "end"
    else:
        detail = nt or ""
    branches = "; ".join(
        f"{t.get('name') or '→'} → {names_by_id.get(t.get('targetStep'), str(t.get('targetStep'))[:8])}"
        for t in trans if t.get("targetStep"))
    return detail, branches, sorted(auds)


def ts(v) -> str:
    """Epoch ms or ISO string -> 'YYYY-MM-DD HH:MM'."""
    if isinstance(v, (int, float)) and v > 1e11:
        return datetime.fromtimestamp(v / 1000, tz=timezone.utc).strftime("%Y-%m-%d %H:%M")
    if isinstance(v, str) and len(v) >= 16:
        return v[:16].replace("T", " ")
    return "" if v is None else str(v)


def resolve_audiences(service: str | None, sandbox: str, ids: set) -> tuple[dict, str]:
    """{id: {name, description, eval, lifecycle}} via the keyring credential
    (the audiences API accepts the service account). Returns (map, note)."""
    if not service or not ids:
        return {}, "audiences not resolved (no keyring credential)" if ids else ""
    try:
        conf = aep_creds.load_creds(service)
        token = census.authenticate(conf)
        api_key = conf.get("api_key") or conf["client_id"]
        auds, complete = census.fetch_all_audiences(token, api_key, conf, sandbox)
    except Exception as e:  # noqa: BLE001
        logger.warning(f"audience lookup failed ({type(e).__name__}: {e}); UUIDs only.")
        return {}, f"audience lookup failed: {type(e).__name__}"
    out = {a["id"]: {"name": a.get("name") or a["id"], "description": a.get("description") or "",
                     "eval": census.audience_eval_type(a), "lifecycle": a.get("lifecycleState") or ""}
           for a in auds if a.get("id")}
    hit = sum(1 for i in ids if i in out)
    note = f"{hit} of {len(ids)} audience UUIDs resolved from {len(auds)} audiences in {sandbox}"
    if not complete:
        note += " (audience list INCOMPLETE)"
    return out, note


def build_workbook(records, failures, sandbox, service, out_path: Path, source_dir: Path):
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment
    from openpyxl.utils import get_column_letter

    head_font = Font(bold=True, color="FFFFFF")
    head_fill = PatternFill("solid", fgColor=_HEADER_BG)
    title_font = Font(bold=True, size=14)
    conf_font = Font(bold=True, size=11, color="C00000")
    note_font = Font(italic=True, color="666666")
    wrap = Alignment(wrap_text=True, vertical="top")
    tables = []

    def confidential(ws):
        ws["A1"] = CONFIDENTIAL
        ws["A1"].font = conf_font
        ws.oddHeader.center.text = f'&"-,Bold"&12&KC00000{CONFIDENTIAL}'
        ws.evenHeader.center.text = ws.oddHeader.center.text

    def header(ws, cols, row):
        for c, nm in enumerate(cols, 1):
            cell = ws.cell(row, c, nm)
            cell.font, cell.fill = head_font, head_fill
            cell.alignment = Alignment(vertical="center", wrap_text=True)
        ws.row_dimensions[row].height = 28
        ws.freeze_panes = ws.cell(row=row + 1, column=1)
        tables.append((ws, row, len(cols)))

    def widths(ws, ws_widths):
        for i, w in enumerate(ws_widths, 1):
            ws.column_dimensions[get_column_letter(i)].width = w

    # ---- gather ---------------------------------------------------------------
    all_ids: set = set()
    for rec in records:
        extract_segment_refs(rec["definition"], all_ids)
    aud, aud_note = resolve_audiences(service, sandbox, all_ids)
    aud_names = {i: v["name"] for i, v in aud.items()}
    datestr = datetime.now().strftime("%Y-%m-%d")
    journeys_rows, steps_rows = [], []
    aud_use: dict = {}
    state_tally, type_tally, channel_tally = Counter(), Counter(), Counter()

    for rec in sorted(records, key=lambda x: ((x["definition"].get("result") or {}).get("name") or "").lower()):
        r = rec["definition"].get("result") or rec["definition"]
        name = r.get("name") or rec["list_name"] or rec["uid"]
        steps = r.get("steps") if isinstance(r.get("steps"), list) else []
        ordered = flow_order(steps, r.get("initialStep"))
        names_by_id = {s.get("nodeId") or s.get("uid"): (s.get("nodeName") or NODE_LABEL.get(s.get("nodeType"), "")) for s in steps}
        meta = r.get("metadata") or {}
        kinds = Counter(s.get("nodeType") for s in steps)
        channels = Counter((s.get("action") or {}).get("surface", {}).get("channel") for s in steps
                           if s.get("nodeType") == "message" and (s.get("action") or {}).get("surface"))
        j_auds: set = set()
        extract_segment_refs(rec["definition"], j_auds)
        for i in j_auds:
            aud_use.setdefault(i, set()).add(name)
        state_tally[r.get("state") or ""] += 1
        type_tally.update(kinds)
        channel_tally.update(channels)
        start = next((s for s in ordered if s.get("nodeType") == "start"), None)
        trigger = describe_step(start, names_by_id, aud_names)[0] if start else ""
        journeys_rows.append([
            name, r.get("state") or "", r.get("reviewState") or "", r.get("description") or "", trigger,
            ", ".join(f"{k}×{v}" for k, v in channels.most_common()),
            "; ".join(sorted(aud_names.get(i, f"(unresolved) {i}") for i in j_auds)),
            " ".join(sorted(j_auds)),
            len(steps), kinds["condition"], kinds["timer"], kinds["message"], kinds["action"] + kinds["campaign"],
            kinds["event"] + kinds["intermediate"],
            meta.get("createdBy") or "", ts(meta.get("createdAt")), meta.get("lastModifiedBy") or "",
            ts(meta.get("lastModifiedAt")), ts(meta.get("lastDeployedAt")), ts(meta.get("stoppedAt")),
            ", ".join(str(t) for t in (r.get("tags") or [])),
            rec["list_name"] if rec["list_name"] and rec["list_name"] != name else "",
            rec["uid"], r.get("authoringFormatVersion") or "", rec["file"],
        ])
        for n, s in enumerate(ordered, 1):
            detail, branches, s_auds = describe_step(s, names_by_id, aud_names)
            steps_rows.append([name, n, NODE_LABEL.get(s.get("nodeType"), s.get("nodeType") or ""),
                               s.get("nodeName") or "", detail, branches,
                               "; ".join(aud_names.get(i, f"(unresolved) {i}") for i in s_auds),
                               s.get("nodeId") or s.get("uid") or ""])

    wb = Workbook()
    # ---- Summary --------------------------------------------------------------
    ws = wb.active
    ws.title = "Summary"
    confidential(ws)
    ws["A2"] = f"AJO Journeys -- {sandbox}  ({datestr})"
    ws["A2"].font = title_font
    ws["A3"] = ("Every Adobe Journey Optimizer journey in the sandbox with its full node-by-node definition -- the same "
                "content as the UI's 'Copy technical details' -- made readable: what starts it, the audiences it reads "
                "(by name), the channels it uses, every step in flow order with conditions rendered in plain English, "
                "and who built / changed / deployed it. Read from the private authoring API the AJO UI itself uses "
                "(unsupported; read-only). Journeys tab = one row per journey; Steps tab = one row per node; Audiences "
                "tab = every audience referenced. Use the header filters.")
    ws["A3"].font = note_font
    ws["A3"].alignment = wrap
    ws.merge_cells("A3:H3")
    ws.row_dimensions[3].height = 78
    r0 = 5
    ws.cell(r0, 1, "Sandbox"); ws.cell(r0, 2, sandbox)
    ws.cell(r0 + 1, 1, "Journeys exported"); ws.cell(r0 + 1, 2, len(records))
    ws.cell(r0 + 2, 1, "Journeys failed (see Failed tab)"); ws.cell(r0 + 2, 2, len(failures))
    ws.cell(r0 + 3, 1, "Steps (nodes) in total"); ws.cell(r0 + 3, 2, sum(type_tally.values()))
    ws.cell(r0 + 4, 1, "Audiences referenced"); ws.cell(r0 + 4, 2, len(all_ids))
    ws.cell(r0 + 5, 1, "Audience resolution"); ws.cell(r0 + 5, 2, aud_note or "n/a")
    ws.cell(r0 + 6, 1, "Source"); ws.cell(r0 + 6, 2, str(source_dir))
    ws.cell(r0 + 7, 1, "Generated by"); ws.cell(r0 + 7, 2, f"ajo_journey_export.py v{SCRIPT_VERSION} ({SCRIPT_DATE})")
    for rr in range(r0, r0 + 8):
        ws.cell(rr, 1).font = Font(bold=True)
    r1 = r0 + 9
    header(ws, ["Journey state", "Journeys"], r1)
    for i, (k, v) in enumerate(state_tally.most_common(), 1):
        ws.cell(r1 + i, 1, k or "(none)"); ws.cell(r1 + i, 2, v)
    r2 = r1 + len(state_tally) + 2
    header(ws, ["Step type", "Nodes"], r2)
    for i, (k, v) in enumerate(type_tally.most_common(), 1):
        ws.cell(r2 + i, 1, NODE_LABEL.get(k, k or "(none)")); ws.cell(r2 + i, 2, v)
    r3 = r2 + len(type_tally) + 2
    header(ws, ["Message channel", "Message nodes"], r3)
    for i, (k, v) in enumerate(channel_tally.most_common(), 1):
        ws.cell(r3 + i, 1, k or "(none)"); ws.cell(r3 + i, 2, v)
    widths(ws, [34, 60])

    # ---- Journeys -------------------------------------------------------------
    js = wb.create_sheet("Journeys")
    confidential(js)
    js["A2"] = f"Journeys -- {sandbox}"; js["A2"].font = title_font
    js["A3"] = ("One row per journey. 'Starts' is the start node in plain English (schedule, event, or audience "
                "qualification). Audiences are named; the UUID column carries the ids. 'List name' is filled only "
                "when the journeys list still shows a stale name (a duplication artefact) different from the real one.")
    js["A3"].font = note_font
    cols = ["Journey", "State", "Review state", "Description", "Starts", "Channels", "Audiences (names)",
            "Audience UUIDs", "Steps", "Conditions", "Waits", "Messages", "Actions", "Events",
            "Created by", "Created", "Last modified by", "Last modified", "Last deployed", "Stopped",
            "Tags", "List name (stale)", "Journey UID", "Format", "File"]
    header(js, cols, 5)
    for i, row in enumerate(journeys_rows, 6):
        for c, v in enumerate(row, 1):
            cell = js.cell(i, c, v)
            if c in (4, 5, 7):
                cell.alignment = wrap
    widths(js, [46, 11, 14, 40, 44, 16, 50, 40, 7, 10, 7, 9, 8, 7, 22, 16, 22, 16, 16, 16, 18, 40, 38, 7, 40])

    # ---- Steps ----------------------------------------------------------------
    ss = wb.create_sheet("Steps")
    confidential(ss)
    ss["A2"] = f"Steps -- every node of every journey, in flow order  --  {sandbox}"; ss["A2"].font = title_font
    ss["A3"] = ("Filter on Journey to read one journey top to bottom. Detail is a plain-English rendering of the node: "
                "waits with their duration, messages with channel and tracking, conditions with each branch's rule "
                "and the audiences it tests (by name). Branches show where each transition leads.")
    ss["A3"].font = note_font
    header(ss, ["Journey", "#", "Type", "Node name", "Detail", "Branches → next node", "Audiences (names)", "Node id"], 5)
    for i, row in enumerate(steps_rows, 6):
        for c, v in enumerate(row, 1):
            cell = ss.cell(i, c, v)
            if c in (5, 6, 7):
                cell.alignment = wrap
    widths(ss, [40, 5, 14, 40, 70, 46, 40, 38])

    # ---- Audiences ------------------------------------------------------------
    au = wb.create_sheet("Audiences")
    confidential(au)
    au["A2"] = f"Audiences referenced by journeys  --  {sandbox}"; au["A2"].font = title_font
    au["A3"] = ("Every audience an inAudience / inSegment condition points at, resolved through the audiences API "
                "with the keyring credential. '(not found)' means the UUID is not in this sandbox's audience list "
                "(deleted, or the journey was copied from another sandbox). " + (aud_note or ""))
    au["A3"].font = note_font
    header(au, ["Audience", "Description", "Evaluation", "Lifecycle", "Audience UUID", "Journeys using it", "Journey names"], 5)
    a_rows = []
    for i in sorted(all_ids, key=lambda x: (aud_names.get(x, "~").lower(), x)):
        a = aud.get(i)
        a_rows.append([a["name"] if a else "(not found)", a["description"] if a else "", a["eval"] if a else "",
                       a["lifecycle"] if a else "", i, len(aud_use.get(i, ())), "; ".join(sorted(aud_use.get(i, ())))])
    for i, row in enumerate(a_rows, 6):
        for c, v in enumerate(row, 1):
            cell = au.cell(i, c, v)
            if c in (2, 7):
                cell.alignment = wrap
    widths(au, [50, 60, 11, 12, 38, 9, 70])

    # ---- Failed ---------------------------------------------------------------
    fs = wb.create_sheet("Failed")
    confidential(fs)
    fs["A2"] = "Journeys the list returned but /latest could not"; fs["A2"].font = title_font
    fs["A3"] = ("Typically abandoned duplications (..._CopyNNN) with no latest version behind them. They are absent "
                "from the other tabs. Empty when everything exported.")
    fs["A3"].font = note_font
    header(fs, ["Journey UID", "List name", "Error"], 5)
    for i, f in enumerate(failures, 6):
        fs.cell(i, 1, f.get("uid")); fs.cell(i, 2, f.get("list_name")); fs.cell(i, 3, f.get("error"))
    widths(fs, [38, 60, 90])

    for w, row, ncols in tables:
        if w.max_row > row:
            w.auto_filter.ref = f"A{row}:{get_column_letter(ncols)}{w.max_row}"
    out_path.parent.mkdir(exist_ok=True)
    try:
        wb.save(out_path)
    except PermissionError:
        alt = out_path.with_name(out_path.stem + f" ({datetime.now():%H%M%S})" + out_path.suffix)
        logger.warning(f"{out_path.name} is open; writing {alt.name} instead.")
        out_path = alt
        wb.save(out_path)
    return out_path


# ----------------------------------------------------------------------------
# Auth: user token from the browser; org id from env or the keyring credential
# ----------------------------------------------------------------------------
def resolve_token(arg_token) -> str:
    tok = arg_token or os.environ.get("AJO_BEARER_TOKEN") or ""
    if not tok and sys.stdin.isatty():
        print(TOKEN_HELP)
        tok = getpass.getpass("  Bearer token > ").strip()
    tok = tok.strip()
    if tok.lower().startswith("bearer "):
        tok = tok[7:].strip()
    if not tok:
        logger.error("No token. Set AJO_BEARER_TOKEN, pass --token=, or paste it at the prompt.")
        sys.exit(1)
    return tok


def resolve_service(arg_service) -> str | None:
    """The keyring credential to use for the org id and audience names."""
    services = aep_creds.list_services()
    if arg_service:
        return census.resolve_service(services, arg_service) or arg_service
    if services and sys.stdin.isatty():
        return census.menu(services)
    return services[0] if services else None


def print_banner():
    bar = ANSI["cyan"] + "=" * 72 + ANSI["reset"]
    print(bar)
    print(f"  {ANSI['bold']}ajo_journey_export v{SCRIPT_VERSION}{ANSI['reset']}   ({SCRIPT_DATE})")
    print(f"  {ANSI['dim']}Every AJO journey definition, every node -- via the UI's private "
          f"authoring API (unsupported; read-only) -- as JSON and a readable workbook.{ANSI['reset']}")
    print(bar)


def main():
    args = sys.argv[1:]
    opt, positional = {}, []
    for a in args:
        if a.startswith("--") and "=" in a:
            k, v = a[2:].split("=", 1)
            opt[k] = v
        elif a.startswith("--"):
            opt[a[2:]] = True
        elif a.startswith("-"):
            continue
        else:
            positional.append(a)
    print_banner()

    from_dir = Path(opt["from-dir"]) if opt.get("from-dir") else None
    service = None
    org = opt.get("org-id") or os.environ.get("AJO_ORG_ID")
    if not org or not from_dir:
        service = resolve_service(positional[0] if positional else None)
    if not org:
        if not service:
            logger.error("No org id: pass --org-id=, set AJO_ORG_ID, or store a credential "
                         "(credential_validator_v2.py store).")
            sys.exit(1)
        org = aep_creds.load_creds(service)["org_id"]

    sandbox = opt.get("sandbox") or os.environ.get("AJO_SANDBOX")
    if not sandbox and from_dir:
        m = re.search(r"ajo-journeys_([^_]+)_", from_dir.name)
        sandbox = m.group(1) if m else "prod"
    if from_dir:
        records = load_export_dir(from_dir)
        failures = []
        print(f"  loaded {len(records)} journey definition(s) from {from_dir}")
        out_dir = from_dir
    else:
        token = resolve_token(opt.get("token"))
        api_key = opt.get("api-key") or os.environ.get("AJO_API_KEY") or DEFAULT_API_KEY
        if not sandbox:
            # The sandbox-management API may refuse a UI token; pick_sandbox falls
            # back to 'prod' on any error, so this never blocks.
            sandbox = census.pick_sandbox(token, api_key, {"org_id": org}, default="prod")[0]
        if sandbox != sandbox.lower():
            logger.warning(f"Sandbox names are case-sensitive on this endpoint; "
                           f"using {sandbox.lower()!r} not {sandbox!r}.")
            sandbox = sandbox.lower()
        headers = {"Authorization": f"Bearer {token}", "x-api-key": api_key, "x-gw-ims-org-id": org,
                   "x-sandbox-name": sandbox, "Content-Type": "application/json"}
        logger.info(f"org {org}; sandbox {sandbox}; api key {api_key}")
        out_dir = Path(opt["out"]) if opt.get("out") else (
            OUTPUT_DIR / f"ajo-journeys_{sandbox}_{datetime.now():%Y-%m-%d}")
        records, failures = export(headers, sandbox, out_dir, opt.get("journey"))

    if opt.get("no-xlsx") or not records:
        return
    xlsx = OUTPUT_DIR / f"AJO Journeys - {sandbox} - {datetime.now():%Y-%m-%d}.xlsx"
    logger.info("building the readable workbook (resolving audience names via the keyring credential)...")
    path = build_workbook(records, failures, sandbox, service, xlsx, out_dir)
    print(f"    Workbook:         {path}")


if __name__ == "__main__":
    main()
