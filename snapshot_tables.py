#!/usr/bin/env python3
"""
snapshot_tables.py
==================
Inventory every Profile / Segment SNAPSHOT EXPORT table across the org -- every
sandbox the credential can see -- with the merge policy each one belongs to and
when the system last evaluated / exported it.

For each sandbox it:

  1. Lists the merge policies (/config/mergePolicies) so a snapshot's
     mergePolicyId tag can be shown as a NAME (and flagged when it's the default).
  2. Pages Catalog /dataSets and keeps the snapshot exports -- the datasets whose
     tags.unifiedProfile carries a ups_snapshot_type:* entry. Reports the friendly
     name AND the Query Service table name (tags['adobe/pqs/table']), which is the
     one you SELECT ... FROM.
  3. Looks up each snapshot's latest successful Catalog batch: that batch's
     started/completed time is when the snapshot was last written.
  4. Reads the sandbox's system evaluation: the batch_segmentation schedule
     (/config/schedules -- the cron time it is MEANT to fire) and the most recent
     scheduler-triggered segment job (/segment/jobs -- when it ACTUALLY ran). The
     snapshot is exported off the back of that run, so the two are shown side by
     side.

Writes ./output/snapshot_tables_<service>_<stamp>.xlsx -- a workbook in the Data
Dictionary house style (Summary tab + Snapshot Tables tab, one row per snapshot
table) -- and the same rows as a .csv. A sandbox that can't be read is reported
as UNREADABLE, never as "no snapshots".

Read-only: it never creates, edits or deletes anything in AEP. Standard library
only, plus the repo's aep_creds for the keyring vault and openpyxl for the XLSX
(without it the CSV is still written).

Usage:
    python snapshot_tables.py                       # interactive cred menu, all sandboxes
    python snapshot_tables.py aep-prod              # pick the service by name
    python snapshot_tables.py aep-prod --sandbox=prod,dev   # only these sandboxes
"""

from __future__ import annotations

import csv
import json
import logging
import ssl
import sys
import urllib.error
import urllib.parse
import urllib.request
from datetime import datetime, timezone
from pathlib import Path

import aep_creds  # keyring-backed credential store (replaces creds/*.json)

# ----------------------------------------------------------------------------
# Constants
# ----------------------------------------------------------------------------
SCRIPT_NAME = "snapshot_tables"
SCRIPT_VERSION = "1.0.0"
SCRIPT_DATE = "2026-09-29"
SCRIPT_AUTHOR = "Barry Mann (barrymann.com)"

SCRIPT_DIR = Path(__file__).resolve().parent
OUTPUT_DIR = SCRIPT_DIR / "output"

IMS_URL = "https://ims-na1.adobelogin.com/ims/token"
PLATFORM = "https://platform.adobe.io"
SANDBOX_LIST_URL = f"{PLATFORM}/data/foundation/sandbox-management/sandboxes"
DATASETS_URL = f"{PLATFORM}/data/foundation/catalog/dataSets"
CATALOG_BATCHES_URL = f"{PLATFORM}/data/foundation/catalog/batches"
UPS_MERGE_POLICIES_URL = f"{PLATFORM}/data/core/ups/config/mergePolicies"
CONFIG_SCHEDULES_URL = f"{PLATFORM}/data/core/ups/config/schedules"
SEGMENT_JOBS_URL = f"{PLATFORM}/data/core/ups/segment/jobs"

# The CORRECT snapshots are the ones on the sandbox's DEFAULT merge policy
# (default == true). Snapshots on other merge policies (added for debugging
# etc.) and datasets with no merge policy are listed but are not highlighted.
_CORRECT_BG = "C6EFCE"                  # Excel's 'Good' green
_CORRECT_FG = "006100"

# Time to run = how long the snapshot's latest successful batch took to write
# (batch started -> completed). RAG bands, in minutes: under AMBER is GREEN,
# AMBER up to RED is AMBER, RED and over is RED.
RAG_AMBER_MIN = 60
RAG_RED_MIN = 180
# Excel's Good / Neutral / Bad cell styles: (fill, font).
RAG_COLOURS = {
    "GREEN": ("C6EFCE", "006100"),
    "AMBER": ("FFEB9C", "9C5700"),
    "RED": ("FFC7CE", "9C0006"),
}

CONFIDENTIAL = "STRICTLY CONFIDENTIAL"
_HEADER_BG = "1F4E78"                   # the Data Dictionary's header blue

DEFAULT_SANDBOX = "prod"
PAGE_LIMIT = 100
# How many recent segment jobs to look through for the last scheduler run. The
# daily run is one job a day; api/FAE jobs can sit in front of it.
RECENT_JOBS = 50

DEFAULT_SCOPES = (
    "openid,AdobeID,read_organizations,"
    "additional_info.projectedProductContext,session"
)

# ----------------------------------------------------------------------------
# ANSI / logging - matches credential_validator.py / batch_fetcher.py style
# ----------------------------------------------------------------------------
if sys.platform == "win32":
    try:
        import ctypes
        kernel32 = ctypes.windll.kernel32
        h = kernel32.GetStdHandle(-11)  # STD_OUTPUT_HANDLE
        mode = ctypes.c_ulong()
        if kernel32.GetConsoleMode(h, ctypes.byref(mode)):
            kernel32.SetConsoleMode(h, mode.value | 0x0004)  # VT processing
    except Exception:
        pass

ANSI = {
    "reset": "\033[0m", "bold": "\033[1m", "dim": "\033[2m",
    "red": "\033[31m", "green": "\033[32m", "yellow": "\033[33m",
    "blue": "\033[34m", "magenta": "\033[35m", "cyan": "\033[36m",
}
LEVEL_COLOR = {
    "DEBUG": ANSI["dim"], "INFO": ANSI["green"],
    "WARNING": ANSI["yellow"],
    "ERROR": ANSI["red"] + ANSI["bold"],
    "CRITICAL": ANSI["red"] + ANSI["bold"],
}


class ColoredFormatter(logging.Formatter):
    def format(self, record):
        color = LEVEL_COLOR.get(record.levelname, "")
        ts = self.formatTime(record, "%H:%M:%S")
        return (
            f"{ANSI['dim']}{ts}{ANSI['reset']} "
            f"{color}[{record.levelname:<7}]{ANSI['reset']} "
            f"{record.getMessage()}"
        )


_handler = logging.StreamHandler(sys.stdout)
_handler.setFormatter(ColoredFormatter())
logging.basicConfig(level=logging.INFO, handlers=[_handler])
logger = logging.getLogger("snapshot_tables")
SSL_CTX = ssl._create_unverified_context()


# ----------------------------------------------------------------------------
# HTTP / IMS / credential helpers  (shared house style)
# ----------------------------------------------------------------------------
def http(url, method="GET", headers=None, data=None, timeout=60):
    """Stdlib-only HTTP. Returns response bytes; raises HTTPError on 4xx/5xx."""
    req = urllib.request.Request(url, data=data, headers=headers or {}, method=method)
    with urllib.request.urlopen(req, context=SSL_CTX, timeout=timeout) as r:
        return r.read()


def flatten_err(text: str, limit: int = 200) -> str:
    return " ".join((text or "").split())[:limit]


def describe_error(e: Exception) -> str:
    if isinstance(e, urllib.error.HTTPError):
        return f"HTTP {e.code}: {flatten_err(e.read().decode(errors='replace'))}"
    return f"{type(e).__name__}: {e}"


def authenticate(conf):
    """client_credentials grant against Adobe IMS -> access token string."""
    payload = urllib.parse.urlencode({
        "grant_type": "client_credentials",
        "client_id": conf["client_id"],
        "client_secret": conf["client_secret"],
        "scope": conf.get("scopes") or DEFAULT_SCOPES,
    }).encode("utf-8")
    body = http(
        conf.get("oauth_url") or IMS_URL,
        method="POST",
        headers={"Content-Type": "application/x-www-form-urlencoded"},
        data=payload,
    )
    return json.loads(body)["access_token"]


def aep_headers(token, conf, sandbox=None):
    headers = {
        "Authorization": f"Bearer {token}",
        "x-api-key": conf.get("api_key") or conf["client_id"],
        "x-gw-ims-org-id": conf["org_id"],
        "Accept": "application/json",
    }
    if sandbox:
        headers["x-sandbox-name"] = sandbox
    return headers


def menu(services):
    """Prompt for ONE credential set. Takes the list of keyring service-name
    strings and returns the chosen name (or None)."""
    print()
    bar = ANSI["cyan"] + "=" * 70 + ANSI["reset"]
    print(bar)
    print(f"  {ANSI['bold']}Credential bank{ANSI['reset']}  "
          f"{ANSI['dim']}(OS keyring vault){ANSI['reset']}")
    print(ANSI["cyan"] + "-" * 70 + ANSI["reset"])
    for i, name in enumerate(services, 1):
        print(f"  {ANSI['bold']}{i:>2}{ANSI['reset']}  "
              f"{ANSI['yellow']}{name}{ANSI['reset']}")
    print(bar)
    raw = input(f"\nPick a credential set by number "
                f"({ANSI['cyan']}1{ANSI['reset']}), blank to quit: ").strip()
    if not raw:
        return None
    if raw.isdigit() and 1 <= int(raw) <= len(services):
        return services[int(raw) - 1]
    logger.warning(f"Invalid choice: {raw}")
    return None


# ----------------------------------------------------------------------------
# Time helpers
# ----------------------------------------------------------------------------
def to_dt(value) -> datetime | None:
    """Best-effort parse of the many timestamp shapes AEP uses into an aware
    UTC datetime: epoch ms (int/str), epoch seconds, or ISO-8601 strings.
    Returns None when it can't be parsed."""
    if value in (None, "", 0):
        return None
    if isinstance(value, (int, float)) or (isinstance(value, str) and value.isdigit()):
        n = float(value)
        if n > 1e12:        # milliseconds
            n /= 1000.0
        try:
            return datetime.fromtimestamp(n, tz=timezone.utc)
        except (OverflowError, OSError, ValueError):
            return None
    if isinstance(value, str):
        s = value.replace("Z", "+00:00")
        try:
            dt = datetime.fromisoformat(s)
            return dt if dt.tzinfo else dt.replace(tzinfo=timezone.utc)
        except ValueError:
            return None
    return None


def first_dt(obj: dict, keys) -> datetime | None:
    for key in keys:
        dt = to_dt(obj.get(key))
        if dt:
            return dt
    return None


def fmt_dt(dt: datetime | None) -> str:
    return dt.strftime("%Y-%m-%d %H:%M:%S") if dt else ""


def fmt_age(dt: datetime | None) -> str:
    """How long ago, e.g. '5h 12m' or '3d 04h'."""
    if not dt:
        return ""
    secs = int((datetime.now(timezone.utc) - dt).total_seconds())
    if secs < 0:
        return "0m"
    if secs < 3600:
        return f"{secs // 60}m"
    if secs < 86400:
        return f"{secs // 3600}h {(secs % 3600) // 60:02d}m"
    return f"{secs // 86400}d {(secs % 86400) // 3600:02d}h"


def fmt_dur(seconds) -> str:
    """Human-friendly duration, e.g. '4m 12s' or '3h 29m'."""
    if seconds is None:
        return ""
    seconds = int(seconds)
    if seconds < 60:
        return f"{seconds}s"
    if seconds < 3600:
        return f"{seconds // 60}m {seconds % 60:02d}s"
    return f"{seconds // 3600}h {(seconds % 3600) // 60:02d}m"


def run_rag(seconds) -> str:
    """GREEN / AMBER / RED for a snapshot's time to run; '' when unknown."""
    if seconds is None:
        return ""
    if seconds >= RAG_RED_MIN * 60:
        return "RED"
    if seconds >= RAG_AMBER_MIN * 60:
        return "AMBER"
    return "GREEN"


def cron_time_utc(cron: str) -> str:
    """Pull an HH:MM (UTC) out of a cron that pins a single daily time. Handles
    Quartz 6-7 field ('sec min hour dom mon dow [year]') and 5-field standard
    ('min hour dom mon dow'). Returns the raw cron when it's recurring/sub-daily."""
    if not isinstance(cron, str) or not cron.strip():
        return ""
    parts = cron.split()
    if len(parts) >= 6:
        minute, hour = parts[1], parts[2]
    elif len(parts) == 5:
        minute, hour = parts[0], parts[1]
    else:
        return cron
    if minute.isdigit() and hour.isdigit():
        return f"{int(hour):02d}:{int(minute):02d}"
    return cron


# ----------------------------------------------------------------------------
# AEP reads
# ----------------------------------------------------------------------------
def list_sandboxes(token, conf):
    """(ok, sandboxes-or-error). Each sandbox is the raw dict (name, type, state)."""
    try:
        data = json.loads(http(SANDBOX_LIST_URL, headers=aep_headers(token, conf)))
        return True, data.get("sandboxes") or []
    except Exception as e:
        return False, describe_error(e)


def is_production(sb: dict) -> bool:
    """Production sandbox? Falls back to the name when the type is unknown
    (a sandbox named on the CLI that the sandbox list didn't return)."""
    stype = (sb.get("type") or "").lower()
    if stype:
        return stype == "production"
    return "prod" in sb["name"].lower()


def section_of(production: bool) -> str:
    return "PRODUCTION" if production else "DEVELOPMENT / PPE"


def fetch_merge_policies(headers) -> dict[str, dict]:
    """{mergePolicyId: policy dict} for the sandbox. Cursor-paged at _page.next;
    results live under 'children' or 'mergePolicies'."""
    out: dict[str, dict] = {}
    start, pages = None, 0
    while pages < 50:
        pages += 1
        params = {"limit": PAGE_LIMIT}
        if start:
            params["start"] = start
        url = f"{UPS_MERGE_POLICIES_URL}?{urllib.parse.urlencode(params)}"
        data = json.loads(http(url, headers=headers)) or {}
        batch = []
        for key in ("children", "mergePolicies"):
            if isinstance(data.get(key), list):
                batch = data[key]
                break
        for p in batch:
            if isinstance(p, dict) and p.get("id"):
                out[str(p["id"])] = p
        nxt = (data.get("_page") or {}).get("next")
        if not batch or len(batch) < PAGE_LIMIT or not nxt:
            break
        start = nxt
    return out


def _ups_tags(ds) -> list[str]:
    """The dataset's tags.unifiedProfile entries as a flat list of strings.
    Handles the dict-of-lists / flat-list / single-value shapes AEP uses."""
    up = (ds.get("tags") or {}).get("unifiedProfile")
    if up is None:
        return []
    if isinstance(up, list):
        return [str(x) for x in up]
    if isinstance(up, dict):
        out = []
        for v in up.values():
            out.extend(v if isinstance(v, list) else [v])
        return [str(x) for x in out]
    return [str(up)]


def _pqs_table(ds: dict) -> str:
    """The dataset's Query Service table name (tags['adobe/pqs/table']) -- the
    normalized SYSTEM name you SELECT ... FROM. '' when not assigned."""
    v = (ds.get("tags") or {}).get("adobe/pqs/table")
    if isinstance(v, list):
        return str(v[0]) if v else ""
    return str(v) if v else ""


def fetch_snapshot_datasets(headers, policies: dict[str, dict]) -> list[dict]:
    """Every snapshot export dataset in the sandbox: the ones whose
    tags.unifiedProfile carries a ups_snapshot_type:* entry. The owning merge
    policy id is among those same tag strings."""
    out = []
    start = 0
    while True:
        url = (f"{DATASETS_URL}?limit={PAGE_LIMIT}&start={start}"
               f"&properties=name,schemaRef,tags,created,updated")
        data = json.loads(http(url, headers=headers))
        if not isinstance(data, dict) or not data:
            break
        for dsid, ds in data.items():
            if not isinstance(ds, dict):
                continue
            ups = _ups_tags(ds)
            stype = next((t.split(":", 1)[1] for t in ups
                          if t.startswith("ups_snapshot_type:")), None)
            if stype is None:
                continue
            mp_id = next((t.split(":", 1)[1] for t in ups
                          if t.startswith("mergePolicyId:")), "")
            if not mp_id:
                # Tag shape varies; fall back to any known policy id in the tags.
                mp_id = next((pid for pid in policies
                              if any(pid in t for t in ups)), "")
            out.append({
                "id": dsid,
                "name": ds.get("name") or dsid,
                "table": _pqs_table(ds),
                "snapshot_type": stype,
                "merge_policy_id": mp_id,
                "schema": (ds.get("schemaRef") or {}).get("id") or "",
                "created": to_dt(ds.get("created")),
            })
        if len(data) < PAGE_LIMIT:
            break
        start += PAGE_LIMIT
    return out


def _batch_record_count(b: dict):
    """Records in a batch from its catalog metrics. '' when unknown."""
    m = b.get("metrics") or {}
    for v in (m.get("outputRecordCount"), m.get("inputRecordCount"),
              b.get("recordCount"), m.get("recordCount")):
        if isinstance(v, (int, float)):
            return int(v)
    return ""


def fetch_latest_batch(headers, dsid: str) -> dict | None:
    """The dataset's most recent SUCCESSFUL batch (with its id folded in), or
    None when it has never had one."""
    url = (f"{CATALOG_BATCHES_URL}?dataSet={urllib.parse.quote(dsid, safe='')}"
           f"&status=success&orderBy=desc:created&limit=1")
    data = json.loads(http(url, headers=headers))
    if not isinstance(data, dict):
        return None
    for bid, b in data.items():
        if isinstance(b, dict):
            return {**b, "id": bid}
    return None


def fetch_system_evaluation(headers) -> dict:
    """The sandbox's system (scheduled) evaluation: what the batch_segmentation
    schedule is set to, and when the scheduler last actually ran. Each half is
    fetched independently; a failure is recorded in 'error' rather than raised,
    so an unreadable schedule never costs us the snapshot rows."""
    out = {"state": "", "cron": "", "time": "", "job_id": "", "job_status": "",
           "job_started": None, "job_ended": None, "error": ""}
    errors = []
    try:
        data = json.loads(http(CONFIG_SCHEDULES_URL, headers=headers)) or {}
        scheds = data if isinstance(data, list) else (
            data.get("children") or data.get("schedules") or [])
        seg = [s for s in scheds if s.get("type") == "batch_segmentation"]
        # Prefer the active one when several exist.
        seg.sort(key=lambda s: s.get("state") != "active")
        if seg:
            cron = (seg[0].get("schedule") or "").strip()
            out.update(state=seg[0].get("state") or "", cron=cron,
                       time=cron_time_utc(cron))
    except Exception as e:
        errors.append(f"schedules {describe_error(e)}")
    try:
        params = {"limit": RECENT_JOBS, "sort": "creationTime:desc"}
        url = f"{SEGMENT_JOBS_URL}?{urllib.parse.urlencode(params)}"
        data = json.loads(http(url, headers=headers)) or {}
        jobs = data.get("children") or data.get("segmentJobs") or []
        jobs = [j for j in jobs if j.get("source") == "scheduler"]
        jobs.sort(key=lambda j: to_dt(j.get("creationTime"))
                  or datetime.min.replace(tzinfo=timezone.utc), reverse=True)
        if jobs:
            j = jobs[0]
            out.update(
                job_id=j.get("id") or "", job_status=j.get("status") or "",
                job_started=first_dt(j, ("startTime", "startedAt", "createEpoch",
                                         "creationTime", "created")),
                job_ended=first_dt(j, ("completedTime", "completedAt", "endTime",
                                       "updateEpoch", "updateTime", "updated")))
    except Exception as e:
        errors.append(f"segment jobs {describe_error(e)}")
    out["error"] = "; ".join(errors)
    return out


# ----------------------------------------------------------------------------
# Collect
# ----------------------------------------------------------------------------
def collect_sandbox(token, conf, sb: dict) -> tuple[list[dict], str]:
    """(rows, error) for one sandbox. error is '' on success; when set, the
    sandbox could not be read and rows is empty."""
    name = sb["name"]
    headers = aep_headers(token, conf, name)
    try:
        policies = fetch_merge_policies(headers)
    except Exception as e:
        # Snapshots can still be listed without policy names.
        logger.warning(f"  {name}: merge policies unreadable ({describe_error(e)}); "
                       f"policy names will be blank.")
        policies = {}
    try:
        snaps = fetch_snapshot_datasets(headers, policies)
    except Exception as e:
        return [], describe_error(e)
    if not snaps:
        logger.info(f"  {name}: no snapshot export tables.")
        return [], ""

    sysev = fetch_system_evaluation(headers)
    if sysev["error"]:
        logger.warning(f"  {name}: system evaluation partly unreadable: {sysev['error']}")

    rows = []
    for s in snaps:
        batch, batch_err = None, ""
        try:
            batch = fetch_latest_batch(headers, s["id"])
        except Exception as e:
            batch_err = describe_error(e)
        b = batch or {}
        pol = policies.get(s["merge_policy_id"]) or {}
        is_default = pol.get("default") is True
        if is_default:
            verdict = "YES"
        elif not pol:
            verdict = "no (no merge policy)"
        else:
            verdict = "no (non-default merge policy)"
        b_start = first_dt(b, ("started", "created"))
        b_end = first_dt(b, ("completed", "updated"))
        run_secs = ((b_end - b_start).total_seconds()
                    if b_start and b_end and b_end >= b_start else None)
        rows.append({
            "production": is_production(sb),
            "run_seconds": run_secs,
            "run_rag": run_rag(run_secs),
            "correct": is_default,
            "correct_table": verdict,
            "sandbox": name,
            "sandbox_type": sb.get("type") or "",
            "snapshot_name": s["name"],
            "table_name": s["table"],
            "dataset_id": s["id"],
            "snapshot_type": s["snapshot_type"],
            "merge_policy_name": pol.get("name") or "",
            "merge_policy_id": s["merge_policy_id"],
            "merge_policy_default": ("yes" if pol.get("default") is True
                                     else "no" if pol else ""),
            "snapshot_batch_id": b.get("id") or "",
            "snapshot_started": b_start,
            "snapshot_completed": b_end,
            "snapshot_records": _batch_record_count(b),
            "snapshot_status": (batch_err and f"UNREADABLE: {batch_err}")
                               or ("ok" if batch else "no successful batch"),
            "sys": sysev,
        })
    # The correct (default merge policy) tables lead each sandbox.
    rows.sort(key=lambda r: not r["correct"])
    logger.info(f"  {name}: {len(rows)} snapshot table(s).")
    return rows, ""


# ----------------------------------------------------------------------------
# Output
# ----------------------------------------------------------------------------
def print_rows(rows: list[dict]) -> None:
    bar = ANSI["cyan"] + "=" * 150 + ANSI["reset"]
    print()
    print(bar)
    print(ANSI["bold"] +
          f"  {'SANDBOX':<16}{'TABLE':<44}{'MERGE POLICY':<34}"
          f"{'SNAPSHOT WRITTEN (UTC)':<24}{'TIME TO RUN':<20}SYSTEM EVAL (UTC)"
          + ANSI["reset"])
    print(ANSI["cyan"] + "-" * 150 + ANSI["reset"])
    section = None
    for r in rows:
        if r["production"] != section:
            section = r["production"]
            print(f"  {ANSI['bold']}{ANSI['blue']}{section_of(section)}{ANSI['reset']}")
        written = r["snapshot_completed"] or r["snapshot_started"]
        pol = r["merge_policy_name"] or r["merge_policy_id"] or "?"
        if r["merge_policy_default"] == "yes":
            pol += " *"
        sysev = r["sys"]
        ev = fmt_dt(sysev["job_started"]) or "?"
        if sysev["time"]:
            ev += f"  (sched {sysev['time']})"
        wcolor = ANSI["green"] if written else ANSI["red"]
        rcolor = {"GREEN": ANSI["green"], "AMBER": ANSI["yellow"],
                  "RED": ANSI["red"] + ANSI["bold"]}.get(r["run_rag"], ANSI["dim"])
        run = f"{fmt_dur(r['run_seconds'])} {r['run_rag']}".strip()
        mark =(f"{ANSI['green']}{ANSI['bold']}>{ANSI['reset']} " if r["correct"]
                else "  ")
        tcolor = ANSI["green"] + ANSI["bold"] if r["correct"] else ""
        print(f"{mark}{ANSI['yellow']}{r['sandbox'][:14]:<16}{ANSI['reset']}"
              f"{tcolor}{(r['table_name'] or r['snapshot_name'])[:42]:<44}{ANSI['reset']}"
              f"{ANSI['magenta']}{pol[:32]:<34}{ANSI['reset']}"
              f"{wcolor}{(fmt_dt(written) or r['snapshot_status'])[:22]:<24}{ANSI['reset']}"
              f"{rcolor}{run:<20}{ANSI['reset']}"
              f"{ev}")
    print(bar)
    print(f"  {ANSI['bold']}{len(rows)} snapshot table(s){ANSI['reset']}  "
          f"{ANSI['dim']}(> and * = on the DEFAULT merge policy, the correct "
          f"one; SYSTEM EVAL = last scheduler "
          f"segment job start){ANSI['reset']}")


def write_csv(rows: list[dict], unreadable: list[tuple[str, str]],
              label: str, stamp: str) -> Path:
    OUTPUT_DIR.mkdir(exist_ok=True)
    path = OUTPUT_DIR / f"snapshot_tables_{label}_{stamp}.csv"
    cols = ["sandbox", "sandbox_type", "correct_table", "snapshot_name",
            "table_name", "dataset_id",
            "snapshot_type", "merge_policy_name", "merge_policy_id",
            "merge_policy_default", "snapshot_status", "snapshot_batch_id",
            "snapshot_started_utc", "snapshot_completed_utc",
            "time_to_run", "time_to_run_seconds", "time_to_run_rag",
            "snapshot_records", "system_eval_schedule_state", "system_eval_schedule_cron",
            "system_eval_schedule_time_utc", "system_eval_last_job_id",
            "system_eval_last_job_status", "system_eval_last_started_utc",
            "system_eval_last_ended_utc", "system_eval_error"]
    with path.open("w", newline="", encoding="utf-8") as fh:
        w = csv.writer(fh)
        w.writerow(cols)
        for r in rows:
            s = r["sys"]
            w.writerow([
                r["sandbox"], r["sandbox_type"], r["correct_table"],
                r["snapshot_name"], r["table_name"], r["dataset_id"],
                r["snapshot_type"],
                r["merge_policy_name"], r["merge_policy_id"],
                r["merge_policy_default"], r["snapshot_status"],
                r["snapshot_batch_id"], fmt_dt(r["snapshot_started"]),
                fmt_dt(r["snapshot_completed"]), fmt_dur(r["run_seconds"]),
                "" if r["run_seconds"] is None else int(r["run_seconds"]),
                r["run_rag"], r["snapshot_records"],
                s["state"], s["cron"], s["time"], s["job_id"], s["job_status"],
                fmt_dt(s["job_started"]), fmt_dt(s["job_ended"]), s["error"],
            ])
        # An unreadable sandbox gets a row of its own so it can't be mistaken
        # for a sandbox with no snapshots.
        for name, err in unreadable:
            row = [""] * len(cols)
            row[0] = name
            row[cols.index("snapshot_status")] = f"SANDBOX UNREADABLE: {err}"
            w.writerow(row)
    return path


def fmt_lag(start: datetime | None, end: datetime | None) -> str:
    """Evaluation start -> snapshot written, e.g. '5h 25m'. '' when the snapshot
    predates the evaluation (it belongs to an earlier run)."""
    if not start or not end or end < start:
        return ""
    secs = int((end - start).total_seconds())
    return f"{secs // 3600}h {(secs % 3600) // 60:02d}m"


def write_xlsx(rows: list[dict], unreadable: list[tuple[str, str]],
               sandboxes: list[dict], conf: dict, label: str, stamp: str):
    """The workbook, in the Data Dictionary house style: confidential banner,
    title + italic note, header-blue filter row, frozen panes, fixed column
    widths and uniform two-line rows. Returns the path, or None when openpyxl
    isn't installed (the CSV is still written)."""
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Font, PatternFill, Alignment
        from openpyxl.utils import get_column_letter
    except ImportError:
        logger.warning("openpyxl not installed -- skipping XLSX (CSV still written).")
        return None

    head_font = Font(bold=True, color="FFFFFF")
    head_fill = PatternFill("solid", fgColor=_HEADER_BG)
    title_font = Font(bold=True, size=14)
    conf_font = Font(bold=True, size=11, color="C00000")
    note_font = Font(italic=True, color="666666")
    bad_font = Font(bold=True, color="C00000")
    good_fill = PatternFill("solid", fgColor=_CORRECT_BG)
    section_fill = PatternFill("solid", fgColor="DCE6F1")
    wrap = Alignment(wrap_text=True, vertical="top")
    top = Alignment(vertical="top")
    ROW_H = 30          # two lines; uniform rows read better than Excel's per-row autofit
    tables = []
    datestr = datetime.now().strftime("%Y-%m-%d")

    def confidential(ws):
        ws["A1"] = CONFIDENTIAL
        ws["A1"].font = conf_font
        ws.oddHeader.center.text = f'&"-,Bold"&12&KC00000{CONFIDENTIAL}'
        ws.evenHeader.center.text = ws.oddHeader.center.text

    def note(ws, text, span, height):
        ws["A3"] = text
        ws["A3"].font = note_font
        ws["A3"].alignment = wrap
        ws.merge_cells(f"A3:{get_column_letter(span)}3")
        ws.row_dimensions[3].height = height

    def header(ws, cols, row):
        for c, nm in enumerate(cols, 1):
            cell = ws.cell(row, c, nm)
            cell.font, cell.fill = head_font, head_fill
            cell.alignment = Alignment(vertical="center", wrap_text=True)
        ws.row_dimensions[row].height = 28
        tables.append((ws, row, len(cols)))

    def rag(cell, grade):
        fill, fg = RAG_COLOURS[grade]
        cell.fill = PatternFill("solid", fgColor=fill)
        cell.font = Font(bold=True, color=fg)

    def band(ws, row, text, ncols):
        """Section divider row, in the Data Dictionary's section blue."""
        for c in range(1, ncols + 1):
            ws.cell(row, c).fill = section_fill
        ws.cell(row, 1, text).font = Font(bold=True, color=_HEADER_BG)
        ws.cell(row, 1).alignment = Alignment(vertical="center")
        ws.row_dimensions[row].height = 22

    def widths(ws, ws_widths):
        for i, w in enumerate(ws_widths, 1):
            ws.column_dimensions[get_column_letter(i)].width = w

    wb = Workbook()
    # House rule: the file's author is Barry, not the library ('openpyxl').
    wb.properties.creator = wb.properties.lastModifiedBy = SCRIPT_AUTHOR
    # ---- Summary --------------------------------------------------------------
    ws = wb.active
    ws.title = "Summary"
    ws.sheet_properties.tabColor = _HEADER_BG
    confidential(ws)
    ws["A2"] = f"AEP Snapshot Tables -- all sandboxes  ({datestr})"
    ws["A2"].font = title_font
    note(ws, "Every Profile / Segment snapshot export table in every sandbox the "
             "credential can see, with the merge policy it belongs to and when the "
             "system last evaluated it. Production sandboxes are at the top; "
             "development / PPE sandboxes are in their own section at the bottom. "
             "'Scheduled evaluation' is the time the sandbox's batch segmentation "
             "schedule is set to fire; 'Last system evaluation' is when the "
             "scheduler actually last ran. Snapshots are written off the back of "
             "that run. 'Time to run' is how long the snapshot took to write: "
             f"GREEN under {RAG_AMBER_MIN} min, AMBER {RAG_AMBER_MIN}-{RAG_RED_MIN} "
             f"min, RED {RAG_RED_MIN} min or more. All times are UTC. Read-only.",
         8, 94)
    r0 = 5
    facts = [
        ("Org", conf["org_id"]),
        ("Credential", label),
        ("Sandboxes read", len(sandboxes) - len(unreadable)),
        ("Sandboxes unreadable", len(unreadable)),
        ("Snapshot tables", len(rows)),
        ("Merge policies in use",
         len({(r["sandbox"], r["merge_policy_id"]) for r in rows if r["merge_policy_id"]})),
        ("Generated (UTC)", datetime.now(timezone.utc).strftime("%Y-%m-%d %H:%M:%S")),
        ("Generated by", f"{SCRIPT_NAME}.py v{SCRIPT_VERSION} ({SCRIPT_DATE})"),
    ]
    for i, (k, v) in enumerate(facts):
        ws.cell(r0 + i, 1, k).font = Font(bold=True)
        ws.cell(r0 + i, 2, v).alignment = Alignment(horizontal="left")
    r1 = r0 + len(facts) + 1
    header(ws, ["Sandbox", "Correct table (default merge policy)",
                "Snapshot written (UTC)", "Time to run", "Time to run RAG",
                "Snapshot tables", "Schedule state",
                "Scheduled evaluation (UTC)", "Last system evaluation started (UTC)",
                "Last system evaluation ended (UTC)", "Last evaluation status",
                "Type"], r1)
    by_sb: dict[str, list[dict]] = {}
    for r in rows:
        by_sb.setdefault(r["sandbox"], []).append(r)
    bad = dict(unreadable)
    rr = r1
    section = None
    for sb in sandboxes:
        if is_production(sb) != section:
            section = is_production(sb)
            rr += 1
            band(ws, rr, section_of(section), 12)
        rr += 1
        name = sb["name"]
        ws.row_dimensions[rr].height = ROW_H
        got = by_sb.get(name) or []
        s = got[0]["sys"] if got else {}
        right = [r for r in got if r["correct"]]
        # The cell lists every correct table; colour it by the slowest of them.
        worst = next((g for g in ("RED", "AMBER", "GREEN")
                      if any(r["run_rag"] == g for r in right)), "")
        if name in bad:
            vals = [name, f"UNREADABLE: {bad[name]}"]
        else:
            vals = [name,
                    "\n".join(r["table_name"] or r["snapshot_name"] for r in right)
                    or "NONE FOUND",
                    "\n".join(fmt_dt(r["snapshot_completed"] or r["snapshot_started"])
                              for r in right),
                    "\n".join(fmt_dur(r["run_seconds"]) for r in right),
                    "\n".join(r["run_rag"] for r in right),
                    len(got), s.get("state") or "",
                    s.get("time") or "", fmt_dt(s.get("job_started")),
                    fmt_dt(s.get("job_ended")), s.get("job_status") or "",
                    sb.get("type") or ""]
        for c, v in enumerate(vals, 1):
            cell = ws.cell(rr, c, v)
            cell.alignment = wrap if c in (2, 3, 4, 5) else top
            if name in bad or (c == 2 and not right):
                cell.font = bad_font
            elif c == 2:
                cell.font = Font(bold=True, color=_CORRECT_FG)
                cell.fill = good_fill
            elif c in (4, 5) and worst:
                rag(cell, worst)
    widths(ws, [24, 66, 22, 14, 14, 16, 16, 22, 26, 26, 20, 16])

    # ---- Snapshot Tables ------------------------------------------------------
    st = wb.create_sheet("Snapshot Tables")
    st.sheet_properties.tabColor = "7030A0"      # the Data Dictionary's Profile purple
    confidential(st)
    st["A2"] = "Snapshot tables -- one row per snapshot export table"
    st["A2"].font = title_font
    note(st, "'Table name' is the Query Service table you SELECT ... FROM. "
             "'Snapshot written' is the completion time of the table's latest "
             "successful batch. 'Evaluation -> snapshot' is the gap between the "
             "last system evaluation starting and this snapshot being written "
             "(blank when the snapshot predates that evaluation). A blank merge "
             "policy means the dataset carries no merge policy tag. The GREEN rows "
             "are the CORRECT tables for each sandbox: the ones on the DEFAULT "
             "merge policy. The rest are on other merge policies (added for "
             "debugging etc.) or carry no merge policy. 'Time to run' is how long "
             "that latest batch took to write (started -> completed): GREEN under "
             f"{RAG_AMBER_MIN} min, AMBER {RAG_AMBER_MIN}-{RAG_RED_MIN} min, RED "
             f"{RAG_RED_MIN} min or more.", 8, 94)
    cols = ["Sandbox", "Correct table", "Table name", "Merge policy",
            "Default policy", "Snapshot written (UTC)", "Time to run",
            "Time to run RAG",
            "Snapshot age", "Records", "Snapshot status",
            "Scheduled evaluation (UTC)", "Last system evaluation started (UTC)",
            "Last system evaluation ended (UTC)", "Last evaluation status",
            "Evaluation -> snapshot", "Snapshot type", "Snapshot name",
            "Sandbox type", "Schedule state", "Schedule cron",
            "Merge policy id", "Dataset id", "Snapshot batch id",
            "Last evaluation job id"]
    header(st, cols, 5)
    st.freeze_panes = "D6"
    i = 5
    section = None
    for r in rows:
        if r["production"] != section:
            section = r["production"]
            i += 1
            band(st, i, section_of(section), len(cols))
        i += 1
        s = r["sys"]
        written = r["snapshot_completed"] or r["snapshot_started"]
        vals = [r["sandbox"], r["correct_table"], r["table_name"],
                r["merge_policy_name"], r["merge_policy_default"],
                fmt_dt(written), fmt_dur(r["run_seconds"]), r["run_rag"],
                fmt_age(written), r["snapshot_records"],
                r["snapshot_status"], s["time"], fmt_dt(s["job_started"]),
                fmt_dt(s["job_ended"]), s["job_status"],
                fmt_lag(s["job_started"], written), r["snapshot_type"],
                r["snapshot_name"], r["sandbox_type"], s["state"], s["cron"],
                r["merge_policy_id"], r["dataset_id"], r["snapshot_batch_id"],
                s["job_id"]]
        st.row_dimensions[i].height = ROW_H
        for c, v in enumerate(vals, 1):
            cell = st.cell(i, c, v)
            cell.alignment = wrap if c in (2, 3, 4, 11, 18) else top
            if c == 10 and isinstance(v, int):
                cell.number_format = "#,##0"
            if r["correct"]:
                cell.fill = good_fill
                if c <= 4:
                    cell.font = Font(bold=True, color=_CORRECT_FG)
            elif c <= 4:
                cell.font = Font(color="808080")
            if c in (7, 8) and r["run_rag"]:
                rag(cell, r["run_rag"])
        if r["snapshot_status"] != "ok":
            st.cell(i, 11).font = bad_font
    widths(st, [16, 18, 46, 30, 10, 20, 13, 12, 12, 14, 18, 14, 20, 20, 14, 14,
                18, 40, 14, 12, 18, 38, 28, 30, 38])

    for w, row, ncols in tables:
        if w.max_row > row:
            w.auto_filter.ref = f"A{row}:{get_column_letter(ncols)}{w.max_row}"
    OUTPUT_DIR.mkdir(exist_ok=True)
    path = OUTPUT_DIR / f"snapshot_tables_{label}_{stamp}.xlsx"
    try:
        wb.save(path)
    except PermissionError:
        alt = path.with_name(path.stem + f" ({datetime.now():%H%M%S})" + path.suffix)
        logger.warning(f"{path.name} is open; writing {alt.name} instead.")
        path = alt
        wb.save(path)
    return path


# ----------------------------------------------------------------------------
# Main
# ----------------------------------------------------------------------------
def parse_args(argv):
    opts = {"name": None, "sandboxes": None}
    for a in argv:
        if a.startswith("--sandbox=") or a.startswith("--sandboxes="):
            val = a.split("=", 1)[1].strip()
            if val and val.lower() != "all":
                opts["sandboxes"] = [s.strip() for s in val.split(",") if s.strip()]
        elif a in ("-h", "--help"):
            print(__doc__)
            sys.exit(0)
        elif a.startswith("-"):
            logger.warning(f"Ignoring unknown option {a}")
        else:
            opts["name"] = a  # keyring service name
    return opts


def banner(conf):
    bar = ANSI["cyan"] + "=" * 72 + ANSI["reset"]
    print(bar)
    print(f"  {ANSI['bold']}{SCRIPT_NAME} v{SCRIPT_VERSION}{ANSI['reset']}   ({SCRIPT_DATE})")
    print(f"  by {SCRIPT_AUTHOR}")
    print(f"  {ANSI['dim']}Snapshot export tables, their merge policies and system "
          f"evaluation time, across all sandboxes (read-only){ANSI['reset']}")
    print(f"  {ANSI['bold']}Org:{ANSI['reset']}      {ANSI['magenta']}{conf['org_id']}{ANSI['reset']}")
    print(bar)


def main():
    opts = parse_args(sys.argv[1:])

    print(aep_creds.source_banner())
    services = aep_creds.list_services()
    if not services:
        logger.error("No credentials found in the keyring vault or the creds/ "
                     "folder. Add one with credential_validator_v2.py store "
                     "(or migrate), or drop a <service>.json in creds/.")
        return

    if opts["name"]:
        try:
            chosen = aep_creds.pick_service(opts["name"])
        except aep_creds.CredsError as e:
            logger.error(str(e))
            return
    else:
        chosen = menu(services)
    if not chosen:
        logger.info("Nothing chosen. Exiting.")
        return

    try:
        conf = aep_creds.load_creds(chosen)
    except aep_creds.CredsError as e:
        logger.error(f"Failed to load credentials for {chosen!r}: {e}")
        return
    banner(conf)

    try:
        token = authenticate(conf)
    except Exception as e:
        logger.error(f"IMS auth FAILED: {describe_error(e)}")
        return
    logger.info("IMS authenticated.")

    # Which sandboxes: everything the credential can see, unless narrowed.
    ok, found = list_sandboxes(token, conf)
    if ok:
        sandboxes = [s for s in found if isinstance(s, dict) and s.get("name")
                     and (s.get("state") or "active") == "active"]
        skipped = len(found) - len(sandboxes)
        logger.info(f"{len(sandboxes)} active sandbox(es)"
                    + (f" ({skipped} inactive skipped)" if skipped else "") + ".")
    else:
        logger.warning(f"Sandbox list unavailable ({found}).")
        sandboxes = []
    if opts["sandboxes"]:
        known = {s["name"]: s for s in sandboxes}
        sandboxes = [known.get(n) or {"name": n} for n in opts["sandboxes"]]
    elif not sandboxes:
        fallback = conf.get("sandbox") or DEFAULT_SANDBOX
        logger.warning(f"Falling back to the single sandbox '{fallback}' -- this "
                       f"is NOT the whole org. Pass --sandbox=a,b,c to name them.")
        sandboxes = [{"name": fallback}]

    # Production sandboxes first ('prod' leading), development / PPE at the bottom.
    sandboxes.sort(key=lambda s: (not is_production(s), s["name"] != "prod",
                                  s["name"]))

    rows, unreadable = [], []
    for sb in sandboxes:
        logger.info(f"Sandbox '{sb['name']}'...")
        got, err = collect_sandbox(token, conf, sb)
        if err:
            logger.error(f"  {sb['name']}: UNREADABLE -- {err}")
            unreadable.append((sb["name"], err))
        rows.extend(got)

    if rows:
        print_rows(rows)
    else:
        logger.warning("No snapshot export tables found.")
    if unreadable:
        print(f"  {ANSI['red']}{ANSI['bold']}{len(unreadable)} sandbox(es) UNREADABLE "
              f"(not the same as having no snapshots):{ANSI['reset']} "
              + ", ".join(n for n, _ in unreadable))

    if rows or unreadable:
        stamp = datetime.now(timezone.utc).strftime("%Y%m%d_%H%M%S")
        path = write_csv(rows, unreadable, chosen, stamp)
        logger.info(f"Wrote {len(rows)} row(s) to {path}")
        xlsx = write_xlsx(rows, unreadable, sandboxes, conf, chosen, stamp)
        if xlsx:
            logger.info(f"Wrote workbook {xlsx}")

    print()
    logger.info("Done.")


if __name__ == "__main__":
    main()
