#!/usr/bin/env python3
"""
audience_renamer.py  (AEP Swiss Army Knife)  -- BETA
====================================================
The babelfish_query_renamer idea, for audiences: give every rule-based audience
in a sandbox a friendly name built from a naming convention.

BETA -- not yet proven end to end:
  * the rename and restore calls (--apply / --restore) have never been sent to
    AEP; the first real run should be --apply --limit=1, then --restore;
  * the Claude step has not been run from this tool. It needs the 'anthropic'
    package and an API key, and a key that is not scoped to a workspace also
    needs an anthropic-workspace-id header, which is not sent yet. The dev
    descriptions on file were written in a Claude Code session instead.

It does NOT fetch or render anything itself. It follows on from a Data
Dictionary run and reads what that run left behind:

  * the full-PQL sidecar (output/data_dictionary/pql/audiences_pql_<sandbox>_
    <date>.jsonl) -- every audience's id, name and rule;
  * the dictionary workbook -- the Audiences tab (existing description,
    evaluation type) and the Field Index (the friendly name of every field).

So the order is:

    python data_dictionary_v3.py <service> --sandbox=dev     # fresh dictionary
    python audience_renamer.py                               # dry run: the plan
    python audience_renamer.py <service> --apply             # rename in AEP
    python audience_renamer.py <service> --restore           # put the names back

Each audience gets three HEADER TERMS, and the naming convention (NAME_FORMAT)
strings them together:

    prefix A      placeholder until the real convention arrives (--prefix-a=)
    prefix B      placeholder until the real convention arrives (--prefix-b=)
    description   a short, friendly description written by Claude

When the convention is agreed, change NAME_FORMAT / the header terms here and
re-run: everything is renamed to match.

THE DESCRIPTION is written by Claude from four things: the audience's CURRENT
NAME (the use-case codes, campaign ids and labels people already put there are
kept), its existing description, its rule, and the dictionary's friendly names
for the fields and audiences the rule mentions -- so a name reads "Clubcard
members who bought cat products twice in 3 months", never a field path or an
id. Descriptions are kept in output/audience_renamer_<sandbox>_descriptions.json
and only re-written when an audience's rule changes, so a re-run costs nothing
and gives the same names.

The Anthropic API key is read from the ANTHROPIC_API_KEY environment variable,
or from the keyring. Store it once (it is never written to disk):

    python -c "import keyring, getpass; keyring.set_password('anthropic', 'api_key', getpass.getpass('Anthropic API key: '))"

Without a key the tool still runs: it uses the descriptions already in the
file and skips any audience that has none.

Audiences with no rule (uploads, Data Distiller, compositions) are left alone
-- there is nothing to describe -- and the workbook says so.

Putting it back: before an audience is renamed, its name is saved in
output/audience_renamer_<sandbox>_previous_names.json. --restore renames every
audience in that file back to the saved name and clears it as it goes. The
file keeps the ORIGINAL name, so renaming twice still restores to the start.

Safety:
  * Dry run by default. Nothing is written to AEP without --apply / --restore.
  * Development sandboxes only (dev, the PPEs). A production sandbox is refused.
  * --apply re-reads each audience first and skips it if its name has changed
    since the dictionary run, or if it is a system audience.

Writes ONE workbook, output/audience_renamer_<sandbox>.xlsx, overwritten each
run.

Usage:
    python audience_renamer.py                          # dry run on dev
    python audience_renamer.py --sandbox=uk-ppe         # dry run on a PPE
    python audience_renamer.py --prefix-a=UK --prefix-b=CRM
    python audience_renamer.py aep-prod --apply         # rename in dev
    python audience_renamer.py aep-prod --apply --limit=1   # just the first one
    python audience_renamer.py aep-prod --restore       # undo: old names back
"""

from __future__ import annotations

import hashlib
import json
import logging
import os
import re
import sys
import urllib.error
from datetime import datetime, timezone
from pathlib import Path

import aep_creds  # keyring-backed credential store (replaces creds/*.json)
# The dictionary owns the audience read, the PQL rendering and the sidecar;
# this tool only consumes its output and borrows its HTTP / IMS helpers.
import data_dictionary_v3 as dd

SCRIPT_NAME = "audience_renamer"
SCRIPT_VERSION = "0.2.0-beta"
SCRIPT_DATE = "2026-10-01"
SCRIPT_AUTHOR = "Barry Mann (barrymann.com)"

SCRIPT_DIR = Path(__file__).resolve().parent
OUTPUT_DIR = SCRIPT_DIR / "output"

DEFAULT_SANDBOX = "dev"

# ----------------------------------------------------------------------------
# THE NAMING CONVENTION -- starter version. Swap this for the real one when it
# arrives; the tokens are the header terms built in build_plan().
# ----------------------------------------------------------------------------
NAME_FORMAT = "{prefix_a}: {prefix_b}: {description}"
PREFIX_A = "prefix A"
PREFIX_B = "prefix B"
MAX_DESCRIPTION = 100           # characters of description that go into a name

# ----------------------------------------------------------------------------
# The AI description
# ----------------------------------------------------------------------------
AI_MODEL = "claude-opus-5-5"
AI_EFFORT = "medium"
AI_KEYRING_SERVICE = "anthropic"        # keyring: service 'anthropic', key 'api_key'
AI_SYSTEM = f"""\
You write the description part of an Adobe Experience Platform audience name.
You are given one audience: its current name, its existing description (often \
blank or a placeholder), how it is evaluated, its segmentation rule, and the \
friendly names of the fields and audiences the rule mentions.

Write ONE short description that a marketer can read at a glance.

- Keep what is good in the current name. Use-case and test codes (UC 5.1, \
TS37, TO24, E4T11), campaign ids, journey and group labels (Control Group, \
Holdout Group), owner tags ([IS], BAM, TDI, PerfTest), warnings (DONT USE), \
"copy", and a trailing Batch / Streaming label were put there on purpose. \
Carry them over, tidied.
- Then say what the rule actually does, in plain English: who is in the \
audience, with the real numbers, time windows and values from the rule. The \
rule is the truth; where the name or the existing description disagrees with \
it, follow the rule.
- Use the friendly field names you are given, never a field path. Never \
include a UUID, key or technical id: name the audience it refers to if you \
are told it, otherwise say "a specific store", "the journey email" and so on.
- Common hygiene clauses (marketing consent, Clubcard member, not a new \
registrant, excluding PerfTest profiles) can be summarised in a word or two \
or dropped when space is short.
- Aim for under 80 characters; never exceed {MAX_DESCRIPTION}.

Reply with the description only: one line, no quotes, no prefixes, no \
explanation."""

logger = logging.getLogger(SCRIPT_NAME)   # dd has already set up the handler
ANSI = dd.ANSI


# ----------------------------------------------------------------------------
# What the Data Dictionary run left behind
# ----------------------------------------------------------------------------
def latest_sidecar(sandbox: str) -> Path | None:
    """Newest audiences_pql_<sandbox>_<yyyymmdd>.jsonl the dictionary wrote."""
    found = sorted(dd.PQL_DIR.glob(f"audiences_pql_{sandbox}_*.jsonl"))
    return found[-1] if found else None


def sidecar_date(path: Path) -> str:
    """The run date in the sidecar's name, as YYYY-MM-DD ('' if it has none)."""
    m = re.search(r"_(\d{4})(\d{2})(\d{2})\.jsonl$", path.name)
    return "-".join(m.groups()) if m else ""


def read_sidecar(path: Path) -> list[dict]:
    with path.open(encoding="utf-8") as fh:
        return [json.loads(line) for line in fh if line.strip()]


def read_dictionary(sandbox: str) -> tuple[dict, dict]:
    """(friendly, audiences) from the sandbox's newest dictionary workbook.

    friendly   {field path: friendly name}, from the Field Index tab
    audiences  {audience id: {description, evaluation, lifecycle}}, from the
               Audiences tab
    Both empty when there is no workbook -- the AI then works from the sidecar
    alone, which still names every audience and carries every rule."""
    books = sorted(dd.OUTPUT_DIR.glob(f"Data Dictionary - {sandbox} - *.xlsx"))
    if not books:
        logger.warning(f"No dictionary workbook for '{sandbox}' in "
                       f"{dd.OUTPUT_DIR}; no friendly field names to offer.")
        return {}, {}
    from openpyxl import load_workbook
    wb = load_workbook(books[-1], read_only=True)

    def table(sheet: str):
        """Rows of a house-style tab as dicts keyed by its row-5 headers."""
        if sheet not in wb.sheetnames:
            return
        rows = wb[sheet].iter_rows(min_row=5, values_only=True)
        headers = next(rows, None) or ()
        for row in rows:
            yield dict(zip(headers, row))

    friendly, audiences = {}, {}
    for r in table("Field Index"):
        path, name = r.get("Field (dot notation)"), r.get("Friendly Name")
        if path and name:
            friendly.setdefault(path.replace("[]", ""), name)
    for r in table("Audiences"):
        if r.get("Audience id"):
            audiences[r["Audience id"]] = {
                "description": r.get("Description") or "",
                "evaluation": r.get("Evaluation") or "",
                "lifecycle": r.get("Lifecycle") or ""}
    logger.info(f"Read {books[-1].name}: {len(friendly)} friendly field names.")
    return friendly, audiences


# ----------------------------------------------------------------------------
# The AI description
# ----------------------------------------------------------------------------
def rule_sha(entry: dict) -> str:
    """Fingerprint of an audience's rule: a description is reused only while
    the rule it was written from is unchanged."""
    return hashlib.sha256((entry.get("raw_pql") or "").encode("utf-8")).hexdigest()[:12]


def descriptions_path(sandbox: str) -> Path:
    return OUTPUT_DIR / f"audience_renamer_{sandbox}_descriptions.json"


def load_descriptions(sandbox: str) -> dict:
    """{audience id: {description, rule_sha, named_from, by}}."""
    path = descriptions_path(sandbox)
    return json.loads(path.read_text(encoding="utf-8")) if path.exists() else {}


def save_descriptions(sandbox: str, descriptions: dict) -> None:
    OUTPUT_DIR.mkdir(exist_ok=True)
    descriptions_path(sandbox).write_text(
        json.dumps(descriptions, indent=2, ensure_ascii=False), encoding="utf-8")


def rule_lookups(entry: dict, friendly: dict, names: dict) -> tuple[dict, dict]:
    """(fields, audiences) the rule mentions that the dictionary can name:
    {field path: friendly name} and {audience id: audience name}."""
    try:
        tree = json.loads(entry.get("raw_pql") or "")
    except ValueError:                  # pql-text: ids appear as plain text
        raw = entry.get("raw_pql") or ""
        return {}, {i: n for i, n in names.items() if i in raw}
    fields, audiences = {}, {}

    def walk(node, inside_path=False):
        if isinstance(node, list):
            for n in node:
                walk(n)
        elif isinstance(node, dict):
            if node.get("nodeType") == "fieldLookup" and not inside_path:
                parts, cur = [], node
                while isinstance(cur, dict) and cur.get("nodeType") == "fieldLookup":
                    parts.append(cur.get("fieldName") or "")
                    cur = cur.get("object")
                path = ".".join(reversed(parts))
                if path in friendly:
                    fields[path] = friendly[path]
            if node.get("nodeType") == "literal" and node.get("value") in names:
                audiences[node["value"]] = names[node["value"]]
            for key, v in node.items():
                walk(v, inside_path=(key == "object"
                                     and node.get("nodeType") == "fieldLookup"))

    walk(tree)
    return fields, audiences


def ai_prompt(entry: dict, info: dict, friendly: dict, names: dict) -> str:
    fields, audiences = rule_lookups(entry, friendly, names)
    lines = [f"Current name: {entry.get('audience_name') or ''}",
             f"Existing description: {' '.join(info.get('description', '').split()) or '(none)'}",
             f"Evaluation: {info.get('evaluation') or '?'}    "
             f"Lifecycle: {info.get('lifecycle') or '?'}",
             "",
             "Rule as the Data Dictionary renders it (<...> marks a part it "
             "could not render; the full rule below has everything):",
             entry.get("readable_pql") or "",
             "",
             "Full rule (Adobe pql/json syntax tree, or PQL text):",
             entry.get("raw_pql") or ""]
    if fields:
        lines += ["", "Friendly names of the fields in the rule:"]
        lines += [f"{path} = {name}" for path, name in sorted(fields.items())]
    if audiences:
        lines += ["", "Audiences the rule refers to by id:"]
        lines += [f"{aid} = {name}" for aid, name in sorted(audiences.items())]
    return "\n".join(lines)


def ai_client():
    """An Anthropic client, or None (with the reason logged) when there is no
    SDK or no key. The key comes from ANTHROPIC_API_KEY, else the keyring."""
    key = None
    if not (os.environ.get("ANTHROPIC_API_KEY") or os.environ.get("ANTHROPIC_AUTH_TOKEN")):
        key = aep_creds.kr_get(AI_KEYRING_SERVICE, "api_key")
        if not key:
            logger.warning("No Anthropic API key (ANTHROPIC_API_KEY, or keyring "
                           f"service '{AI_KEYRING_SERVICE}' key 'api_key') -- "
                           "no new descriptions will be written. See --help.")
            return None
    try:
        import anthropic
    except ImportError:
        logger.warning("The 'anthropic' package is not installed -- no new "
                       "descriptions will be written (pip install anthropic).")
        return None
    return anthropic.Anthropic(api_key=key) if key else anthropic.Anthropic()


def tidy_description(text: str) -> str:
    """One line, no wrapping quotes, cut to MAX_DESCRIPTION."""
    text = " ".join((text or "").split()).strip("\"'")
    if len(text) > MAX_DESCRIPTION:
        text = text[:MAX_DESCRIPTION - 3].rstrip() + "..."
    return text


def ai_describe(client, prompt: str) -> str:
    """One description from Claude. Raises on an API failure; returns '' when
    the model declined or said nothing."""
    response = client.beta.messages.create(
        model=AI_MODEL,
        max_tokens=16000,
        # A request the model's safety classifiers decline is re-run on
        # Anthropic's recommended fallback model instead of coming back empty.
        betas=["server-side-fallback-2026-07-01"],
        fallbacks="default",
        output_config={"effort": AI_EFFORT},
        system=AI_SYSTEM,
        messages=[{"role": "user", "content": prompt}],
    )
    if response.stop_reason == "refusal":
        return ""
    return tidy_description(next((b.text for b in response.content
                                  if b.type == "text"), ""))


def describe_audiences(entries: list[dict], sandbox: str) -> dict:
    """The descriptions file brought up to date: Claude writes one for every
    rule-based audience that has none, or whose rule has changed since. Saved
    after each, so a run that stops half way keeps what it paid for."""
    descriptions = load_descriptions(sandbox)
    todo = [e for e in entries if e.get("readable_pql")
            and (descriptions.get(e["audience_id"]) or {}).get("rule_sha") != rule_sha(e)]
    if not todo:
        return descriptions
    logger.info(f"{len(todo)} audience(s) need a description.")
    client = ai_client()
    if not client:
        return descriptions
    import anthropic

    friendly, info = read_dictionary(sandbox)
    names = {e["audience_id"]: e.get("audience_name") or "" for e in entries}
    for i, e in enumerate(todo, 1):
        aid = e["audience_id"]
        try:
            text = ai_describe(client, ai_prompt(e, info.get(aid, {}), friendly, names))
        except anthropic.AuthenticationError:
            logger.error("Anthropic rejected the API key -- stopping the AI pass.")
            break
        except anthropic.RateLimitError:
            logger.error("Anthropic rate limit hit -- stopping the AI pass; "
                         "re-run to carry on from here.")
            break
        except anthropic.APIStatusError as ex:
            logger.warning(f"  [{i}/{len(todo)}] {e['audience_name']!r}: "
                           f"API error {ex.status_code}: {ex.message}")
            continue
        except anthropic.APIConnectionError as ex:
            logger.warning(f"  [{i}/{len(todo)}] {e['audience_name']!r}: "
                           f"network error: {ex}")
            continue
        if not text:
            logger.warning(f"  [{i}/{len(todo)}] {e['audience_name']!r}: "
                           f"no description returned.")
            continue
        descriptions[aid] = {"description": text, "rule_sha": rule_sha(e),
                             "named_from": e.get("audience_name") or "",
                             "by": AI_MODEL}
        save_descriptions(sandbox, descriptions)
        logger.info(f"  [{i}/{len(todo)}] {e['audience_name']!r} -> {text!r}")
    return descriptions


# ----------------------------------------------------------------------------
# Header terms and the new name
# ----------------------------------------------------------------------------
def build_plan(entries: list[dict], descriptions: dict, opts: dict) -> list[dict]:
    """One row per audience: RENAME / UNCHANGED / SKIP, with the reason."""
    rows = []
    for e in entries:
        described = descriptions.get(e.get("audience_id")) or {}
        row = {"id": e.get("audience_id") or "", "old": e.get("audience_name") or "",
               "new": "", "readable": e.get("readable_pql") or "",
               "action": "RENAME", "why": "", "result": "",
               "by": described.get("by", ""),
               "terms": {"prefix_a": "", "prefix_b": "", "description": ""}}
        if not row["readable"]:
            row["action"], row["why"] = "SKIP", "no rule (not a rule-based audience)"
        elif described.get("rule_sha") != rule_sha(e):
            row["action"], row["why"] = "SKIP", "no description yet (needs the AI pass)"
        else:
            row["terms"] = {"prefix_a": opts["prefix_a"], "prefix_b": opts["prefix_b"],
                            "description": tidy_description(described["description"])}
            row["new"] = NAME_FORMAT.format(**row["terms"])
        rows.append(row)

    # Two audiences can be given the same name; number the later ones.
    taken = {r["old"].casefold() for r in rows if r["action"] == "SKIP"}
    for r in rows:
        if r["action"] != "RENAME":
            continue
        if r["new"] == r["old"]:
            r["action"], r["why"] = "UNCHANGED", "already matches the convention"
        else:
            base, n = r["new"], 1
            while r["new"].casefold() in taken:
                n += 1
                r["new"] = f"{base} ({n})"
        taken.add((r["new"] if r["action"] == "RENAME" else r["old"]).casefold())
    return rows


# ----------------------------------------------------------------------------
# Apply (the only part that writes to AEP)
# ----------------------------------------------------------------------------
def http_error(e: Exception) -> str:
    if isinstance(e, urllib.error.HTTPError):
        return f"HTTP {e.code}: {dd.flatten_err(e.read().decode(errors='replace'))}"
    return f"{type(e).__name__}: {e}"


def is_development(token, conf, sandbox: str) -> tuple[bool, str]:
    """(ok, why-not). Only a sandbox AEP itself lists as type=development."""
    ok, found = dd.list_sandboxes(token, conf)
    if not ok:
        return False, f"could not list sandboxes ({found})"
    sb = next((s for s in found if s.get("name") == sandbox), None)
    if not sb:
        return False, f"no sandbox named {sandbox!r}"
    if sb.get("type") != "development":
        return False, f"{sandbox!r} is a {sb.get('type')} sandbox"
    return True, ""


def get_audience(token, conf, sandbox: str, audience_id: str) -> dict:
    body, _ = dd.http(f"{dd.AUDIENCES_URL}/{audience_id}",
                      headers=dd.aep_headers(token, conf, sandbox))
    return json.loads(body) or {}


def set_name(token, conf, sandbox: str, audience_id: str, name: str) -> None:
    """The one write this tool makes: PATCH an audience's name."""
    patch = json.dumps([{"op": "add", "path": "/name", "value": name}])
    dd.http(f"{dd.AUDIENCES_URL}/{audience_id}", method="PATCH",
            data=patch.encode("utf-8"),
            headers={**dd.aep_headers(token, conf, sandbox),
                     "Content-Type": "application/json"})


def rename_audience(token, conf, sandbox: str, row: dict, previous: dict) -> str:
    """Rename one audience; returns the result text for the workbook. The old
    name goes into the previous-names file BEFORE the rename is sent."""
    try:
        current = get_audience(token, conf, sandbox, row["id"])
    except Exception as e:
        return f"FAILED to read: {http_error(e)}"
    if current.get("isSystem"):
        return "SKIPPED: system audience"
    if (current.get("name") or "") != row["old"]:
        return "SKIPPED: name changed since the dictionary run"
    remembered = row["id"] in previous
    if not remembered:                  # keep the ORIGINAL across re-renames
        previous[row["id"]] = row["old"]
        save_previous(sandbox, previous)
    try:
        set_name(token, conf, sandbox, row["id"], row["new"])
    except Exception as e:
        if not remembered:              # never renamed, so nothing to restore
            del previous[row["id"]]
            save_previous(sandbox, previous)
        return f"FAILED: {http_error(e)}"
    return "RENAMED"


# ----------------------------------------------------------------------------
# The previous-names file (what --restore puts back)
# ----------------------------------------------------------------------------
def previous_path(sandbox: str) -> Path:
    return OUTPUT_DIR / f"audience_renamer_{sandbox}_previous_names.json"


def load_previous(sandbox: str) -> dict:
    """{audience id: the name it had before this tool first renamed it}."""
    path = previous_path(sandbox)
    if not path.exists():
        return {}
    return json.loads(path.read_text(encoding="utf-8"))


def save_previous(sandbox: str, previous: dict) -> None:
    """Rewritten after every change, so a run that dies half way still knows
    what it renamed. Removed once there is nothing left to put back."""
    path = previous_path(sandbox)
    if not previous:
        path.unlink(missing_ok=True)
        return
    OUTPUT_DIR.mkdir(exist_ok=True)
    path.write_text(json.dumps(previous, indent=2, ensure_ascii=False),
                    encoding="utf-8")


# ----------------------------------------------------------------------------
# Workbook
# ----------------------------------------------------------------------------
def write_xlsx(rows: list[dict], sandbox: str, sidecar: Path, opts: dict) -> Path | None:
    """Summary + Rename plan, in the Data Dictionary house style. One file per
    sandbox, overwritten each run."""
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Alignment, Font, PatternFill
        from openpyxl.utils import get_column_letter
    except ImportError:
        logger.error("openpyxl not installed -- nothing written "
                     "(pip install -r requirements.txt).")
        return None

    head_font = Font(bold=True, color="FFFFFF")
    head_fill = PatternFill("solid", fgColor=dd._HEADER_BG)
    note_font = Font(italic=True, color="666666")
    wrap = Alignment(wrap_text=True, vertical="top")
    top = Alignment(vertical="top")
    action_font = {"RENAME": Font(bold=True, color="006100"),
                   "SKIP": Font(color="808080"), "UNCHANGED": Font(color="808080")}

    def start(ws, title, note, span, height):
        ws["A1"] = dd.CONFIDENTIAL
        ws["A1"].font = Font(bold=True, size=11, color="C00000")
        ws["A2"] = title
        ws["A2"].font = Font(bold=True, size=14)
        ws["A3"] = note
        ws["A3"].font, ws["A3"].alignment = note_font, wrap
        ws.merge_cells(f"A3:{get_column_letter(span)}3")
        ws.row_dimensions[3].height = height

    wb = Workbook()
    # House rule: the file's author is Barry, not the library ('openpyxl').
    wb.properties.creator = wb.properties.lastModifiedBy = SCRIPT_AUTHOR
    mode = "APPLIED to AEP" if opts["apply"] else "DRY RUN - nothing changed in AEP"

    ws = wb.active
    ws.title = "Summary"
    start(ws, f"Audience renamer  -  {sandbox}  ({mode})",
          "Proposed audience names, built from the naming convention below. The "
          "description is written by AI from each audience's current name, its "
          "rule and the Data Dictionary's friendly names. Audiences with no "
          "rule are left alone. Old names are kept in "
          f"{previous_path(sandbox).name} for --restore.", 4, 48)
    count = lambda a: sum(1 for r in rows if r["action"] == a)
    facts = [("Sandbox", sandbox), ("Mode", mode),
             ("Naming convention", NAME_FORMAT),
             ("Prefix A", opts["prefix_a"]), ("Prefix B", opts["prefix_b"]),
             ("Dictionary sidecar", sidecar.name),
             ("Audiences in the sidecar", len(rows)),
             ("To rename", count("RENAME")),
             ("Unchanged", count("UNCHANGED")),
             ("Skipped: no rule",
              sum(1 for r in rows if r["why"].startswith("no rule"))),
             ("Skipped: no description yet",
              sum(1 for r in rows if r["why"].startswith("no description"))),
             ("Generated (UTC)", f"{datetime.now(timezone.utc):%Y-%m-%d %H:%M:%S}"),
             ("Generated by", f"{SCRIPT_NAME}.py v{SCRIPT_VERSION} ({SCRIPT_DATE})")]
    if opts["apply"]:
        facts.insert(11, ("Renamed in AEP",
                          sum(1 for r in rows if r["result"] == "RENAMED")))
    for i, (k, v) in enumerate(facts, 5):
        ws.cell(i, 1, k).font = Font(bold=True)
        ws.cell(i, 2, v).alignment = Alignment(horizontal="left")
    ws.column_dimensions["A"].width = 36
    ws.column_dimensions["B"].width = 70

    ps = wb.create_sheet("Rename plan")
    start(ps, "Rename plan -- one row per audience",
          "'Prefix A', 'Prefix B' and 'Description' are the header terms; 'New "
          "name' is those run through the naming convention. 'Described by' is "
          "what wrote the description. 'Result' is filled in on an --apply run.",
          6, 30)
    cols = ["Action", "Current name", "New name", "Prefix A", "Prefix B",
            "Description", "Described by", "Why skipped", "Result",
            "PQL (readable)", "Audience id"]
    for c, name in enumerate(cols, 1):
        cell = ps.cell(5, c, name)
        cell.font, cell.fill = head_font, head_fill
        cell.alignment = Alignment(vertical="center", wrap_text=True)
    ps.row_dimensions[5].height = 28
    order = {"RENAME": 0, "UNCHANGED": 1, "SKIP": 2}
    for i, r in enumerate(sorted(rows, key=lambda r: (order[r["action"]],
                                                      r["old"].casefold())), 6):
        t = r["terms"]
        vals = [r["action"], r["old"], r["new"], t["prefix_a"], t["prefix_b"],
                t["description"], r["by"], r["why"], r["result"],
                dd.fit_cell(r["readable"])[0], r["id"]]
        ps.row_dimensions[i].height = 30
        for c, v in enumerate(vals, 1):
            cell = ps.cell(i, c, v)
            cell.alignment = wrap if c in (2, 3, 6) else top
        ps.cell(i, 1).font = action_font[r["action"]]
        if r["result"].startswith("FAILED"):
            ps.cell(i, 9).font = Font(bold=True, color="C00000")
    for i, w in enumerate([12, 46, 60, 12, 12, 50, 18, 34, 30, 70, 38], 1):
        ps.column_dimensions[get_column_letter(i)].width = w
    ps.freeze_panes = "C6"
    if ps.max_row > 5:
        ps.auto_filter.ref = f"A5:{get_column_letter(len(cols))}{ps.max_row}"

    OUTPUT_DIR.mkdir(exist_ok=True)
    path = OUTPUT_DIR / f"audience_renamer_{sandbox}.xlsx"
    try:
        wb.save(path)
    except PermissionError:
        alt = path.with_name(path.stem + f" ({datetime.now():%H%M%S})" + path.suffix)
        logger.warning(f"{path.name} is open; writing {alt.name} instead.")
        path = alt
        wb.save(path)
    return path


# ----------------------------------------------------------------------------
# Main
# ----------------------------------------------------------------------------
def parse_args(argv):
    opts = {"name": None, "sandbox": DEFAULT_SANDBOX, "apply": False,
            "restore": False, "limit": None,
            "prefix_a": PREFIX_A, "prefix_b": PREFIX_B}
    for a in argv:
        if a.startswith("--sandbox="):
            opts["sandbox"] = a.split("=", 1)[1].strip()
        elif a.startswith("--prefix-a="):
            opts["prefix_a"] = a.split("=", 1)[1].strip()
        elif a.startswith("--prefix-b="):
            opts["prefix_b"] = a.split("=", 1)[1].strip()
        elif a.startswith("--limit="):
            opts["limit"] = int(a.split("=", 1)[1])
        elif a == "--apply":
            opts["apply"] = True
        elif a == "--restore":
            opts["restore"] = True
        elif a in ("-h", "--help"):
            print(__doc__)
            sys.exit(0)
        elif a.startswith("-"):
            logger.warning(f"Ignoring unknown option {a}")
        else:
            opts["name"] = a  # keyring service name
    return opts


def banner(sandbox: str, opts: dict):
    bar = ANSI["cyan"] + "=" * 72 + ANSI["reset"]
    if opts["restore"]:
        mode = f"{ANSI['red']}{ANSI['bold']}RESTORE -- puts the old names back in AEP"
    elif opts["apply"]:
        mode = f"{ANSI['red']}{ANSI['bold']}APPLY -- renames audiences in AEP"
    else:
        mode = f"{ANSI['dim']}dry run -- nothing is changed in AEP"
    print(bar)
    print(f"  {ANSI['bold']}{SCRIPT_NAME} v{SCRIPT_VERSION}{ANSI['reset']}   ({SCRIPT_DATE})")
    print(f"  by {SCRIPT_AUTHOR}")
    print(f"  {ANSI['bold']}Sandbox:{ANSI['reset']}  {ANSI['magenta']}{sandbox}{ANSI['reset']}")
    print(f"  {mode}{ANSI['reset']}")
    print(bar)


def connect(opts: dict):
    """(token, conf) for a write run, or None. Authenticates and refuses any
    sandbox that is not a development one."""
    try:
        service = aep_creds.resolve_service(opts["name"])
        conf = aep_creds.load_creds(service)
    except aep_creds.CredsError as e:
        logger.error(str(e))
        return None
    try:
        token = dd.authenticate(conf)
    except Exception as e:
        logger.error(f"IMS auth FAILED: {http_error(e)}")
        return None
    ok, why = is_development(token, conf, opts["sandbox"])
    if not ok:
        logger.error(f"Refusing to rename: {why}. Development sandboxes only.")
        return None
    return token, conf


def apply_plan(rows: list[dict], opts: dict) -> bool:
    """Rename every RENAME row. Returns False when it could not start (nothing
    was changed)."""
    sandbox = opts["sandbox"]
    conn = connect(opts)
    if not conn:
        return False
    token, conf = conn

    todo = [r for r in rows if r["action"] == "RENAME"]
    if opts["limit"] is not None:
        todo = todo[:opts["limit"]]
    previous = load_previous(sandbox)
    logger.info(f"Renaming {len(todo)} audience(s) in '{sandbox}'...")
    for r in todo:
        r["result"] = rename_audience(token, conf, sandbox, r, previous)
        level = logging.INFO if r["result"] == "RENAMED" else logging.WARNING
        logger.log(level, f"  {r['result']}: {r['old']!r} -> {r['new']!r}")
    done = sum(1 for r in todo if r["result"] == "RENAMED")
    logger.info(f"{done} of {len(todo)} renamed.")
    if previous:
        logger.info(f"Old names kept in {previous_path(sandbox)} "
                    f"({len(previous)}) -- --restore puts them back.")
    return True


def restore(opts: dict) -> None:
    """Rename every audience in the previous-names file back to its saved
    name. Each one is dropped from the file as it is put back, so whatever is
    left in the file afterwards is what still needs restoring."""
    sandbox = opts["sandbox"]
    previous = load_previous(sandbox)
    if not previous:
        logger.info(f"Nothing to restore: no {previous_path(sandbox).name}.")
        return
    conn = connect(opts)
    if not conn:
        return
    token, conf = conn

    logger.info(f"Restoring {len(previous)} audience name(s) in '{sandbox}'...")
    done = 0
    for audience_id, old in list(previous.items()):
        try:
            now = get_audience(token, conf, sandbox, audience_id).get("name") or ""
            if now != old:
                set_name(token, conf, sandbox, audience_id, old)
        except Exception as e:
            logger.warning(f"  FAILED: {old!r} ({audience_id}): {http_error(e)}")
            continue
        logger.info(f"  RESTORED: {now!r} -> {old!r}" if now != old
                    else f"  already {old!r}")
        del previous[audience_id]
        save_previous(sandbox, previous)
        done += 1
    logger.info(f"{done} restored, {len(previous)} still to restore.")


def main():
    opts = parse_args(sys.argv[1:])
    sandbox = opts["sandbox"]
    banner(sandbox, opts)

    if opts["restore"]:
        restore(opts)
        return

    sidecar = latest_sidecar(sandbox)
    if not sidecar:
        logger.error(f"No Data Dictionary sidecar for '{sandbox}' in {dd.PQL_DIR}. "
                     f"Run the dictionary first:  python data_dictionary_v3.py "
                     f"<service> --sandbox={sandbox}")
        return
    logger.info(f"Reading {sidecar.name}")
    if sidecar_date(sidecar) != f"{datetime.now():%Y-%m-%d}":
        logger.warning(f"That dictionary run is from {sidecar_date(sidecar) or '?'}, "
                       f"not today -- re-run the dictionary for current names.")

    entries = read_sidecar(sidecar)
    rows = build_plan(entries, describe_audiences(entries, sandbox), opts)
    n = {a: sum(1 for r in rows if r["action"] == a)
         for a in ("RENAME", "UNCHANGED", "SKIP")}
    logger.info(f"{len(rows)} audience(s): {n['RENAME']} to rename, "
                f"{n['UNCHANGED']} unchanged, {n['SKIP']} skipped.")
    for r in rows:
        if r["action"] == "RENAME":
            print(f"  {r['old']}\n    {ANSI['green']}-> {r['new']}{ANSI['reset']}")

    if opts["apply"]:
        if not apply_plan(rows, opts):
            return
    else:
        logger.info("Dry run. Add a credential name and --apply to rename in AEP.")

    xlsx = write_xlsx(rows, sandbox, sidecar, opts)
    if xlsx:
        logger.info(f"Wrote workbook {xlsx}")


if __name__ == "__main__":
    main()
