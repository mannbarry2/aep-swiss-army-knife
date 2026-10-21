#!/usr/bin/env python3
"""
role_permissions_grid.py  (AEP Swiss Army Knife)
================================================
The AEP role -> permission grid, and nothing else: no people, no members, no
audit trail. Built for process design -- the access-request form where a
consultant picks a ROLE, and the guidance notes that tell them which role
covers the job they already know they need to do ("manage datasets",
"publish orchestrated campaigns" ...).

Reads the Access Control API at org level (GETs only):
  roles                     name, description, type, sandboxes, labels, and
                            the permission sets each one grants
  permission-set catalogue  every permission AEP offers: category, name, what
                            it grants (resource: actions), description
  sandboxes                 names -> the titles the console shows

Output: ONE workbook, overwritten each run, named after the organisation it
describes:  output/AEP Role Permissions - <Client>.xlsx
  Which role    one row per permission in the catalogue -> every role that
                grants it. The guidance-notes table: filter the Permission
                column for the job, read off the role. Permissions that no
                role grants are flagged -- the form cannot offer them yet.
  Grid          permissions down, roles across, a tick where the role grants
                it: the console's permissions grid with every role side by side.
  Roles         every role: description, type, labels, permission count, and a
                tick per sandbox it reaches.
  Grid data     long form, one row per role x permission, for loading into the
                form itself.

The client name is never hard-coded here (it does not belong in this repo);
it comes from --client, else a 'client' key stored on the credential, else
the credential's own name.

Usage:
    python role_permissions_grid.py --client=Acme
    python role_permissions_grid.py aep-prod --client=Acme
"""

from __future__ import annotations

import sys
from collections import defaultdict
from datetime import datetime, timezone
from pathlib import Path

import aep_creds
import data_dictionary_v3 as dd
import permissions_audit as pa
from house_xlsx import Book, stable_name

SCRIPT_NAME = "role_permissions_grid"
SCRIPT_VERSION = "1.0.0"
SCRIPT_DATE = "2026-10-21"
SCRIPT_AUTHOR = "Barry Mann (barrymann.com)"

OUTPUT_DIR = Path(__file__).resolve().parent / "output"
TICK = "✓"
NONE = "(no role grants this)"
logger = pa.logger


def parse_args(argv):
    o = {"credential": "aep-prod", "client": ""}
    for a in argv:
        k, _, v = a.partition("=")
        if k == "--client":
            o["client"] = v.strip()
        elif a in ("-h", "--help"):
            print(__doc__)
            sys.exit(0)
        elif not a.startswith("-"):
            o["credential"] = a
        else:
            logger.warning(f"ignoring {a}")
    return o


def grants_text(p: dict) -> str:
    return "; ".join(f"{x.get('resource')}: {','.join(x.get('actions') or [])}"
                     for x in (p.get("permissions") or []))


def main():
    o = parse_args(sys.argv[1:])
    service = aep_creds.pick_service(o["credential"])
    conf = aep_creds.load_creds(service)
    client = o["client"] or aep_creds.client_label(conf, service)
    if not o["client"] and client == service:
        logger.warning(f"No --client given: the file will be named after the "
                       f"credential ({service!r}), not the organisation.")
    token = dd.authenticate(conf)["access_token"]
    h = dd.aep_headers(token, conf)          # org-level: no sandbox header

    print("=" * 72)
    print(f"  {SCRIPT_NAME} v{SCRIPT_VERSION}   ({SCRIPT_DATE})  by {SCRIPT_AUTHOR}")
    print(f"  Role -> permission grid, no people.  {client}  "
          f"(org {conf['org_id']}, read-only)")
    print("=" * 72)

    roles = sorted(pa.fetch_roles(h), key=lambda r: (r.get("name") or "").lower())
    psets = pa.fetch_permission_sets(h)
    ok, sbx = dd.list_sandboxes(token, conf)
    title_of = {s.get("name"): s.get("title") or s.get("name")
                for s in (sbx if ok else [])}
    if not ok:
        logger.warning(f"sandbox titles unavailable ({sbx}); showing names.")
    logger.info(f"{len(roles)} role(s), {len(psets)} permission set(s) in the catalogue")

    # Permission sets a role grants that the catalogue does not describe
    # (retired or custom) still belong on the grid -- under their bare id.
    for r in roles:
        for ps in r.get("permissionSets") or []:
            psets.setdefault(ps, {"id": ps, "name": ps,
                                  "category": "(not in catalogue)"})

    granted_by = defaultdict(list)           # permission-set id -> [role names]
    for r in roles:
        for ps in r.get("permissionSets") or []:
            granted_by[ps].append(r.get("name") or r["id"])

    perms = sorted(psets.values(), key=lambda p: ((p.get("category") or "").lower(),
                                                  (p.get("name") or p["id"]).lower()))
    orphans = [p for p in perms if not granted_by.get(p["id"])]

    def sandboxes(r):
        names = r.get("sandboxes") or []
        return ", ".join(title_of.get(n, n) for n in names) or "(all)"

    def labels(r):
        return ", ".join((r.get("subjectAttributes") or {}).get("labels") or [])

    stamp = f"{datetime.now(timezone.utc):%Y-%m-%d %H:%M} UTC"
    book = Book(f"AEP role permissions  -  {client}",
                "Which role grants which permission. Roles only: no people, no "
                "members. For the access-request form and its guidance notes.",
                generator=f"{SCRIPT_NAME}.py v{SCRIPT_VERSION}")

    # ---- Which role? (the guidance-notes table) -----------------------------
    book.sheet("Which role",
               ["Category", "Permission", "Roles that grant it", "Roles",
                "What it grants (resource: actions)", "Description", "Permission id"],
               [[p.get("category") or "", p.get("name") or p["id"],
                 "\n".join(granted_by.get(p["id"], [])) or NONE,
                 len(granted_by.get(p["id"], [])), grants_text(p),
                 p.get("description") or "", p["id"]] for p in perms],
               widths=[26, 40, 50, 8, 60, 70, 40], wrap_cols=(3, 5, 6),
               row_height=30, tab_colour="2D7D46", red_when={3: NONE},
               note="Find the job in the Permission column (filter it), then read "
                    "off the role(s) that grant it. A role's reach is also limited "
                    "to its sandboxes -- see the Roles tab.",
               facts=[("Organisation", client), ("Org id", conf["org_id"]),
                      ("Roles", len(roles)), ("Permissions in the catalogue", len(perms)),
                      ("Granted by no role", len(orphans)), ("Generated", stamp)])

    # ---- Grid: permissions down, roles across --------------------------------
    role_names = [r.get("name") or r["id"] for r in roles]
    has = {(ps, r.get("name") or r["id"]) for r in roles
           for ps in r.get("permissionSets") or []}
    ws = book.sheet("Grid",
                    ["Category", "Permission", "Roles"] + role_names,
                    [[p.get("category") or "", p.get("name") or p["id"],
                      len(granted_by.get(p["id"], []))]
                     + [TICK if (p["id"], rn) in has else "" for rn in role_names]
                     for p in perms],
                    widths=[24, 40, 7] + [5] * len(role_names), row_height=15,
                    tab_colour="1F4E78",
                    note="Permissions down, roles across: a tick where the role "
                         "grants the permission. Role names run up the header row.")
    if ws is not None:
        from openpyxl.styles import Alignment
        hr = book._tables[-1][1]
        ws.row_dimensions[hr].height = 220
        for c in range(4, 4 + len(role_names)):
            ws.cell(hr, c).alignment = Alignment(text_rotation=90, vertical="bottom",
                                                 horizontal="center")
            for rr in range(hr + 1, hr + 1 + len(perms)):
                ws.cell(rr, c).alignment = Alignment(horizontal="center")
        ws.freeze_panes = ws.cell(hr + 1, 4)

    # ---- Roles: one row each, a tick per sandbox ------------------------------
    sbx_names = sorted({n for r in roles for n in (r.get("sandboxes") or [])},
                       key=lambda n: title_of.get(n, n).lower())
    book.sheet("Roles",
               ["Role", "Description", "Type", "Labels", "Permissions"]
               + [title_of.get(n, n) for n in sbx_names],
               [[r.get("name") or r["id"], r.get("description") or "",
                 r.get("roleType") or "", labels(r), len(r.get("permissionSets") or [])]
                + [TICK if n in (r.get("sandboxes") or []) else "" for n in sbx_names]
                for r in roles],
               widths=[44, 70, 14, 24, 11] + [12] * len(sbx_names),
               wrap_cols=(2,), row_height=30, tab_colour="7030A0",
               note="Every role, with the sandboxes it reaches (a tick per "
                    "sandbox) and its attribute-based-access labels.")

    # ---- Grid data: long form for the form itself -----------------------------
    book.sheet("Grid data",
               ["Role", "Category", "Permission", "Permission id", "Sandboxes",
                "Labels", "What it grants (resource: actions)"],
               [[r.get("name") or r["id"], psets[ps].get("category") or "",
                 psets[ps].get("name") or ps, ps, sandboxes(r), labels(r),
                 grants_text(psets[ps])]
                for r in roles for ps in sorted(r.get("permissionSets") or [],
                                                key=lambda x: ((psets[x].get("category") or "").lower(),
                                                               (psets[x].get("name") or x).lower()))],
               widths=[44, 26, 40, 40, 50, 24, 60], tab_colour="C55A11",
               note="One row per role x permission: the grid in a shape a form "
                    "or a lookup can load directly.")

    path = book.save(OUTPUT_DIR / f"{stable_name(f'AEP Role Permissions - {client}')}.xlsx")
    print()
    logger.info(f"{len(roles)} roles x {len(perms)} permissions; "
                f"{len(orphans)} permission(s) granted by no role.")
    if path:
        logger.info(f"Wrote {path}")


if __name__ == "__main__":
    main()
