#!/usr/bin/env python3
"""
aep_usage_log.py  (AEP Swiss Army Knife)
========================================
Who is actively using AEP, per human user, over the last N days -- read from
the Audit events API. Replaces the UI export, which caps at 10,000 events and
is flooded by a tech account deleting from the AJO Suppression List.

This is ACTIVITY, not sign-ins: an audit event is something a user did (read
a schema, edit a journey, run a segment job). A user who logged in and did
nothing leaves no trace here.

Version 1.2.0 (2026-10-02)

Changelog
---------
1.2.0  2026-10-02  Survives a stall: every request has a hard 2-minute limit
                   and is retried, and each finished day is cached under
                   output/aep_usage_log_cache/<sandbox>/, so a stopped run
                   resumes and a later run only fetches new days. Asks
                   Windows not to idle-sleep while it runs: three 90-day
                   runs had frozen mid-way when the machine slept.
1.1.0  2026-10-02  What people do: every asset type is mapped to a work area
                   (Content, Journeys & campaigns, Audiences, Data ...);
                   by_user gains a suggested role, the share of events that
                   changed something, events per area and top changes; new
                   by_role sheet. raw_events is now opt-in (--raw): it was
                   ~95% of a 40 MB workbook.
1.0.0  2026-10-02  First version. Adaptive windows under the API's 1,000-event
                   cap; tech-account noise excluded server-side and counted
                   separately; Core events only; workbook with raw_events,
                   by_user, by_user_month, by_assetType_action and
                   excluded_noise.

How the API behaves (measured against prod on 2026-10-02 -- the docs differ)
---------------------------------------------------------------------------
* A query returns at most 1,000 events, and page.totalElements is capped at
  1,000 too: it reads 1,000 for a whole day and for a single busy hour. So the
  tool never trusts a count of 1,000. Each day is queried as one window; any
  window that hits the cap is split (1 h, 10 min, 1 min, 10 s, 1 s) until
  every piece is under it. A 1-second window still at the cap is reported as
  INCOMPLETE rather than silently short.
* Timestamps must be UTC with milliseconds (2026-09-30T00:00:00.000Z); at most
  two timestamp filters.
* The user filter is 'user' (not 'userEmail'). It takes ONE value, so only one
  account can be excluded server-side. '!=' works for user, assetType and
  action; each key can appear once.
* Core vs Enhanced: there is no eventType field. Each returned event IS the
  Core event; its Enhanced events ride inside it under 'enhancedEvents' with
  the same requestId. Counting top-level events therefore counts Core only.

How the noise is handled
------------------------
One prod day held 44,924 events, 41,047 of them from a single tech account.
The tool finds that account (the busiest @techacct.adobe.com user in a sample
of the most recent day) and its bulk asset type (AJO Suppression List), then
runs three streams per day:

  A  everyone except that account (user!=)   fetched in full
  B  that account, other asset types         fetched in full (Work Orders etc)
  C  that account, its bulk asset type        COUNTED only, never fetched

Stream C is ~4 million rows over 90 days -- more than an Excel sheet holds --
so it is counted (cheap: limit=3 probes) and shown in excluded_noise. Any
other @techacct.adobe.com events that come through A and B are dropped from
the people sheets client-side and counted in excluded_noise too, unless
--no-exclude-techacct is given (stream C is always counted only).

Credentials come from the shared aep_creds layer (keyring service 'aep-prod'
by default; manage with credential_validator_v2.py). Read-only: every call is
a GET.

What people do
--------------
Every asset type is mapped to a work area (AREAS below). Each person's
suggested role comes from the area where they CHANGE things (create, update,
publish, activate ...), not where they look: Profile views alone are a quarter
of all activity and say little about anyone's job. Someone who changes almost
nothing is a "Viewer". The mapping and the role names are plain constants --
edit them when the team's real job titles are known.

Cache: output/aep_usage_log_cache/<sandbox>/<YYYY-MM-DD>.json.gz, one per
finished day. Delete the folder to force a full re-fetch.

Output: output/aep_usage_log_<sandbox>_<YYYYMMDD>.xlsx  (add --raw for the
one-row-per-event sheet; it makes the file ~20x larger)

Usage:
    python aep_usage_log.py                         # aep-prod, prod, last 90 days
    python aep_usage_log.py aep-prod --sandbox=prod --days=30
    python aep_usage_log.py --days=7 --no-noise-count      # skip stream C
    python aep_usage_log.py --noise-user=<id>@techacct.adobe.com \\
        --noise-asset-type="AJO Suppression List"
"""

from __future__ import annotations

import argparse
import concurrent.futures
import gzip
import json
import logging
import ssl
import sys
import time
import urllib.error
import urllib.parse
import urllib.request
import uuid
from collections import Counter, defaultdict
from datetime import datetime, timedelta, timezone
from pathlib import Path

import aep_creds  # keyring-backed credential store (replaces creds/*.json)

SCRIPT_NAME = "aep_usage_log"
SCRIPT_VERSION = "1.2.0"
SCRIPT_DATE = "2026-10-02"
SCRIPT_AUTHOR = "Barry Mann (barrymann.com)"

SCRIPT_DIR = Path(__file__).resolve().parent
OUTPUT_DIR = SCRIPT_DIR / "output"

IMS_URL = "https://ims-na1.adobelogin.com/ims/token"
AUDIT_URL = "https://platform.adobe.io/data/foundation/audit/events"
DEFAULT_SCOPES = (
    "openid,AdobeID,read_organizations,"
    "additional_info.projectedProductContext,session"
)
DEFAULT_SERVICE = "aep-prod"

API_CAP = 1000          # max page size AND the ceiling on page.totalElements
# How a window at the cap is cut up, coarsest first.
SPLITS = (timedelta(hours=1), timedelta(minutes=10), timedelta(minutes=1),
          timedelta(seconds=10), timedelta(seconds=1))
TECHACCT = "@techacct.adobe.com"
MAX_ATTEMPTS = 6        # per request, for 429 / 5xx / network errors / hangs
REQUEST_DEADLINE = 120  # seconds; a request still running after this is abandoned

CONFIDENTIAL = "STRICTLY CONFIDENTIAL"
_HEADER_BG = "1F4E78"   # the Data Dictionary's header blue
XL_MAX_ROWS = 1_048_576

# ----------------------------------------------------------------------------
# ANSI / logging - matches credential_validator.py style
# ----------------------------------------------------------------------------
if sys.platform == "win32":
    try:
        import ctypes
        kernel32 = ctypes.windll.kernel32
        h = kernel32.GetStdHandle(-11)
        mode = ctypes.c_ulong()
        if kernel32.GetConsoleMode(h, ctypes.byref(mode)):
            kernel32.SetConsoleMode(h, mode.value | 0x0004)
    except Exception:
        pass

ANSI = {
    "reset": "\033[0m", "bold": "\033[1m", "dim": "\033[2m",
    "red": "\033[31m", "green": "\033[32m", "yellow": "\033[33m",
    "blue": "\033[34m", "magenta": "\033[35m", "cyan": "\033[36m",
}
LEVEL_COLOR = {
    "DEBUG": ANSI["dim"], "INFO": ANSI["green"],
    "WARNING": ANSI["yellow"],
    "ERROR": ANSI["red"] + ANSI["bold"],
    "CRITICAL": ANSI["red"] + ANSI["bold"],
}


class ColoredFormatter(logging.Formatter):
    def format(self, record):
        color = LEVEL_COLOR.get(record.levelname, "")
        ts = self.formatTime(record, "%H:%M:%S")
        return (
            f"{ANSI['dim']}{ts}{ANSI['reset']} "
            f"{color}[{record.levelname:<7}]{ANSI['reset']} "
            f"{record.getMessage()}"
        )


_handler = logging.StreamHandler(sys.stdout)
_handler.setFormatter(ColoredFormatter())
logging.basicConfig(level=logging.INFO, handlers=[_handler])
logger = logging.getLogger(SCRIPT_NAME)
SSL_CTX = ssl._create_unverified_context()


# ----------------------------------------------------------------------------
# HTTP / IMS
# ----------------------------------------------------------------------------
def http(url, method="GET", headers=None, data=None, timeout=90):
    req = urllib.request.Request(url, data=data, headers=headers or {}, method=method)
    with urllib.request.urlopen(req, context=SSL_CTX, timeout=timeout) as r:
        return r.read()


def flatten_err(text: str, limit: int = 200) -> str:
    return " ".join((text or "").split())[:limit]


def authenticate(conf) -> str:
    """client_credentials grant against Adobe IMS -> access token string."""
    payload = urllib.parse.urlencode({
        "grant_type": "client_credentials",
        "client_id": conf["client_id"],
        "client_secret": conf["client_secret"],
        "scope": conf.get("scopes") or DEFAULT_SCOPES,
    }).encode("utf-8")
    body = http(conf.get("oauth_url") or IMS_URL, method="POST", data=payload,
                headers={"Content-Type": "application/x-www-form-urlencoded"})
    return json.loads(body)["access_token"]


class AuditClient:
    """GETs against the audit events endpoint, with backoff on 429 / 5xx."""

    def __init__(self, token: str, conf: dict, sandbox: str):
        self.token, self.conf, self.sandbox = token, conf, sandbox
        self.requests = 0
        # Requests run on a worker thread so a hung one can be walked away
        # from; an abandoned thread finishes (or dies) on its own.
        self._pool = concurrent.futures.ThreadPoolExecutor(max_workers=4)

    def get(self, params: list[tuple[str, object]]) -> dict:
        url = f"{AUDIT_URL}?{urllib.parse.urlencode(params)}"
        for attempt in range(1, MAX_ATTEMPTS + 1):
            headers = {
                "Authorization": f"Bearer {self.token}",
                "x-api-key": self.conf.get("api_key") or self.conf["client_id"],
                "x-gw-ims-org-id": self.conf["org_id"],
                "x-sandbox-name": self.sandbox,
                "x-request-id": str(uuid.uuid4()),
                "Accept": "application/json",
            }
            self.requests += 1
            try:
                job = self._pool.submit(http, url, "GET", headers)
                return json.loads(job.result(timeout=REQUEST_DEADLINE))
            except urllib.error.HTTPError as e:
                body = flatten_err(e.read().decode(errors="replace"))
                if e.code != 429 and e.code < 500:
                    raise RuntimeError(f"HTTP {e.code}: {body}") from None
                wait = _retry_after(e) or min(2 ** attempt, 60)
                why = f"HTTP {e.code}"
            except (urllib.error.URLError, TimeoutError, ConnectionError, OSError) as e:
                wait, why = min(2 ** attempt, 60), type(e).__name__
            if attempt == MAX_ATTEMPTS:
                raise RuntimeError(f"gave up after {attempt} attempts ({why})")
            logger.warning(f"    {why}; retrying in {wait}s "
                           f"(attempt {attempt}/{MAX_ATTEMPTS})")
            time.sleep(wait)
        raise AssertionError("unreachable")


def _retry_after(e: urllib.error.HTTPError) -> int | None:
    try:
        return max(1, int(e.headers.get("Retry-After", "")))
    except (TypeError, ValueError):
        return None


# ----------------------------------------------------------------------------
# Windows under the cap
# ----------------------------------------------------------------------------
def iso_ms(dt: datetime) -> str:
    """The only timestamp form the API accepts: UTC with milliseconds."""
    return dt.strftime("%Y-%m-%dT%H:%M:%S.") + f"{dt.microsecond // 1000:03d}Z"


def window_params(a: datetime, b: datetime, filters: list[str], limit: int):
    return ([("limit", limit),
             ("property", f"timestamp>={iso_ms(a)}"),
             ("property", f"timestamp<{iso_ms(b)}")]
            + [("property", f) for f in filters])


def walk_windows(client: AuditClient, a: datetime, b: datetime,
                 filters: list[str], fetch: bool, stats: dict,
                 splits=SPLITS) -> list[dict]:
    """Every event in [a, b) matching the filters (fetch=True), or just the
    count in stats['counted'] (fetch=False). A window at the 1,000 cap is cut
    into the next finer step and each piece walked in turn."""
    page = client.get(window_params(a, b, filters, API_CAP if fetch else 3))
    total = (page.get("page") or {}).get("totalElements") or 0
    if total >= API_CAP and splits:
        stats["splits"] += 1
        out, cur = [], a
        while cur < b:
            nxt = min(cur + splits[0], b)
            out += walk_windows(client, cur, nxt, filters, fetch, stats, splits[1:])
            cur = nxt
        return out
    if total >= API_CAP:
        stats["incomplete"].append(f"{iso_ms(a)} -> {iso_ms(b)}")
        logger.warning(f"    a 1-second window is still at the {API_CAP} cap "
                       f"({iso_ms(a)}); some events in it are missing.")
    if not fetch:
        stats["counted"] += total
        return []
    events = list((page.get("_embedded") or {}).get("events") or [])
    # One page holds the whole window once it is under the cap; page on anyway
    # in case the server ever serves less than 'limit' per page.
    while len(events) < total and page.get("queryId"):
        page = client.get([("queryId", page["queryId"]), ("start", len(events)),
                           ("limit", API_CAP)])
        more = (page.get("_embedded") or {}).get("events") or []
        if not more:
            break
        events += more
    return events


def day_windows(days: int) -> list[tuple[datetime, datetime]]:
    """1-day UTC windows covering the last `days` days up to now; today is
    the last, partial window."""
    now = datetime.now(timezone.utc).replace(microsecond=0)
    today = now.replace(hour=0, minute=0, second=0)
    out = []
    for i in range(days - 1, -1, -1):
        a = today - timedelta(days=i)
        out.append((a, min(a + timedelta(days=1), now)))
    return out


# ----------------------------------------------------------------------------
# Noise
# ----------------------------------------------------------------------------
def detect_noise(client: AuditClient, windows) -> tuple[str, str]:
    """(noise_user, noise_asset_type) from one page of the most recent full
    day: the busiest tech account, if it is at least half the sample, and its
    commonest asset type. ('', '') when nothing dominates."""
    a, b = windows[-2] if len(windows) > 1 else windows[-1]
    page = client.get(window_params(a, b, [], API_CAP))
    sample = (page.get("_embedded") or {}).get("events") or []
    users = Counter(e.get("userEmail") or "" for e in sample)
    top = next(((u, n) for u, n in users.most_common() if u.endswith(TECHACCT)), None)
    if not top or top[1] < len(sample) / 2:
        return "", ""
    assets = Counter(e.get("assetType") or "" for e in sample if e.get("userEmail") == top[0])
    return top[0], assets.most_common(1)[0][0]


# ----------------------------------------------------------------------------
# Events -> rows
# ----------------------------------------------------------------------------
RAW_COLUMNS = ["timestamp", "userEmail", "assetType", "action", "status",
               "assetName", "assetId", "permissionResource", "permissionType",
               "actorType", "failureCode", "sandboxName", "region",
               "userIpAddresses", "enhancedEvents", "requestId", "id", "authId",
               "aoTaskId", "sandboxId", "imsOrgId", "version"]


def parse_ts(value: str) -> datetime | None:
    """'2026-09-30T23:59:58.100+0000' -> naive UTC datetime (Excel has no tz)."""
    try:
        dt = datetime.strptime(value, "%Y-%m-%dT%H:%M:%S.%f%z")
    except (TypeError, ValueError):
        return None
    return dt.astimezone(timezone.utc).replace(tzinfo=None)


def raw_row(e: dict) -> list:
    """One raw_events row: every Core field; the nested Enhanced events are
    reduced to how many there were."""
    vals = []
    for col in RAW_COLUMNS:
        v = e.get(col)
        if col == "timestamp":
            v = parse_ts(v) or v
        elif col == "userIpAddresses":
            v = ", ".join(map(str, v or []))
        elif col == "enhancedEvents":
            v = len(v or [])
        vals.append(v if v is not None else "")
    return vals


# ----------------------------------------------------------------------------
# One day, and the per-day cache
# ----------------------------------------------------------------------------
# A finished UTC day's audit events never change, so each one is kept in a
# small gzip file. A run that stops part-way picks up where it left off, and a
# later run only fetches the days it has not seen. Events are cached BEFORE the
# tech-account split, so --exclude-techacct can change without a re-fetch; the
# noise settings are part of the key, so changing them does re-fetch.
CACHE_KEEP_FIELDS = [c for c in RAW_COLUMNS if c != "enhancedEvents"]


def cache_dir(sandbox: str) -> Path:
    return OUTPUT_DIR / "aep_usage_log_cache" / sandbox


def load_day(sandbox: str, day: str, key: dict) -> dict | None:
    path = cache_dir(sandbox) / f"{day}.json.gz"
    try:
        with gzip.open(path, "rt", encoding="utf-8") as fh:
            got = json.load(fh)
    except (OSError, ValueError):
        return None
    return got if got.get("key") == key else None


def save_day(sandbox: str, day: str, key: dict, got: dict) -> None:
    folder = cache_dir(sandbox)
    folder.mkdir(parents=True, exist_ok=True)
    tmp = folder / f"{day}.json.gz.tmp"
    with gzip.open(tmp, "wt", encoding="utf-8") as fh:
        json.dump({**got, "key": key}, fh)
    tmp.replace(folder / f"{day}.json.gz")      # never leave a half-written day


def fetch_day(client: AuditClient, a: datetime, b: datetime, noise_user: str,
              noise_asset: str, count_noise: bool, stats: dict) -> dict:
    """Everything the report needs for one day: the fetched events of streams
    A and B (slimmed to the report's fields, Enhanced events reduced to a
    count) and the stream-C count."""
    streams = [([f"user!={noise_user}"] if noise_user else [], True)]
    if noise_user:
        streams.append(([f"user=={noise_user}"]
                        + ([f"assetType!={noise_asset}"] if noise_asset else []), True))
        if count_noise:
            streams.append(([f"user=={noise_user}", f"assetType=={noise_asset}"], False))
    events, counted = [], 0
    day_stats = {"splits": 0, "counted": 0, "incomplete": []}
    for filters, fetch in streams:
        day_stats["counted"] = 0
        for e in walk_windows(client, a, b, filters, fetch, day_stats):
            slim = {k: e.get(k) for k in CACHE_KEEP_FIELDS if e.get(k) is not None}
            slim["enhancedEvents"] = [None] * len(e.get("enhancedEvents") or [])
            events.append(slim)
        counted += 0 if fetch else day_stats["counted"]
    stats["splits"] += day_stats["splits"]
    return {"events": events, "counted": counted,
            "incomplete": day_stats["incomplete"]}


# ----------------------------------------------------------------------------
# Collect
# ----------------------------------------------------------------------------
def collect(client: AuditClient, opts) -> dict:
    windows = day_windows(opts.days)
    noise_user, noise_asset = opts.noise_user, opts.noise_asset_type
    if noise_user is None:
        noise_user, noise_asset_found = detect_noise(client, windows)
        noise_asset = noise_asset or noise_asset_found
    elif not noise_asset:
        noise_asset = ""
    if noise_user:
        logger.info(f"Noise account: {noise_user}  (bulk asset type: "
                    f"{noise_asset or 'none'})")
    else:
        logger.info("No dominant tech account found; fetching everything.")

    kept, noise = [], Counter()          # noise[(date, user, assetType, how)]
    stats = {"splits": 0, "counted": 0, "incomplete": []}
    count_noise = bool(noise_user and noise_asset and not opts.no_noise_count)
    cache_key = {"noise_user": noise_user, "noise_asset": noise_asset,
                 "count_noise": count_noise}
    n_cached = 0
    for n, (a, b) in enumerate(windows, 1):
        t0, req0 = time.time(), client.requests
        day = a.date().isoformat()
        whole_day = b - a == timedelta(days=1)
        got = load_day(opts.sandbox, day, cache_key) if whole_day else None
        if got is None:
            got = fetch_day(client, a, b, noise_user, noise_asset, count_noise, stats)
            if whole_day:          # past days never change; today is still filling
                save_day(opts.sandbox, day, cache_key, got)
        else:
            n_cached += 1
        stats["incomplete"] += got["incomplete"]
        day_kept = day_noise = 0
        if got["counted"]:
            noise[(day, noise_user, noise_asset, "counted only")] += got["counted"]
            day_noise += got["counted"]
        for e in got["events"]:
            user = e.get("userEmail") or ""
            if opts.exclude_techacct and user.endswith(TECHACCT):
                noise[(day, user, e.get("assetType") or "", "fetched")] += 1
                day_noise += 1
            else:
                kept.append(e)
                day_kept += 1
        flag = (f"  {ANSI['red']}INCOMPLETE windows: "
                f"{len(got['incomplete'])}{ANSI['reset']}" if got["incomplete"] else "")
        source = ("from cache" if client.requests == req0 else
                  f"{client.requests - req0} requests, {time.time() - t0:.1f}s")
        logger.info(f"[{n:>3}/{len(windows)}] {day}: {day_kept:>6,} kept, "
                    f"{day_noise:>7,} noise  ({source}){flag}")
    if n_cached:
        logger.info(f"{n_cached} day(s) came from {cache_dir(opts.sandbox)}.")
    return {"kept": kept, "noise": noise, "windows": windows,
            "noise_user": noise_user, "noise_asset": noise_asset,
            "incomplete": stats["incomplete"], "splits": stats["splits"]}


# ----------------------------------------------------------------------------
# What people do: asset type -> work area -> suggested role
# ----------------------------------------------------------------------------
AREAS = {
    "Content": ["AJO Fragment", "AJO Content Template", "AJO Saved Expression",
                "AJO Message Preset", "Simulated User"],
    "Journeys & campaigns": ["Journey", "Campaign", "Orchestrated Campaign",
                             "Experiment", "Journey Event", "Journey Custom Action",
                             "Journey Fragment", "AJO Saved Condition"],
    "Decisioning": ["Ranking Strategy", "ExD Offer Item", "ExD Decision Policy",
                    "ExD Collections", "ExD Selection Strategy"],
    "Audiences": ["Segment", "Segment Job", "Composition", "Destination",
                  "Profiles Export"],
    "Data": ["Dataset", "Schema", "Class", "Field Group", "Source Data Flow",
             "Identity Namespace", "Datastreams", "Merge Policy",
             "Profile Settings", "Dule Policy"],
    "Query": ["Query", "Query Template", "Scheduled Query"],
    "Profile look-up": ["Profile", "Identity Graph"],
    "Admin": ["Role", "Account", "Audits", "CMK Config", "Package", "Work Order",
              "AJO Subdomain", "AJO Seedlist"],
}
AREA_OF = {t: area for area, types in AREAS.items() for t in types}
AREA_NAMES = list(AREAS) + ["Other"]
# The job a person whose changes are mostly in an area is probably doing.
ROLE_OF = {
    "Content": "Content executive",
    "Journeys & campaigns": "Journey / campaign manager",
    "Decisioning": "Offer decisioning",
    "Audiences": "Audience builder",
    "Data": "Data engineer",
    "Query": "Analyst (Query Service)",
    "Profile look-up": "Profile support",
    "Admin": "Platform admin",
    "Other": "Other",
}
LOOK_ACTIONS = {"View", "Read"}       # everything else changed or ran something
VIEWER_BELOW = 0.05                   # under 5% changes -> a Viewer
FEW_CHANGES = 20                      # too few changes to name a role on
SECOND_ROLE_FROM = 0.30               # a 2nd area with 30%+ of changes is named too


def area_of(asset_type: str) -> str:
    return AREA_OF.get(asset_type, "Other")


def suggested_role(changes: Counter, looks: Counter) -> str:
    """A best guess at someone's job from where they change things."""
    n_change, n_look = sum(changes.values()), sum(looks.values())
    if n_change < FEW_CHANGES or n_change / (n_change + n_look) < VIEWER_BELOW:
        where = (looks + changes).most_common(1)
        return f"Viewer ({where[0][0]})" if where else "Viewer"
    ranked = changes.most_common()
    role = ROLE_OF[ranked[0][0]]
    if len(ranked) > 1 and ranked[1][1] / n_change >= SECOND_ROLE_FROM:
        role += f" + {ROLE_OF[ranked[1][0]]}"
    return role


# ----------------------------------------------------------------------------
# Summaries
# ----------------------------------------------------------------------------
def summarise(kept: list[dict]) -> dict:
    by_user = defaultdict(lambda: {"events": 0, "days": set(), "first": None,
                                   "last": None, "assets": Counter(),
                                   "changes": Counter(), "looks": Counter(),
                                   "areas": Counter(), "what": Counter()})
    by_month = defaultdict(Counter)
    by_asset_action = defaultdict(lambda: {"events": 0, "users": set()})
    for e in kept:
        user = e.get("userEmail") or "(no user)"
        ts = parse_ts(e.get("timestamp"))
        u = by_user[user]
        u["events"] += 1
        asset, action = e.get("assetType") or "", e.get("action") or ""
        u["assets"][asset] += 1
        u["areas"][area_of(asset)] += 1
        if action in LOOK_ACTIONS:
            u["looks"][area_of(asset)] += 1
        else:
            u["changes"][area_of(asset)] += 1
            u["what"][f"{action} {asset}"] += 1
        if ts:
            u["days"].add(ts.date())
            u["first"] = min(u["first"] or ts, ts)
            u["last"] = max(u["last"] or ts, ts)
            by_month[user][ts.strftime("%Y-%m")] += 1
        k = (e.get("assetType") or "", e.get("action") or "")
        by_asset_action[k]["events"] += 1
        by_asset_action[k]["users"].add(user)
    for d in by_user.values():
        d["role"] = suggested_role(d["changes"], d["looks"])
    return {"by_user": by_user, "by_month": by_month,
            "by_asset_action": by_asset_action}


# ----------------------------------------------------------------------------
# Workbook
# ----------------------------------------------------------------------------
def write_xlsx(result: dict, opts, conf: dict, path: Path) -> Path:
    """Write-only workbook (raw_events can run to hundreds of thousands of
    rows), in the Data Dictionary style: confidential banner, title, italic
    note, blue filterable header, frozen panes, fixed widths."""
    from openpyxl import Workbook
    from openpyxl.cell import WriteOnlyCell
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    wb = Workbook(write_only=True)
    # House rule: the file's author is Barry, not the library ('openpyxl').
    wb.properties.creator = wb.properties.lastModifiedBy = SCRIPT_AUTHOR
    head_font, head_fill = Font(bold=True, color="FFFFFF"), PatternFill("solid", fgColor=_HEADER_BG)
    HEADER_ROW = 5

    def sheet(name, title, note, cols, widths, rows, formats=None):
        ws = wb.create_sheet(name)
        for i, w in enumerate(widths, 1):
            ws.column_dimensions[get_column_letter(i)].width = w
        ws.freeze_panes = f"A{HEADER_ROW + 1}"
        c = WriteOnlyCell(ws, CONFIDENTIAL)
        c.font = Font(bold=True, size=11, color="C00000")
        ws.append([c])
        c = WriteOnlyCell(ws, title)
        c.font = Font(bold=True, size=14)
        ws.append([c])
        c = WriteOnlyCell(ws, note)
        c.font = Font(italic=True, color="666666")
        c.alignment = Alignment(wrap_text=False, vertical="top")
        ws.append([c])
        ws.append([])
        header = []
        for name_ in cols:
            c = WriteOnlyCell(ws, name_)
            c.font, c.fill = head_font, head_fill
            header.append(c)
        ws.append(header)
        n = 0
        for row in rows:
            if formats:
                row = [_fmt(ws, v, formats.get(i)) for i, v in enumerate(row)]
            ws.append(row)
            n += 1
        ws.auto_filter.ref = f"A{HEADER_ROW}:{get_column_letter(len(cols))}{HEADER_ROW + max(n, 1)}"
        return n

    def _fmt(ws, value, number_format):
        if not number_format:
            return value
        c = WriteOnlyCell(ws, value)
        c.number_format = number_format
        return c

    kept, noise = result["kept"], result["noise"]
    s = summarise(kept)
    a, b = result["windows"][0][0], result["windows"][-1][1]
    span = f"{a:%Y-%m-%d} to {b:%Y-%m-%d %H:%M} UTC ({opts.days} days)"
    DT = "yyyy-mm-dd hh:mm:ss"

    # ---- summary -----------------------------------------------------------
    noise_total = sum(noise.values())
    facts = [
        ("What this is", "ACTIVITY per user from the AEP audit log -- things "
                         "people did, NOT sign-ins. A user who logged in and "
                         "did nothing does not appear."),
        ("Org", conf["org_id"]), ("Sandbox", opts.sandbox), ("Period", span),
        ("People (users with activity)", len(s["by_user"])),
        ("Activity events kept", len(kept)),
        ("Tech-account events excluded", noise_total),
        ("Noise account (excluded server-side)", result["noise_user"] or "(none)"),
        ("Its bulk asset type (counted only)", result["noise_asset"] or "(none)"),
        ("Tech accounts kept?", "no" if opts.exclude_techacct else
         "yes, except the noise account's bulk asset type"),
        ("Windows split to stay under the API cap", result["splits"]),
        ("INCOMPLETE windows (still at the cap at 1 s)", len(result["incomplete"])),
        ("Event type", "Core only (Enhanced events are nested in each Core "
                       "event and are not counted separately)"),
        ("Generated (UTC)", f"{datetime.now(timezone.utc):%Y-%m-%d %H:%M:%S}"),
        ("Generated by", f"{SCRIPT_NAME}.py v{SCRIPT_VERSION} ({SCRIPT_DATE})"),
    ]
    if opts.no_noise_count and result["noise_user"]:
        facts.insert(7, ("NOTE", "--no-noise-count: the noise account's bulk "
                                 "events were NOT counted"))
    sheet("summary", f"AEP activity log  -  {opts.sandbox}  ({span})",
          "Who is actively using AEP, per user. Activity, not sign-ins.",
          ["Fact", "Value"], [44, 110], facts)

    # ---- raw_events (opt-in with --raw: ~95% of the file) ---------------
    if opts.raw:
        raw = sorted(kept, key=lambda e: e.get("timestamp") or "")
        if len(raw) > XL_MAX_ROWS - 10:
            logger.warning(f"{len(raw):,} events is more than a sheet holds; "
                           f"raw_events keeps the newest {XL_MAX_ROWS - 10:,}.")
            raw = raw[-(XL_MAX_ROWS - 10):]
        sheet("raw_events", f"Raw activity events  -  {opts.sandbox}",
              "One row per Core audit event kept (tech accounts excluded). "
              "'enhancedEvents' is how many Enhanced events rode inside it. "
              "Timestamps are UTC.",
              RAW_COLUMNS,
              [20, 36, 26, 14, 10, 36, 30, 26, 16, 10, 12, 10, 8, 18, 10, 34, 38,
               38, 14, 38, 34, 8],
              (raw_row(e) for e in raw), {0: DT})

    # ---- by_user -----------------------------------------------------------
    users = sorted(s["by_user"].items(), key=lambda kv: -kv[1]["events"])
    sheet("by_user", f"What each person does  -  {span}",
          "Busiest first. 'Suggested role' is a best guess from the work area "
          "where the person CHANGES things (create, update, publish, "
          "activate...); 'Viewer' means they mostly look. 'Changes %' is the "
          "share of their events that changed or ran something. The area "
          "columns count all their events. Times are UTC.",
          ["userEmail", "Suggested role", "Total events", "Changes %",
           "Distinct days active", "First seen", "Last seen"] + AREA_NAMES
          + ["Top 3 changes", "Top 3 assetTypes"],
          [36, 44, 12, 10, 12, 18, 18] + [11] * len(AREA_NAMES) + [70, 60],
          ([u, d["role"], d["events"],
            sum(d["changes"].values()) / d["events"] if d["events"] else 0,
            len(d["days"]), d["first"], d["last"]]
           + [d["areas"].get(a, 0) for a in AREA_NAMES]
           + [", ".join(f"{w} ({n:,})" for w, n in d["what"].most_common(3)),
              ", ".join(f"{t} ({n:,})" for t, n in d["assets"].most_common(3))]
           for u, d in users), {3: "0%", 5: DT, 6: DT})

    # ---- by_role -----------------------------------------------------------
    roles = defaultdict(list)
    for u, d in users:
        roles[d["role"]].append((u, d))
    cutoff = b.replace(tzinfo=None) - timedelta(days=30)
    sheet("by_role", "People by suggested role",
          "How many people do each kind of work. 'Active last 30 days' counts "
          "those seen in the final 30 days of the period.",
          ["Suggested role", "People", "Active last 30 days", "Events",
           "People (busiest first)"],
          [44, 10, 18, 12, 140],
          ([r, len(ps), sum(1 for _, d in ps if d["last"] and d["last"] >= cutoff),
            sum(d["events"] for _, d in ps), ", ".join(u for u, _ in ps)]
           for r, ps in sorted(roles.items(), key=lambda kv: -len(kv[1]))))

    # ---- by_user_month -----------------------------------------------------
    months = sorted({m for c in s["by_month"].values() for m in c})
    sheet("by_user_month", "Activity by user and month",
          "Events per user per calendar month (UTC). The first and last "
          "months are partial.",
          ["userEmail"] + months + ["Total"], [40] + [12] * len(months) + [12],
          ([u] + [s["by_month"][u].get(m, 0) for m in months] + [d["events"]]
           for u, d in users))

    # ---- by_assetType_action -----------------------------------------------
    aa = sorted(s["by_asset_action"].items(), key=lambda kv: -kv[1]["events"])
    sheet("by_assetType_action", "Activity by asset type and action",
          "What people did. 'Users' is how many distinct people did it.",
          ["assetType", "action", "Events", "Users"], [36, 16, 12, 10],
          ([t, act, d["events"], len(d["users"])] for (t, act), d in aa))

    # ---- excluded_noise ----------------------------------------------------
    sheet("excluded_noise", "Excluded tech-account events",
          "Kept in view so the suppression-list deletes stay visible. 'counted "
          "only' rows were counted, never fetched (millions of rows); "
          "'fetched' rows were pulled and dropped from the people sheets.",
          ["Date (UTC)", "userEmail", "assetType", "Events", "How"],
          [12, 56, 30, 12, 14],
          ([d, u, t, n, how] for (d, u, t, how), n in sorted(noise.items())))

    if result["incomplete"]:
        sheet("incomplete_windows", "Windows still at the API cap",
              "Each of these 1-second windows held 1,000+ events; only the "
              "first 1,000 were read. Counts for that second are a floor.",
              ["Window (UTC)"], [60], ([w] for w in result["incomplete"]))

    path.parent.mkdir(parents=True, exist_ok=True)
    try:
        wb.save(path)
    except PermissionError:
        alt = path.with_name(path.stem + f" ({datetime.now():%H%M%S})" + path.suffix)
        logger.warning(f"{path.name} is open; writing {alt.name} instead.")
        wb.save(alt)
        path = alt
    return path


# ----------------------------------------------------------------------------
# Main
# ----------------------------------------------------------------------------
def parse_args(argv):
    p = argparse.ArgumentParser(
        description="AEP activity (not sign-ins) per user, from the audit log.")
    p.add_argument("service", nargs="?", default=DEFAULT_SERVICE,
                   help=f"keyring credential service (default {DEFAULT_SERVICE})")
    p.add_argument("--days", type=int, default=90, help="days back (default 90)")
    p.add_argument("--sandbox", default="prod", help="sandbox (default prod)")
    p.add_argument("--exclude-techacct", default=True,
                   action=argparse.BooleanOptionalAction,
                   help="drop @techacct.adobe.com users from the people sheets "
                        "(default on)")
    p.add_argument("--noise-user", default=None,
                   help="tech account to exclude server-side (default: detect)")
    p.add_argument("--noise-asset-type", default=None,
                   help="its bulk asset type, counted only (default: detect)")
    p.add_argument("--raw", action="store_true",
                   help="add the raw_events sheet (one row per event; ~20x the size)")
    p.add_argument("--no-noise-count", action="store_true",
                   help="do not count the noise account's bulk events (faster)")
    opts = p.parse_args(argv)
    if opts.days < 1:
        p.error("--days must be at least 1")
    return opts


def keep_awake(on: bool) -> None:
    """Ask Windows not to idle-sleep while the run is going (on=True), and
    hand that back afterwards (on=False). Three 90-day runs froze mid-way when
    the machine went to sleep. Changes no power settings; a closed lid still
    sleeps. A no-op off Windows."""
    if sys.platform != "win32":
        return
    try:
        import ctypes
        ES_CONTINUOUS, ES_SYSTEM_REQUIRED = 0x80000000, 0x00000001
        ctypes.windll.kernel32.SetThreadExecutionState(
            ES_CONTINUOUS | (ES_SYSTEM_REQUIRED if on else 0))
    except Exception:
        pass


def banner(conf, opts):
    bar = ANSI["cyan"] + "=" * 72 + ANSI["reset"]
    print(bar)
    print(f"  {ANSI['bold']}{SCRIPT_NAME} v{SCRIPT_VERSION}{ANSI['reset']}   ({SCRIPT_DATE})")
    print(f"  by {SCRIPT_AUTHOR}")
    print(f"  {ANSI['dim']}AEP ACTIVITY per user from the audit log (not sign-ins). "
          f"Read-only.{ANSI['reset']}")
    print(f"  {ANSI['bold']}Org:{ANSI['reset']}      {ANSI['magenta']}{conf['org_id']}{ANSI['reset']}")
    print(f"  {ANSI['bold']}Sandbox:{ANSI['reset']}  {ANSI['magenta']}{opts.sandbox}{ANSI['reset']}"
          f"   last {opts.days} day(s)")
    print(bar)


def main():
    opts = parse_args(sys.argv[1:])
    print(aep_creds.source_banner(), file=sys.stderr)
    try:
        conf = aep_creds.load_creds(aep_creds.resolve_service(opts.service))
    except aep_creds.CredsError as e:
        logger.error(str(e))
        return
    banner(conf, opts)
    try:
        token = authenticate(conf)
    except Exception as e:
        logger.error(f"IMS auth FAILED: {type(e).__name__}: {e}")
        return
    client = AuditClient(token, conf, opts.sandbox)

    t0 = time.time()
    keep_awake(True)
    try:
        result = collect(client, opts)
    except RuntimeError as e:
        logger.error(f"Audit API: {e}")
        return
    finally:
        keep_awake(False)
    logger.info(f"{len(result['kept']):,} activity events kept, "
                f"{sum(result['noise'].values()):,} tech-account events "
                f"excluded; {client.requests:,} requests in "
                f"{(time.time() - t0) / 60:.1f} min.")
    if result["incomplete"]:
        logger.warning(f"{len(result['incomplete'])} window(s) INCOMPLETE -- "
                       f"see the incomplete_windows sheet.")

    path = OUTPUT_DIR / f"aep_usage_log_{opts.sandbox}_{datetime.now():%Y%m%d}.xlsx"
    path = write_xlsx(result, opts, conf, path)
    logger.info(f"Wrote {path}")


if __name__ == "__main__":
    main()
